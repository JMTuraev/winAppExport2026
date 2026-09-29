"""Хатлар — bo'limning chiquvchi/kiruvchi xatlari reyestri va «xotira»si (faqat ofis tarmog'ida, internetsiz).

Ma'lumot: data\\xatlar\\xatlar.db (WAL) + fayllar data\\xatlar\\<file_path> (telegramdan\\files\\…, qolda\\YYYY\\MM\\…).
Har xat — bitta karta (sarlavha, qisqa mazmun, raqam, sana, kimga, qaysi xatga javob, asosiy raqamlar) va istalgancha
bog'lanish (letter_links): korxona (INN) · soha · tuman · davlat · kiruvchi raqam. Qidiruv — FTS5, kirill/lotin farqsiz
(asl matn + norm() transliteratsiyasi bir indeksda). Import (Telegram, Claude) va qo'lda qo'shish bir xil formatda yozadi."""
from __future__ import annotations
import io, json, os, re, sqlite3, threading, time, zipfile, datetime as dt
from pathlib import Path
from app import db

_LOCK = threading.Lock()
MAX_FILE = 80 * 1024 * 1024
MAX_LINKS = 60
KINDS = {"chiquvchi_xat": "Чиқувчи хат", "kiruvchi_xat": "Кирувчи хат", "malumotnoma": "Маълумотнома", "bayonnoma": "Баённома",
         "hisobot": "Даврий ҳисобот", "ilova": "Илова", "korxona_hujjati": "Корхона ҳужжати", "boshqa": "Бошқа", "shovqin": "Шовқин"}
SCOPES = {"korxona": "Корхона", "soha": "Соҳа", "tuman": "Туман", "davlat": "Давлат", "umumiy": "Умумий"}
LINK_KINDS = ("company", "sector", "district", "country", "reply")
EDIT_FIELDS = ("title", "summary", "kind", "doc_type", "number", "date", "sender", "recipient", "scope", "note", "signer", "executor")

SCHEMA = """
CREATE TABLE IF NOT EXISTS letters (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  source TEXT NOT NULL DEFAULT 'manual', src_msg_id INTEGER, direction TEXT NOT NULL DEFAULT 'chiquvchi',
  file_path TEXT, file_name TEXT, file_size INTEGER, tg_date TEXT, group_id TEXT,
  role TEXT DEFAULT 'main', parent_id INTEGER,
  kind TEXT, doc_type TEXT, number TEXT, date TEXT, sender TEXT, recipient TEXT, title TEXT, summary TEXT,
  scope TEXT, ocr_quality TEXT, confidence TEXT, note TEXT, signer TEXT, executor TEXT,
  flags TEXT, facts TEXT, topics TEXT, reply_to TEXT, companies_unmatched TEXT,
  hidden INTEGER DEFAULT 0, reviewed_by TEXT, reviewed_at TEXT, created_by TEXT, created_at TEXT, updated_at TEXT,
  UNIQUE(source, src_msg_id)
);
CREATE TABLE IF NOT EXISTS letter_links (
  letter_id INTEGER NOT NULL, kind TEXT NOT NULL, ref TEXT NOT NULL, label TEXT, role TEXT,
  UNIQUE(letter_id, kind, ref)
);
CREATE INDEX IF NOT EXISTS ix_links_ref ON letter_links(kind, ref);
CREATE TABLE IF NOT EXISTS letter_text (letter_id INTEGER PRIMARY KEY, text TEXT);
CREATE VIRTUAL TABLE IF NOT EXISTS letters_fts USING fts5(title, body, meta, tokenize = 'unicode61 remove_diacritics 2');
"""

# ------------------------------------------------------------ kirill ↔ lotin (import skripti candidates.py bilan bir xil)
CYR = {"а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "j", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l",
       "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "x", "ц": "ts", "ч": "ch", "ш": "sh",
       "щ": "sh", "ъ": "", "ы": "i", "ь": "", "э": "e", "ю": "yu", "я": "ya", "ў": "o", "қ": "q", "ғ": "g", "ҳ": "h", "і": "i"}

def norm(s: str) -> str:
    s = (s or "").lower()
    s = "".join(CYR.get(ch, ch) for ch in s)
    s = re.sub(r"[‘’`ʻʼ'´]", "", s)
    s = s.replace("kh", "x").replace("textile", "tekstil").replace("textil", "tekstil")
    s = s.replace("buxara", "buxoro").replace("buhara", "buxoro")
    s = re.sub(r"c(?!h)(?=[aoulrtk])", "k", s)
    s = re.sub(r"c(?!h)(?=[eiy])", "s", s)
    s = s.replace("ts", "s")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return f" {s.strip()} "

# ------------------------------------------------------------ db
def base() -> Path: return db.BASE / "xatlar"
def db_path() -> Path: return base() / "xatlar.db"
def now() -> str: return dt.datetime.now().isoformat(timespec="seconds")

def connect() -> sqlite3.Connection:
    p = db_path(); p.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(p), timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL"); con.execute("PRAGMA synchronous=NORMAL")
    con.executescript(SCHEMA)
    return con

def init() -> None:
    con = connect(); con.close()

def _j(v, default=None):
    if v in (None, ""): return [] if default is None else default
    try: return json.loads(v)
    except Exception: return [] if default is None else default

# ------------------------------------------------------------ export.db ma'lumotnomalari (tumanlar, korxonalar)
_REF: dict = {"t": 0}

def refs() -> dict:
    """Tumanlar va korxonalar (nom, tuman, soha) — 5 daqiqalik kesh."""
    if time.time() - _REF["t"] < 300: return _REF
    lc = db.connect()
    try:
        _REF["districts"] = {r[0]: r[1] for r in lc.execute("SELECT code, name_uz FROM districts")}
        _REF["companies"] = {r[0]: {"inn": r[0], "name": r[1], "district": r[2], "tarmoq": r[3]} for r in lc.execute(
            "SELECT inn, coalesce(nullif(full_name,''), name), district_code, tarmoq FROM companies WHERE coalesce(excluded,0)=0")}
        try: _REF["countries"] = [r[0] for r in lc.execute("SELECT name FROM countries ORDER BY name")]
        except sqlite3.OperationalError: _REF["countries"] = []
        _REF["sectors"] = [r[0] for r in lc.execute("""SELECT trim(tarmoq) FROM companies WHERE tarmoq IS NOT NULL AND trim(tarmoq)<>''
                                                       GROUP BY trim(tarmoq) ORDER BY count(*) DESC""")]
    finally: lc.close()
    _REF["cores"] = None
    _REF["t"] = time.time()
    return _REF

def _label(kind: str, ref: str, label: str | None) -> str:
    if kind == "district": return refs()["districts"].get(ref, label or ref)
    if kind == "company":
        c = refs()["companies"].get(ref)
        return (c["name"] if c else None) or label or ref
    return label or ref

# ------------------------------------------------------------ ro'yxat / qidiruv
def _fts_query(q: str) -> str | None:
    toks = [t for t in norm(q).split() if t]
    toks = [t for t in toks if len(t) >= 2 or t.isdigit()]
    if not toks: return None
    return " AND ".join(f'"{t}"*' for t in toks[:12])

def _snippet(text: str, q: str, width: int = 230) -> str | None:
    """Asl matndan (kirill bo'lsa kirillda) so'rov so'zlari uchragan joy — belgilash UI'da."""
    toks = [t for t in norm(q).split() if len(t) >= 2]
    if not toks or not text: return None
    t = re.sub(r"\s+", " ", text)
    for m in re.finditer(r"[^\s«»“”\"(),.:;!?]+", t):
        w = norm(m.group(0)).strip()
        if w and any(w.startswith(x) or (" " + x) in (" " + w) for x in toks):
            a = max(0, m.start() - width // 3); b = min(len(t), a + width)
            return ("…" if a else "") + t[a:b].strip() + ("…" if b < len(t) else "")
    return None

def _links_for(con, ids: list[int]) -> dict[int, list[dict]]:
    out: dict[int, list[dict]] = {i: [] for i in ids}
    if not ids: return out
    for part in range(0, len(ids), 500):
        chunk = ids[part:part + 500]
        for r in con.execute(f"SELECT letter_id, kind, ref, label, role FROM letter_links WHERE letter_id IN ({','.join('?' * len(chunk))})", chunk):
            out[r["letter_id"]].append({"kind": r["kind"], "ref": r["ref"], "label": _label(r["kind"], r["ref"], r["label"]), "role": r["role"]})
    return out

def _row(r: sqlite3.Row, links: list[dict] | None = None, full: bool = False) -> dict:
    d = {k: r[k] for k in r.keys()}
    for k in ("flags", "facts", "topics", "reply_to", "companies_unmatched"): d[k] = _j(d.get(k))
    d["ext"] = (Path(d["file_name"]).suffix.lower().lstrip(".") if d.get("file_name") else "")
    d["kind_name"] = KINDS.get(d.get("kind") or "", d.get("kind") or "")
    d["reviewed"] = bool(d.get("reviewed_at"))
    if not full and d.get("summary") and len(d["summary"]) > 420: d["summary"] = d["summary"][:420].rstrip() + "…"
    if links is not None: d["links"] = links
    return d

def list_letters(con, q: dict) -> dict:
    """Filtrlar: q, kind (vergul), scope, year, company (INN), sector, district, country, chk=1 (tekshirilmagan),
    att=1 (ilova/nusxalar ham), hidden=1 (arxiv ham), sort=date|new, limit, offset."""
    where, args = [], []
    text = (q.get("q") or "").strip()
    fq = _fts_query(text) if text else None
    if not q.get("hidden"): where.append("l.hidden=0")
    if not q.get("att") and not text: where.append("coalesce(l.role,'main')='main'")
    kinds = [k for k in (q.get("kind") or "").split(",") if k]
    if kinds: where.append(f"l.kind IN ({','.join('?' * len(kinds))})"); args += kinds
    if q.get("scope"): where.append("l.scope=?"); args.append(q["scope"])
    if q.get("year"): where.append("substr(coalesce(l.date, l.tg_date, l.created_at),1,4)=?"); args.append(str(q["year"]))
    if q.get("chk"): where.append("l.confidence='tekshirish' AND l.reviewed_at IS NULL")
    if q.get("source"): where.append("l.source=?"); args.append(q["source"])
    for key, kind in (("company", "company"), ("sector", "sector"), ("district", "district"), ("country", "country"), ("reply", "reply")):
        v = q.get(key)
        if v:
            # xatning o'zi yoki uning ilovasi bog'langan bo'lsa ham
            where.append(f"""(EXISTS (SELECT 1 FROM letter_links k WHERE k.letter_id=l.id AND k.kind='{kind}' AND k.ref=?)
                          OR EXISTS (SELECT 1 FROM letters c JOIN letter_links k ON k.letter_id=c.id WHERE c.parent_id=l.id AND k.kind='{kind}' AND k.ref=?))""")
            args += [v, v]
    order = "coalesce(l.date, substr(l.tg_date,1,10), substr(l.created_at,1,10)) DESC, l.id DESC"
    if q.get("sort") == "new": order = "l.id DESC"
    limit = max(1, min(200, int(q.get("limit") or 50))); offset = max(0, int(q.get("offset") or 0))
    wsql = (" AND ".join(where)) or "1"
    if fq:
        try:
            base_sql = f"FROM letters_fts f JOIN letters l ON l.id=f.rowid WHERE letters_fts MATCH ? AND {wsql}"
            total = con.execute(f"SELECT count(*) {base_sql}", [fq] + args).fetchone()[0]
            rows = con.execute(f"""SELECT l.* {base_sql} ORDER BY (coalesce(l.role,'main')='main') DESC, bm25(letters_fts, 10.0, 1.0, 4.0), l.id DESC
                                   LIMIT ? OFFSET ?""", [fq] + args + [limit, offset]).fetchall()
        except sqlite3.OperationalError:
            fq = None
    if not fq:
        if text:            # FTS ishlamasa — sarlavha/mazmun bo'yicha oddiy qidiruv
            like = "%" + text.lower() + "%"
            wsql += " AND (lower(coalesce(l.title,'')) LIKE ? OR lower(coalesce(l.summary,'')) LIKE ? OR coalesce(l.number,'') LIKE ? OR coalesce(l.file_name,'') LIKE ?)"
            args += [like, like, like, like]
        total = con.execute(f"SELECT count(*) FROM letters l WHERE {wsql}", args).fetchone()[0]
        rows = con.execute(f"SELECT l.* FROM letters l WHERE {wsql} ORDER BY {order} LIMIT ? OFFSET ?", args + [limit, offset]).fetchall()
    ids = [r["id"] for r in rows]
    links = _links_for(con, ids)
    kids = {}
    if ids:
        for r in con.execute(f"SELECT parent_id, count(*) FROM letters WHERE parent_id IN ({','.join('?' * len(ids))}) GROUP BY parent_id", ids):
            kids[r[0]] = r[1]
    parents = {}
    pids = [r["parent_id"] for r in rows if r["parent_id"]]
    if pids:
        for r in con.execute(f"SELECT id, title FROM letters WHERE id IN ({','.join('?' * len(pids))})", pids): parents[r[0]] = r[1]
    texts = {}
    if text and ids:
        for r in con.execute(f"SELECT letter_id, substr(text,1,60000) FROM letter_text WHERE letter_id IN ({','.join('?' * len(ids))})", ids): texts[r[0]] = r[1]
    items = []
    for r in rows:
        d = _row(r, links[r["id"]])
        d["n_att"] = kids.get(r["id"], 0)
        if r["parent_id"]: d["parent_title"] = parents.get(r["parent_id"])
        if text:
            d["snippet"] = _snippet((r["summary"] or "") + " " + (texts.get(r["id"]) or ""), text)
        items.append(d)
    return {"items": items, "total": total, "limit": limit, "offset": offset}

def facets(con) -> dict:
    f = {"total": con.execute("SELECT count(*) FROM letters WHERE hidden=0").fetchone()[0],
         "hidden": con.execute("SELECT count(*) FROM letters WHERE hidden=1").fetchone()[0],
         "chk": con.execute("SELECT count(*) FROM letters WHERE hidden=0 AND confidence='tekshirish' AND reviewed_at IS NULL").fetchone()[0],
         "kinds": {r[0]: r[1] for r in con.execute("SELECT kind, count(*) FROM letters WHERE hidden=0 GROUP BY kind")},
         "scopes": {r[0]: r[1] for r in con.execute("SELECT scope, count(*) FROM letters WHERE hidden=0 GROUP BY scope")},
         "years": [r[0] for r in con.execute("""SELECT DISTINCT substr(coalesce(date, tg_date, created_at),1,4) y FROM letters
                                                 WHERE hidden=0 AND y IS NOT NULL ORDER BY y DESC""")],
         "last": con.execute("SELECT max(coalesce(date, substr(tg_date,1,10))) FROM letters WHERE hidden=0").fetchone()[0]}
    R = refs()
    f["districts"] = [{"code": r[0], "name": R["districts"].get(r[0], r[0]), "n": r[1]} for r in con.execute(
        "SELECT ref, count(DISTINCT letter_id) FROM letter_links WHERE kind='district' GROUP BY ref ORDER BY 2 DESC")]
    f["sectors"] = [{"name": r[0], "n": r[1]} for r in con.execute(
        "SELECT ref, count(DISTINCT letter_id) FROM letter_links WHERE kind='sector' GROUP BY ref ORDER BY 2 DESC")]
    f["countries"] = [{"name": r[0], "n": r[1]} for r in con.execute(
        "SELECT ref, count(DISTINCT letter_id) FROM letter_links WHERE kind='country' GROUP BY ref ORDER BY 2 DESC")]
    f["kind_names"] = KINDS; f["scope_names"] = SCOPES
    return f

# ------------------------------------------------------------ bitta xat
def get_letter(con, lid: int) -> dict:
    r = con.execute("SELECT * FROM letters WHERE id=?", (lid,)).fetchone()
    if not r: raise ValueError("Xat topilmadi")
    d = _row(r, _links_for(con, [lid])[lid], full=True)
    t = con.execute("SELECT text FROM letter_text WHERE letter_id=?", (lid,)).fetchone()
    d["text"] = (t[0] or "")[:200000] if t else ""
    d["children"] = [_row(c, None) for c in con.execute("SELECT * FROM letters WHERE parent_id=? ORDER BY id", (lid,))]
    d["parent"] = None
    if r["parent_id"]:
        p = con.execute("SELECT * FROM letters WHERE id=?", (r["parent_id"],)).fetchone()
        if p: d["parent"] = _row(p, None)
    fp = _file(r)
    d["file_ok"] = bool(fp and fp.is_file())
    # shu xatga javob bo'lgan / shu raqamni tilga olgan boshqa xatlar
    refs_ = [x["ref"] for x in d["links"] if x["kind"] == "reply"]
    d["related"] = []
    if refs_:
        d["related"] = [_row(x, None) for x in con.execute(f"""SELECT DISTINCT l.* FROM letters l JOIN letter_links k ON k.letter_id=l.id
            WHERE k.kind='reply' AND k.ref IN ({','.join('?' * len(refs_))}) AND l.id<>? AND l.hidden=0 ORDER BY l.id DESC LIMIT 12""", refs_ + [lid])]
    return d

def _file(r) -> Path | None:
    """data\\xatlar ichidagi fayl (nisbiy yo'l; «..» va absolyut yo'l — rad etiladi)."""
    if not r or not r["file_path"]: return None
    rel = Path(str(r["file_path"]).replace("\\", "/"))
    if rel.is_absolute() or ".." in rel.parts or rel.drive: return None
    return base() / rel

def file_path(con, lid: int) -> tuple[Path, str]:
    r = con.execute("SELECT id, file_path, file_name FROM letters WHERE id=?", (lid,)).fetchone()
    f = _file(r)
    if not f or not f.is_file(): raise ValueError("Fayl topilmadi")
    return f, r["file_name"] or f.name

# ------------------------------------------------------------ korxona profili
def company_letters(con, inn: str) -> dict:
    R = refs(); c = R["companies"].get(inn) or {}
    direct = con.execute("""SELECT DISTINCT coalesce(p.id, l.id) AS id FROM letter_links k JOIN letters l ON l.id=k.letter_id
                            LEFT JOIN letters p ON p.id=l.parent_id AND coalesce(l.role,'main')<>'main'
                            WHERE k.kind='company' AND k.ref=? AND l.hidden=0""", (inn,)).fetchall()
    dids = [r[0] for r in direct]
    rows = con.execute(f"SELECT * FROM letters WHERE id IN ({','.join('?' * len(dids)) or 'NULL'}) ORDER BY coalesce(date, substr(tg_date,1,10)) DESC, id DESC", dids).fetchall() if dids else []
    links = _links_for(con, [r["id"] for r in rows])
    roles = {r[0]: r[1] for r in con.execute("SELECT letter_id, role FROM letter_links WHERE kind='company' AND ref=?", (inn,))}
    items = []
    for r in rows:
        d = _row(r, links[r["id"]]); d["my_role"] = (roles.get(r["id"]) or "tilga_olingan") if r["id"] in roles else "ilovada"; items.append(d)
    related, why = [], []
    conds, args = [], []
    if c.get("tarmoq"): conds.append("(k.kind='sector' AND k.ref=?)"); args.append(c["tarmoq"].strip()); why.append(c["tarmoq"].strip())
    if c.get("district"): conds.append("(k.kind='district' AND k.ref=?)"); args.append(c["district"]); why.append(R["districts"].get(c["district"], c["district"]))
    if conds:
        rr = con.execute(f"""SELECT DISTINCT l.* FROM letter_links k JOIN letters l ON l.id=k.letter_id
                             WHERE ({' OR '.join(conds)}) AND l.hidden=0 AND coalesce(l.role,'main')='main'
                             {('AND l.id NOT IN (' + ','.join('?' * len(dids)) + ')') if dids else ''}
                             ORDER BY coalesce(l.date, substr(l.tg_date,1,10)) DESC, l.id DESC LIMIT 60""", args + dids).fetchall()
        rl = _links_for(con, [r["id"] for r in rr])
        related = [_row(r, rl[r["id"]]) for r in rr]
    return {"inn": inn, "items": items, "related": related, "related_by": why}

def company_count(con, inn: str) -> int:
    return con.execute("""SELECT count(DISTINCT coalesce(l.parent_id, l.id)) FROM letter_links k JOIN letters l ON l.id=k.letter_id
                          WHERE k.kind='company' AND k.ref=? AND l.hidden=0""", (inn,)).fetchone()[0]

# ------------------------------------------------------------ yozish (tahrir, tasdiq, arxiv, qo'shish)
def _clean_links(links) -> list[tuple]:
    out, seen = [], set()
    for x in (links or [])[:MAX_LINKS]:
        k = str(x.get("kind") or ""); ref = str(x.get("ref") or "").strip()
        if k not in LINK_KINDS or not ref: continue
        if k == "company" and not re.fullmatch(r"\d{9}", ref): continue
        if (k, ref) in seen: continue
        seen.add((k, ref))
        out.append((k, ref[:200], (str(x.get("label") or ref))[:300], (x.get("role") or None)))
    return out

def _reindex(con, lid: int) -> None:
    r = con.execute("SELECT * FROM letters WHERE id=?", (lid,)).fetchone()
    if not r: return
    links = con.execute("SELECT kind, ref, label FROM letter_links WHERE letter_id=?", (lid,)).fetchall()
    t = con.execute("SELECT text FROM letter_text WHERE letter_id=?", (lid,)).fetchone()
    text = (t[0] if t else "") or ""
    topics = _j(r["topics"]); unm = _j(r["companies_unmatched"])
    meta = " ".join(str(x or "") for x in [r["number"], r["doc_type"], r["file_name"], " ".join(topics), " ".join(unm), r["sender"], r["recipient"],
                                           r["signer"], r["executor"], " ".join(f"{k['ref']} {k['label'] or ''}" for k in links),
                                           " ".join(_label(k["kind"], k["ref"], k["label"]) for k in links if k["kind"] == "district")])
    title = r["title"] or ""; body = (r["summary"] or "") + "\n" + text
    try:
        con.execute("DELETE FROM letters_fts WHERE rowid=?", (lid,))
        con.execute("INSERT INTO letters_fts(rowid, title, body, meta) VALUES(?,?,?,?)",
                    (lid, title + " ‖ " + norm(title), body + "\n‖\n" + norm(body), meta + " ‖ " + norm(meta)))
    except sqlite3.OperationalError: pass

def update_letter(con, lid: int, body: dict, user: dict) -> dict:
    r = con.execute("SELECT * FROM letters WHERE id=?", (lid,)).fetchone()
    if not r: raise ValueError("Xat topilmadi")
    sets, args = [], []
    for k in EDIT_FIELDS:
        if k in body:
            v = body[k]; v = (str(v).strip() or None) if v is not None else None
            if k == "kind" and v and v not in KINDS: raise ValueError("Noto'g'ri tur")
            if k == "scope" and v and v not in SCOPES: raise ValueError("Noto'g'ri kesim")
            if k == "date" and v and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", v): raise ValueError("Sana formati: YYYY-MM-DD")
            sets.append(f"{k}=?"); args.append(v[:4000] if isinstance(v, str) else v)
    if "hidden" in body: sets.append("hidden=?"); args.append(1 if body["hidden"] else 0)
    if "parent_id" in body:
        pid = body["parent_id"]
        if pid in (None, "", 0): sets += ["parent_id=?", "role=?"]; args += [None, "main"]
        else:
            pid = int(pid)
            if pid == lid or not con.execute("SELECT 1 FROM letters WHERE id=?", (pid,)).fetchone(): raise ValueError("Asosiy xat topilmadi")
            sets += ["parent_id=?", "role=?"]; args += [pid, body.get("role") if body.get("role") in ("attachment", "duplicate") else "attachment"]
    if body.get("review"):
        sets += ["confidence=?", "reviewed_by=?", "reviewed_at=?"]; args += ["aniq", user.get("login"), now()]
    for k in ("topics", "facts", "reply_to", "companies_unmatched", "flags"):
        if k in body and isinstance(body[k], list): sets.append(f"{k}=?"); args.append(json.dumps(body[k], ensure_ascii=False))
    with _LOCK:
        if sets:
            sets.append("updated_at=?"); args.append(now())
            con.execute(f"UPDATE letters SET {', '.join(sets)} WHERE id=?", args + [lid])
        if "links" in body:
            con.execute("DELETE FROM letter_links WHERE letter_id=?", (lid,))
            con.executemany("INSERT OR IGNORE INTO letter_links VALUES (?,?,?,?,?)", [(lid, *l) for l in _clean_links(body["links"])])
        _reindex(con, lid)
        con.commit()
    return get_letter(con, lid)

# ------------------------------------------------------------ qo'lda qo'shish: yuklash → matn → avtomatik nomzodlar → saqlash
def _safe_name(name: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", (name or "fayl").strip()).strip(". ") or "fayl"
    return name[:120]

def extract_text(path: Path) -> tuple[str, str]:
    """(matn, izoh). PDF — matn qatlami (pypdf), DOCX — XML, XLSX — openpyxl. Skan PDF'da matn bo'lmaydi (OCR serverda yo'q)."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            try: from pypdf import PdfReader
            except ImportError: return "", "PDF matnini o'qish moduli (pypdf) yo'q"
            rd = PdfReader(str(path)); parts = []
            for i, pg in enumerate(rd.pages[:40]):
                try: parts.append(pg.extract_text() or "")
                except Exception: parts.append("")
            t = "\n".join(parts).strip()
            if len(re.sub(r"\s", "", t)) < 40: return t, "skan (matn qatlami yo'q) — maydonlarni qo'lda to'ldiring"
            return t, f"{len(rd.pages)} sahifa"
        if ext == ".docx":
            with zipfile.ZipFile(path) as z:
                xml = z.read("word/document.xml").decode("utf-8", "ignore")
            xml = re.sub(r"</w:p>", "\n", xml); xml = re.sub(r"<w:tab/>", "\t", xml)
            t = re.sub(r"<[^>]+>", "", xml)
            import html as _h
            return _h.unescape(t).strip(), "Word"
        if ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True); parts = []
            for ws in wb.worksheets[:6]:
                parts.append(f"[varaq: {ws.title}]"); n = 0
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if vals: parts.append(" | ".join(vals)); n += 1
                    if n >= 150: break
            return "\n".join(parts), "Excel"
        if ext in (".txt", ".csv"):
            return path.read_text(encoding="utf-8", errors="ignore"), "matn"
    except Exception as e:
        return "", f"matn o'qilmadi: {e}"
    return "", "bu turdagi fayl matni o'qilmaydi — maydonlarni qo'lda to'ldiring"

_DIST_AL = {
    "1706401": ["buxoro shahri", "buxoro sh", "g buxoro"], "1706403": ["kogon shahri", "kogon sh"], "1706219": ["kogon tumani"],
    "1706207": ["buxoro tumani"], "1706212": ["vobkent"], "1706215": ["gijduvon", "gijduvan"], "1706230": ["qorakol", "karakul"],
    "1706232": ["qorovulbozor", "karaulbazar"], "1706240": ["peshku"], "1706242": ["romitan", "ramitan"], "1706246": ["jondor"],
    "1706258": ["shofirkon", "shafirkan"], "1706204": ["olot", "alat"]}
_COUNTRY_AL = {"Россия": ["rossiya", "russia"], "Қозоғистон": ["qozogiston", "kazaxstan"], "Хитой": ["xitoy", "kitay", "china"],
               "Туркия": ["turkiya", "turtsiya", "turkey"], "Қирғизистон": ["qirgiziston", "kirgiziya"], "Тожикистон": ["tojikiston", "tadjikistan"],
               "Туркманистон": ["turkmaniston", "turkmenistan"], "Афғонистон": ["afgoniston", "afganistan"], "Эрон": ["eron", "iran"],
               "Беларусь": ["belarus"], "Германия": ["germaniya", "germany"], "Польша": ["polsha", "poland"], "Озарбайжон": ["ozarbayjon", "azerbaydjan"],
               "Покистон": ["pokiston", "pakistan"], "Ҳиндистон": ["hindiston", "indiya", "india"], "БАА": ["birlashgan arab", "oae", "uae"],
               "Япония": ["yaponiya", "japan"], "Корея": ["koreya", "korea"], "Малайзия": ["malayziya", "malaysia"], "Украина": ["ukraina"],
               "Латвия": ["latviya"], "Литва": ["litva"], "Грузия": ["gruziya"], "Италия": ["italiya"], "Миср": ["misr", "egipet"]}
_SECTOR_KW = {"Тўқимачилик саноати": ["toqimachilik", "tekstil", "ip kalava", "trikotaj", "tikuv", "gazlama", "paxta tolasi", "denim"],
              "Мева-сабзавот маҳсулотлари": ["meva sabzavot", "meva va sabzavot", "plodoovosh", "karantin"],
              "Қурилиш саноати": ["qurilish material", "sement", "gips", "gipsokarton", "keramika"],
              "Озиқ-овқат саноати": ["oziq ovqat", "yog moy", "konserva", "qandolat"], "Кимё саноати": ["kimyo", "ximiya", "polietilen", "polimer"],
              "Чарм саноати": ["charm", "poyabzal"], "Ипакчилик саноати": ["ipakchilik", "pilla", "kokon"],
              "Электротехника саноати": ["elektrotexnika", "kabel", "transformator"], "Нефт маҳсулотлари": ["neft", "benzin", "dizel"],
              "Мебель": ["mebel"], "Фармацевтика": ["farmatsevtika", "farmasevtika", "dori vosita"]}
_LEGAL = {"mchj", "mas", "xk", "qk", "ooo", "oao", "zao", "ao", "aj", "at", "sp", "ip", "ltd", "llc", "xususiy", "korxona", "korxonasi",
          "firma", "firmasi", "fx", "jshsh", "ok", "masuliyati", "cheklangan", "jamiyati", "jamiyat", "aksiyadorlik", "qoshma", "xorijiy",
          "unitar", "davlat", "oilaviy", "fermer", "xojaligi", "va", "the", "group"}
_GENERIC = {"buxoro", "bukhara", "tekstil", "invest", "agro", "trade", "servis", "service", "savdo", "teks", "gold", "lyuks", "luks", "lux",
            "eko", "eco", "export", "eksport", "nur", "baraka", "star", "plus", "kogon", "vobkent", "romitan", "jondor", "peshku", "olot",
            "qorakol", "gijduvon", "shofirkon", "fayz", "omad", "baxt", "sifat", "orzu", "umid", "yangi", "oltin", "global", "universal",
            "mega", "best", "asia", "osiyo", "sharq", "sanoat", "qurilish", "mebel", "farm", "sut", "non", "avto", "auto", "ipak", "silk",
            "kotton", "paxta", "klaster", "kluster", "agroklaster", "chinor", "turon", "progress", "nigora", "mirzo", "tinchlik"}

def _cores() -> dict[str, set]:
    R = refs()
    if R.get("cores") is not None: return R["cores"]
    cores: dict[str, set] = {}
    for inn, c in R["companies"].items():
        nm = c["name"] or ""
        qm = re.findall(r"[\"“”«»„]([^\"“”«»„]{3,})[\"“”«»„]", nm)
        toks = [t for t in norm(qm[0] if qm else nm).split() if t not in _LEGAL]
        core = " ".join(toks)
        if len(core) < 5 or (len(toks) == 1 and toks[0] in _GENERIC) or (len(toks) < 3 and all(t in _GENERIC for t in toks)): continue
        cores.setdefault(f" {core} ", set()).add(inn)
    R["cores"] = cores
    return cores

_MONTHS = {"yanvar": 1, "fevral": 2, "mart": 3, "aprel": 4, "may": 5, "iyun": 6, "iyul": 7, "avgust": 8, "sentyabr": 9, "sentabr": 9,
           "oktyabr": 10, "oktabr": 10, "noyabr": 11, "dekabr": 12}

def detect(text: str, fname: str = "") -> dict:
    """Avtomatik nomzodlar (import skripti bilan bir xil qoidalar): INN, korxona nomi, tuman, davlat, soha, raqam, sana, javob raqami."""
    R = refs(); nt = norm(text + " " + fname)
    comp = {}
    for m in re.finditer(r"(?<!\d)(\d{3})[  ]?(\d{3})[  ]?(\d{3})(?!\d)", text):
        v = "".join(m.groups())
        if v in R["companies"]: comp[v] = "INN"
    for core, inns in _cores().items():
        if core in nt:
            for inn in inns: comp.setdefault(inn, "nom")
    companies = [{"inn": i, "name": R["companies"][i]["name"], "by": b} for i, b in comp.items()]
    companies.sort(key=lambda x: (x["by"] != "INN", x["name"] or ""))
    head = text[:2500]
    dists = sorted({code for code, al in _DIST_AL.items() for a in al if f" {a}" in norm(head[400:] if len(head) > 600 else head)})
    countries = set()
    for cn in R.get("countries", []):
        k = norm(cn)
        if len(k.strip()) >= 4 and k in nt: countries.add(cn)
    for cn, al in _COUNTRY_AL.items():
        if any(f" {a}" in nt for a in al): countries.add(cn)
    sectors = sorted({s for s, kws in _SECTOR_KW.items() if any(f" {k}" in nt for k in kws)})
    num = None
    m = re.search(r"(?<![\w/-])(\d{1,2}(?:[-/]\d{1,4}){1,4})[-\s]*(?:son|сон)\b", head, re.I)
    if m: num = m.group(1)
    date = None
    m = re.search(r"(\d{1,2})[./](\d{1,2})[./](20\d\d)", head)
    if m: date = f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    else:
        nh = norm(head)
        m = re.search(r" (\d{1,2}) (yanvar|fevral|mart|aprel|may|iyun|iyul|avgust|sentyabr|sentabr|oktyabr|oktabr|noyabr|dekabr)\w* (20\d\d)", nh) \
            or re.search(r" (20\d\d) (?:yil )?(\d{1,2}) (yanvar|fevral|mart|aprel|may|iyun|iyul|avgust|sentyabr|sentabr|oktyabr|oktabr|noyabr|dekabr)", nh)
        if m:
            g = m.groups()
            if g[0].startswith("20"): y, d_, mo = g[0], g[1], g[2]
            else: d_, mo, y = g
            try: date = dt.date(int(y), _MONTHS[mo], int(d_)).isoformat()
            except Exception: date = None
    reply = []
    m = re.search(r"(\d{1,3}[-_]\d{1,6})[-\s]*(?:хатига|xatiga|хатга|xatga|хати|xati)", fname, re.I)
    if m: reply.append({"reg_no": m.group(1).replace("_", "-")})
    m = re.search(r"(\d{4}\s*й(?:ил)?\.?\s*\d{1,2}\s*\w+даги|\d{1,2}\.\d{1,2}\.\d{4}\s*й?\.?\s*(?:даги)?)?\s*(\d{1,2}(?:[-/]\d{1,6}){1,5})[-\s]*(?:сонли|сон|sonli|son)\s*(?:хат|xat)", head, re.I)
    if m and not any(r.get("ext_no") == m.group(2) for r in reply):
        (reply[0] if reply else (reply.append({}) or reply[0])).update({"ext_no": m.group(2)})
    kind = "chiquvchi_xat"
    lf = norm(fname)
    if " ilova" in lf: kind = "ilova"
    elif " murojaat" in lf and " javob" not in lf: kind = "kiruvchi_xat"
    stem = Path(fname).stem
    generic = re.fullmatch(r"(?i)(document|scan|сканировать|img|image|file|fayl|doc|[0-9a-f\-]{16,}|\d+)[\s_()\d-]*", stem or "")
    title = "" if generic else re.sub(r"[_]+", " ", stem).strip()
    return {"companies": companies[:40], "companies_total": len(companies), "districts": [{"code": d, "name": R["districts"].get(d, d)} for d in dists],
            "countries": sorted(countries), "sectors": sectors, "number": num, "date": date, "reply_to": reply, "kind": kind, "title": title}

def upload(raw: bytes, fname: str, user: dict) -> dict:
    if not raw: raise ValueError("Fayl bo'sh")
    today = dt.date.today()
    d = base() / "qolda" / f"{today:%Y}" / f"{today:%m}"; d.mkdir(parents=True, exist_ok=True)
    safe = _safe_name(fname)
    stored = d / f"{dt.datetime.now():%Y%m%d_%H%M%S}_{os.getpid() % 1000}_{safe}"
    stored.write_bytes(raw)
    text, info = extract_text(stored)
    (stored.parent / (stored.name + ".txt")).write_text(text, encoding="utf-8")
    rel = stored.relative_to(base()).as_posix()
    return {"token": rel, "file_name": fname, "size": len(raw), "text_chars": len(text), "text_info": info,
            "text_preview": text[:3000], "suggest": detect(text, fname)}

def create(con, body: dict, user: dict) -> dict:
    tok = str(body.get("token") or "").replace("\\", "/")
    rel = Path(tok)
    if not tok.startswith("qolda/") or rel.is_absolute() or ".." in rel.parts or rel.drive: raise ValueError("Avval faylni yuklang")
    f = base() / rel
    if not f.is_file(): raise ValueError("Avval faylni yuklang")
    if con.execute("SELECT 1 FROM letters WHERE file_path=?", (tok,)).fetchone(): raise ValueError("Bu fayl allaqachon saqlangan")
    title = (body.get("title") or "").strip()
    if not title: raise ValueError("Sarlavha kiriting")
    kind = body.get("kind") or "chiquvchi_xat"
    if kind not in KINDS: raise ValueError("Noto'g'ri tur")
    tp = f.parent / (f.name + ".txt")
    text = tp.read_text(encoding="utf-8", errors="ignore") if tp.exists() else ""
    date = (body.get("date") or "").strip() or None
    if date and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date): raise ValueError("Sana formati: YYYY-MM-DD")
    links = _clean_links(body.get("links"))
    scope = body.get("scope") if body.get("scope") in SCOPES else (
        "korxona" if any(l[0] == "company" for l in links) else "soha" if any(l[0] == "sector" for l in links) else
        "tuman" if any(l[0] == "district" for l in links) else "davlat" if any(l[0] == "country" for l in links) else "umumiy")
    reply = [r for r in (body.get("reply_to") or []) if isinstance(r, dict) and (r.get("reg_no") or r.get("ext_no"))]
    with _LOCK:
        cur = con.execute("""INSERT INTO letters(source, direction, file_path, file_name, file_size, role, parent_id, kind, doc_type, number, date,
                             sender, recipient, title, summary, scope, confidence, note, flags, facts, topics, reply_to, companies_unmatched, hidden,
                             reviewed_by, reviewed_at, created_by, created_at, updated_at)
                             VALUES('manual',?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,0,?,?,?,?,?)""",
                          ("kiruvchi" if kind == "kiruvchi_xat" else "chiquvchi", tok, body.get("file_name") or f.name, f.stat().st_size,
                           "attachment" if body.get("parent_id") else "main", int(body["parent_id"]) if body.get("parent_id") else None,
                           kind, body.get("doc_type") or None, (body.get("number") or "").strip() or None, date, body.get("sender") or None,
                           body.get("recipient") or None, title[:500], (body.get("summary") or "").strip() or None, scope, "aniq",
                           body.get("note") or None, "[]", "[]", json.dumps(body.get("topics") or [], ensure_ascii=False),
                           json.dumps(reply, ensure_ascii=False), "[]", user.get("login"), now(), user.get("login"), now(), now()))
        lid = cur.lastrowid
        rl = list(links)
        for r in reply:
            for key in ("reg_no", "ext_no"):
                if r.get(key): rl.append(("reply", str(r[key])[:60], r.get("from") or None, key))
        con.executemany("INSERT OR IGNORE INTO letter_links VALUES (?,?,?,?,?)", [(lid, *l) for l in rl])
        con.execute("INSERT OR REPLACE INTO letter_text VALUES (?,?)", (lid, text))
        _reindex(con, lid)
        con.commit()
    return get_letter(con, lid)

def lookup(q: str, limit: int = 10) -> dict:
    """Bog'lanish muharriri uchun: korxonalar (nom/INN), sohalar, tumanlar, davlatlar."""
    R = refs(); q = (q or "").strip(); nq = norm(q).strip()
    comps = []
    if q:
        digits = re.sub(r"\D", "", q)
        for inn, c in R["companies"].items():
            if (digits and len(digits) >= 3 and inn.startswith(digits)) or (nq and nq in norm(c["name"] or "")):
                comps.append({"inn": inn, "name": c["name"], "district": R["districts"].get(c["district"] or "", ""), "tarmoq": c["tarmoq"]})
                if len(comps) >= limit: break
    sec = [s for s in R["sectors"] if not nq or nq in norm(s)][: (limit if q else 40)]
    dis = [{"code": k, "name": v} for k, v in R["districts"].items() if k != "1706" and (not nq or nq in norm(v))]
    cns = sorted(set(R.get("countries", [])) | set(_COUNTRY_AL))
    cns = [c for c in cns if not nq or nq in norm(c) or any(nq in a for a in _COUNTRY_AL.get(c, []))][: (limit if q else 60)]
    return {"companies": comps, "sectors": sec, "districts": dis, "countries": cns}

def stats(con) -> dict:
    f = base()
    return {"letters": con.execute("SELECT count(*) FROM letters").fetchone()[0],
            "manual": con.execute("SELECT count(*) FROM letters WHERE source='manual'").fetchone()[0],
            "db_mb": round(db_path().stat().st_size / 1048576, 1) if db_path().exists() else 0}
