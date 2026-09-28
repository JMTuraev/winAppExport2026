"""Header-based readers for .xls / .xlsx (columns are found by name, never by position)."""
from __future__ import annotations
import re, io, datetime as dt
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
from pathlib import Path

def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()

def read_sheet(path: str | Path, sheet: str | int = 0, max_rows: int | None = None, data: bytes | None = None) -> list[list]:
    """Return all rows of a sheet as lists of python values (str/float/None).
    data: the file's bytes — then the file is not opened again (Windows keeps no lock on it)."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data, on_demand=True) if data is not None else xlrd.open_workbook(str(path), on_demand=True)
        sh = wb.sheet_by_index(sheet) if isinstance(sheet, int) else wb.sheet_by_name(sheet)
        out = []
        for r in range(sh.nrows if max_rows is None else min(sh.nrows, max_rows)):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = dt.datetime(*xlrd.xldate_as_tuple(v, wb.datemode)).date().isoformat()
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR) or v == "":
                    v = None
                elif isinstance(v, str):
                    v = _CTRL.sub("", v).strip()
                row.append(v)
            out.append(row)
        return out
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data) if data is not None else str(path), read_only=True, data_only=True, keep_links=False)
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        out = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if max_rows is not None and i >= max_rows: break
            row = [re.sub(r"_x00[01][0-9a-fA-F]_", "", v) if isinstance(v, str) else v for v in row]
            out.append([(v.date().isoformat() if isinstance(v, dt.datetime) else (None if isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")) else (v.replace("\x1a","").strip() if isinstance(v,str) else v))) for v in row])
        wb.close()
        return out

def sheet_names(path, data: bytes | None = None) -> list[str]:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data, on_demand=True) if data is not None else xlrd.open_workbook(str(path), on_demand=True)
        n = wb.sheet_names()
        try: wb.release_resources()
        except Exception: pass
        return n
    import openpyxl; wb = openpyxl.load_workbook(io.BytesIO(data) if data is not None else str(path), read_only=True, keep_links=False); n = wb.sheetnames; wb.close(); return n

class Header:
    """Locate columns by (fuzzy) header text."""
    def __init__(self, rows: list[list], required: list[str], scan: int = 15):
        self.row_idx = None
        for i, row in enumerate(rows[:scan]):
            cells = [_norm(v) for v in row]
            if all(any(_norm(req) in c for c in cells) for req in required):
                self.row_idx = i; self.cells = cells; break
        if self.row_idx is None:
            raise ValueError("Sarlavha qatori topilmadi. Kerakli ustunlar: " + ", ".join(required))
    def col(self, *names: str, required=True) -> int | None:
        for name in names:
            n = _norm(name)
            for i, c in enumerate(self.cells):
                if c == n: return i
            for i, c in enumerate(self.cells):
                if n and n in c: return i
        if required:
            raise ValueError(f"Ustun topilmadi: {names[0]}")
        return None

def inn_str(v) -> str:
    if v is None: return ""
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return str(v).strip().split(".")[0] if str(v).strip().replace(".","").isdigit() else str(v).strip()

def num(v) -> float:
    if v is None or v == "": return 0.0
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(" ", "").replace(",", "."))
    except ValueError: return 0.0


def sheet_key(name: str) -> str:
    """Sheet identity by letters only: digits, commas and dots are ignored («СВОД (карантин) 288,7» -> «свод (карантин)»)."""
    return re.sub(r"\s+", " ", re.sub(r"[\d.,]+", " ", str(name or ""))).strip().lower()

def find_sheet(names, target: str):
    """Pick the sheet whose letters match target. Several matches: prefer one whose name carries a number (the plan figure), the last such."""
    key = sheet_key(target)
    cands = [n for n in names if sheet_key(n) == key]
    if not cands: return None
    with_num = [n for n in cands if re.search(r"\d", n)]
    return (with_num or cands)[-1]
