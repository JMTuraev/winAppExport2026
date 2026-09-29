"""Read-side computations shared by the UI and the Excel export.

Rules
- sanoat  = номманом opening (kind 'sanoat') + saved GTD rows (kind 'sanoat') after the opening date
- meva    = karantin opening (СВОД (карантин)) + saved GTD rows (kind 'meva', grown region 1706) after the opening date
- memo    = companies kept outside the main total (БНПЗ, SDK) — shown separately
"""
from __future__ import annotations
import datetime as dt, json
from collections import defaultdict
from . import db

MONTH_NAMES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
REGION = "1706"

def d2(s: str) -> dt.date: return dt.date.fromisoformat(s)

def opening_date(con) -> str: return db.get_setting(con, "opening_through_date")

def districts(con):
    return [dict(r) for r in con.execute("SELECT code,name_uz,order_no FROM districts WHERE length(code)=7 ORDER BY order_no, name_uz")]

def prev_year_days(con, D: str) -> int:
    off = int(db.get_setting(con, "prev_year_day_offset", "0"))
    return d2(D).timetuple().tm_yday + off

def prev_year_amount(con, code: str, kind: str, D: str, year: int | None = None, fset=None) -> float:
    """«Ўтган йил шу даврга» — ФАҚАТ Sozlamalardaги қўлда киритилган ўтган йил базасидан (prognoz.prev_year_base,
    prev_year_base_days), D санасигача кунларга мутаносиб. Божхона базасидаги 2025 items ҳеч қаерда ишлатилмайди
    (Жафар қарори 19.09.2026: вазирлик/ҳокимият йилни бошқа кўрсаткич билан якунлаган бўлиши мумкин).
    code — '1706' (вилоят) ёки тuman коди; kind — 'all' | 'sanoat' | 'meva'.
    fset — танланган «Амалда» тўплами (29.09.2026, Dashboard/Свод танлови): фақат шу тўплам даврлари; бўлмаса — асосий + база."""
    y = year or int(D[:4])
    r = con.execute("SELECT prev_year_base, prev_year_base_days FROM prognoz WHERE year=? AND district_code=? AND kind=?", (y, code, kind)).fetchone()
    base = (r["prev_year_base"], r["prev_year_base_days"]) if r and r["prev_year_base"] else None
    # 21.09.2026: аввало Sozlamalar → «Даврий натижалар» (N ойлик расмий якунлар) — ой охирида аниқ рақам,
    # оралиқ кунда даврлар (ва эски база якори) орасида интерполяция; давр йўқ бўлса — эски база, кунга мутаносиб
    from . import periods
    pc = periods.prev_cum(con, code, kind, D, base, fset)
    if pc is not None: return pc[0]
    if fset or not base: return 0.0
    return base[0] / (base[1] or 273) * prev_year_days(con, D)

def last_data_date(con) -> str | None:
    r = con.execute("SELECT max(report_date) FROM gtd_rows WHERE voided=0").fetchone()[0]
    vals = [v for v in (r, opening_date(con)) if v]
    return max(vals) if vals else None

def first_data_date(con) -> str | None:
    r = con.execute("SELECT min(report_date) FROM gtd_rows WHERE voided=0").fetchone()[0]
    return opening_date(con) or r

def kind_skip(kind): return kind == "exclude"

def not_bukhara_inns(con) -> set:
    """Companies marked «Бухоро эмас — ҳисобга киритилмасин»: nothing of theirs is counted — neither the monthly
    customs base (opening_balances / items) nor the daily GTD rows. The flag is applied at read time, so
    switching it in «Korxonalar» takes effect everywhere (dashboard, top-10, districts, map, Excel) at once."""
    return {r["inn"] for r in con.execute("SELECT inn FROM companies WHERE bukhara=0")}

def bukhara_windows(con) -> dict:
    """{inn: (from, to)} — companies whose Bukhara period is limited (moved into or out of the region).
    A row counts only when from <= date <= to; an open side is '' / '9999'. Ignored for bukhara=0 (nothing counts)."""
    out = {}
    for r in con.execute("SELECT inn, bukhara_from, bukhara_to FROM companies WHERE bukhara IS NOT 0 AND (bukhara_from IS NOT NULL OR bukhara_to IS NOT NULL)"):
        out[r["inn"]] = (r["bukhara_from"] or "", r["bukhara_to"] or "9999")
    return out

def in_window(win: dict, inn: str, d: str) -> bool:
    w = win.get(inn)
    return True if not w else (w[0] <= (d or "") <= w[1])

def window_base_months(con, inn: str, w, year: int) -> list:
    """Monthly sanoat amounts of one company from the customs base rows inside its Bukhara period (meva rows excluded,
    like opening_balances). Falls back to whole months of opening_balances when the base has no row detail."""
    m = [0.0] * 12
    rows = con.execute("SELECT rdate, kind, value FROM items WHERE regime='ЭК' AND inn=? AND year=?", (inn, int(year))).fetchall()
    if rows:
        for r in rows:
            if w[0] <= r["rdate"] <= w[1] and r["kind"] == "sanoat": m[int(r["rdate"][5:7]) - 1] += r["value"] or 0
        return m
    for r in con.execute("SELECT month, amount FROM opening_balances WHERE inn=? AND year=? AND kind='sanoat'", (inn, year)):
        first, last = f"{year}-{r['month']:02d}-01", f"{year}-{r['month']:02d}-31"
        if w[0] <= last and first <= w[1]: m[r["month"] - 1] += r["amount"]
    return m

def company_amounts(con, D: str):
    """{(inn, kind): {"m": [12 floats], "day": float}} for kinds sanoat/meva(номманом legacy)/memo, as of date D (inclusive)."""
    through = opening_date(con) or ""; year = int(D[:4])
    skip = not_bukhara_inns(con); win = bukhara_windows(con)
    out = defaultdict(lambda: {"m": [0.0] * 12, "day": 0.0, "ytd_memo": 0.0})
    for r in con.execute("SELECT inn,kind,month,amount FROM opening_balances WHERE year=? AND kind NOT IN ('bnpz','exclude')", (year,)):
        if r["inn"] in skip or (r["inn"] in win and r["kind"] == "sanoat"): continue
        if r["kind"] == "memo": out[(r["inn"], "memo")]["ytd_memo"] += r["amount"]
        else: out[(r["inn"], r["kind"])]["m"][r["month"] - 1] += r["amount"]
    for inn, w in win.items():          # limited Bukhara period: the base is re-summed by date inside the period
        if inn in skip: continue
        mm = window_base_months(con, inn, w, year)
        if any(mm):
            for i, v in enumerate(mm): out[(inn, "sanoat")]["m"][i] += v
    if D == through:
        for inn, kind, amt in json.loads(db.get_setting(con, "opening_last_day", "[]")):
            if inn in skip or not in_window(win, inn, through): continue
            out[(inn, kind)]["day"] += amt
    for r in con.execute("""SELECT inn, kind, substr(report_date,6,2) m, report_date d, sum(stat_usd) s FROM gtd_rows
                            WHERE voided=0 AND kind IN ('sanoat','memo') AND report_date>? AND report_date<=? AND substr(report_date,1,4)=?
                            GROUP BY inn, kind, m, d""", (through, D, str(year))):
        if r["inn"] in skip or not in_window(win, r["inn"], r["d"]): continue
        k = (r["inn"], r["kind"])
        if r["kind"] == "memo": out[k]["ytd_memo"] += r["s"]
        else: out[k]["m"][int(r["m"]) - 1] += r["s"]
        if r["d"] == D: out[k]["day"] += r["s"]
    # meva-sabzavot of a Bukhara exporter goes to its номманом «Мева-сабзавотлар» row in Excel — count it here the same way
    from .daily import is_bukhara
    meva_inns = {i for (i,) in con.execute("SELECT inn FROM template_inns WHERE category='meva'")} or \
                {i for (i,) in con.execute("SELECT inn FROM companies WHERE kind='meva' AND source='nomma-nom'")}
    comps = {r["inn"]: dict(r) for r in con.execute("SELECT inn, in_base, bukhara, source FROM companies")}
    for r in con.execute("""SELECT inn, kind, substr(report_date,6,2) m, report_date d, sum(stat_usd) s FROM gtd_rows
                            WHERE voided=0 AND kind IN ('meva','meva_x') AND report_date>? AND report_date<=? AND substr(report_date,1,4)=?
                            GROUP BY inn, kind, m, d""", (through, D, str(year))):
        if r["inn"] in skip or not in_window(win, r["inn"], r["d"]): continue
        if r["kind"] == "meva" and not is_bukhara(comps.get(r["inn"]), r["inn"], meva_inns): continue
        k = (r["inn"], "meva")
        out[k]["m"][int(r["m"]) - 1] += r["s"]
        if r["d"] == D: out[k]["day"] += r["s"]
    return out

def bnpz(con, D: str):
    """БНПЗ (only for «СВОД 450 млн», Қоровулбозор): ytd / current month / day as of D."""
    through = opening_date(con) or ""; year = int(D[:4]); month = int(D[5:7])
    out = {"ytd": 0.0, "month": 0.0, "day": 0.0}
    for r in con.execute("SELECT month, sum(amount) a FROM opening_balances WHERE kind='bnpz' AND year=? GROUP BY month", (year,)):
        out["ytd"] += r["a"]
        if r["month"] == month: out["month"] += r["a"]
    for r in con.execute("""SELECT report_date d, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind='bnpz' AND report_date>? AND report_date<=?
                            AND substr(report_date,1,4)=? GROUP BY d""", (through, D, str(year))):
        out["ytd"] += r["s"]
        if int(r["d"][5:7]) == month: out["month"] += r["s"]
        if r["d"] == D: out["day"] += r["s"]
    return out

def karantin(con, D: str):
    """{district_code: {"ytd", "month", "day"}} meva-sabzavot by certificate region rule."""
    through = opening_date(con) or ""; year = int(D[:4]); month = int(D[5:7])
    res = defaultdict(lambda: {"ytd": 0.0, "month": 0.0, "day": 0.0})
    for r in con.execute("SELECT district_code, ytd, month, month_amt FROM karantin_opening WHERE year=?", (year,)):
        res[r["district_code"]]["ytd"] += r["ytd"]
        if r["month"] == month: res[r["district_code"]]["month"] += r["month_amt"]
    if D == through:
        for dc, amt in json.loads(db.get_setting(con, "karantin_last_day", "{}")).items(): res[dc]["day"] += amt
    for r in con.execute("""SELECT district_code dc, report_date d, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind='meva'
                            AND report_date>? AND report_date<=? AND substr(report_date,1,4)=? GROUP BY dc, d""", (through, D, str(year))):
        dc = r["dc"] or "?"
        res[dc]["ytd"] += r["s"]
        if int(r["d"][5:7]) == month: res[dc]["month"] += r["s"]
        if r["d"] == D: res[dc]["day"] += r["s"]
    return res

def karantin_detail(con, D: str):
    """Карантин рўйхати ичи — туман кесимида: {district_code: {"opening": {...}, "companies": [...]}}.
    opening  — карантин свод қолдиғи (opening_through_date ҳолатига, туман жами, корхонасиз);
    companies — қолдиқ санасидан кейинги ГТД (kind='meva', туман — етиштирилган жой) экспортёр бўйича, ойма-ой.
    Қолдиқ + корхоналар = karantin() қатори (номма-номда «Мева-сабзавот (карантин рўйхати)»)."""
    through = opening_date(con) or ""; year = int(D[:4]); month = int(D[5:7])
    cmap = companies_map(con); nb = not_bukhara_inns(con)
    res = defaultdict(lambda: {"opening": None, "companies": []})
    for r in con.execute("SELECT district_code, ytd, month, month_amt FROM karantin_opening WHERE year=? AND ytd<>0", (year,)):
        res[r["district_code"]]["opening"] = {"ytd": round(r["ytd"], 3), "month": r["month"], "month_amt": round(r["month_amt"] or 0, 3),
                                              "day": 0.0, "through": through}
    if D == through:
        for dc, amt in json.loads(db.get_setting(con, "karantin_last_day", "{}")).items():
            if res[dc]["opening"]: res[dc]["opening"]["day"] += amt
    agg = {}
    for r in con.execute("""SELECT district_code dc, inn, max(exporter) exporter, substr(report_date,6,2) m, sum(stat_usd) s,
                                   sum(CASE WHEN report_date=? THEN stat_usd ELSE 0 END) dsum
                            FROM gtd_rows WHERE voided=0 AND kind='meva' AND report_date>? AND report_date<=? AND substr(report_date,1,4)=?
                            GROUP BY dc, inn, m""", (D, through, D, str(year))):
        dc = r["dc"] or "?"
        a = agg.setdefault((dc, r["inn"]), {"inn": r["inn"], "name": (cmap.get(r["inn"], {}).get("name") or r["exporter"] or r["inn"]).strip(),
                                            "m": [0.0] * 12, "day": 0.0, "bukhara": r["inn"] in cmap and r["inn"] not in nb})
        a["m"][int(r["m"]) - 1] += r["s"]; a["day"] += r["dsum"]
    for (dc, _), a in agg.items():
        a["m"] = [round(x, 3) for x in a["m"][:month]]; a["day"] = round(a["day"], 3); a["total"] = round(sum(a["m"]), 3)
        res[dc]["companies"].append(a)
    for v in res.values(): v["companies"].sort(key=lambda a: -a["total"])
    return res

def karantin_monthly(con, year: int, D: str | None = None):
    """Meva by month where known: opening month (from karantin svod) + GTD months. Earlier months are unknown (None)."""
    through = opening_date(con) or ""; om = int(through[5:7]) if through and int(through[:4]) == year else 0
    vals = [None] * 12
    if om:
        tot = con.execute("SELECT sum(month_amt) FROM karantin_opening WHERE year=?", (year,)).fetchone()[0] or 0
        if tot: vals[om - 1] = tot
    for r in con.execute("""SELECT substr(report_date,6,2) m, district_code dc, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind='meva'
                            AND report_date>? AND report_date<=? AND substr(report_date,1,4)=? GROUP BY m, dc""", (through, D or f"{year}-12-31", str(year))):
        i = int(r["m"]) - 1; vals[i] = (vals[i] or 0) + r["s"]
    return vals

def companies_map(con):
    return {r["inn"]: dict(r) for r in con.execute("SELECT * FROM companies")}

def prognoz(con, year: int, D: str | None = None, plan=None):
    """Reja: Sozlamalardagi versiyali reja (D sanasida amaldagi) bo'lsa — undan; bo'lmasa o'z jadvalingizning СВОД prognozi.
    plan: None — sana bo'yicha (avto); "table" — faqat jadval (shablon СВОДи) prognozi; int — tanlangan versiya (kunlik hisobot modali).
    «O'tgan yil bazasi» (prev_year_base) har doim jadvaldan."""
    p = defaultdict(dict)
    for r in con.execute("SELECT * FROM prognoz WHERE year=?", (year,)):
        p[r["district_code"]][r["kind"]] = dict(r)
    if plan == "table": return p
    from . import plans
    pv = plans.plan_for(con, D or f"{year}-12-31", plan)
    if pv:
        for code, kinds in pv.items():
            for kind, v in kinds.items():
                old = p[code].get(kind) or {}
                p[code][kind] = {**old, **v, "district_code": code, "year": year, "kind": kind, "plan_source": "versions"}
    return p

def fset_param(v):
    """URL qiymati → «Амалда» тўплами id (ёки None = асосий)."""
    v = str(v or "").strip()
    return int(v) if v.isdigit() and int(v) > 0 else None

def plan_param(v):
    """URL/so'rov qiymati → prognoz(plan=...): '', 'auto' → None; 'table' → 'table'; raqam → versiya id."""
    v = str(v or "").strip()
    if v in ("", "auto", "None", "null"): return None
    if v == "table": return "table"
    return int(v) if v.isdigit() else None

def dashboard(con, D: str | None = None, plan=None, fset=None):
    D = D or last_data_date(con)
    if not D:
        return {"empty": True}
    if fset:   # «Амалда» тўплами (29.09.2026): йўқ (ўчирилган) бўлса — асосий
        from . import periods
        if not periods.get_set(con, fset): fset = None
    if isinstance(plan, int):   # танланган режа бу санага мос эмас (ўтган давр блоки) — экранда стандартга қайтади; Excel'да require_cover тўхтатади
        from . import plans as _pl
        _v = _pl.get_version(con, plan)
        if not _pl.covers(_v, D): plan = None
    year = int(D[:4]); month = int(D[5:7])
    amounts = company_amounts(con, D); cmap = companies_map(con); kar = karantin(con, D)
    days = prev_year_days(con, D)
    comps = []
    for (inn, kind), a in amounts.items():
        if kind == "memo": continue
        c = cmap.get(inn, {})
        if kind == "meva":  # номманом legacy meva rows are not part of the karantin rule; they are listed but flagged
            pass
        comps.append({"inn": inn, "name": c.get("name") or inn, "d": c.get("district_code"), "t": kind, "tarmoq": (c.get("tarmoq") or "").strip(),
                      "m": [round(x, 3) for x in a["m"][:month]], "day": round(a["day"], 3), "total": round(sum(a["m"]), 3)})
    pr = prognoz(con, year, D, plan)
    dist = []
    for d in districts(con):
        code = d["code"]; p = pr.get(code, {})
        dist.append({"code": code, "name": d["name_uz"],
                     "plan": {k: {"year": v.get("year_plan"), "m9": v.get("m9_plan"), "period": v.get("period_plan"), "month": v.get("cur_month_plan"),
                                  "prev": prev_year_amount(con, code, k, D, fset=fset), "plan_source": v.get("plan_source"),
                                  "months": json.loads(v["months"]) if v.get("months") else None} for k, v in p.items()},
                     "meva": {k: round(v, 3) for k, v in kar.get(code, {"ytd": 0, "month": 0, "day": 0}).items()}})
    region = {k: {"year": v.get("year_plan"), "m9": v.get("m9_plan"), "period": v.get("period_plan"), "month": v.get("cur_month_plan"),
                  "prev": prev_year_amount(con, REGION, k, D, fset=fset), "plan_source": v.get("plan_source")} for k, v in pr.get(REGION, {}).items()}
    from . import periods
    prev_info = periods.prev_info(con, D, fset)
    memo = []
    for (inn, kind), a in amounts.items():
        if kind == "memo": memo.append({"inn": inn, "name": cmap.get(inn, {}).get("name"), "ytd": round(a["ytd_memo"], 3), "day": round(a["day"], 3)})
    return {"date": D, "year": year, "month": month, "opening_date": opening_date(con), "prev_year_days": days,
            "prev_year_label": db.get_setting(con, "prev_year_base_label"), "prev_year_source": prev_info,
            "companies": comps, "districts": dist, "region_plan": region, "meva_monthly": karantin_monthly(con, year, D), "memo": memo,
            "karantin_detail": karantin_detail(con, D), "plan_info": _plan_info(con, pr),
            "selection": {"plan": plan if plan is not None else "auto", "fset": fset}}

def _plan_info(con, pr) -> dict:
    """Qaysi reja ishlatildi (versiya yoki jadval prognozi) — sarlavha/izoh uchun."""
    r = ((pr or {}).get(REGION) or {}).get("all") or {}
    vid = r.get("version_id")
    if r.get("plan_source") == "versions" and vid:
        v = con.execute("SELECT id, title, effective_from FROM plan_versions WHERE id=?", (vid,)).fetchone()
        if v: return {"source": "version", "id": v["id"], "title": v["title"], "effective_from": v["effective_from"]}
    return {"source": "table", "title": "Жадвал (шаблон СВОДи) прогнози", "label": db.get_setting(con, "template_label")}

def calendar(con, year: int, month: int):
    through = opening_date(con)
    rows = {r["d"]: dict(r) for r in con.execute("""SELECT report_date d, sum(CASE WHEN kind='sanoat' THEN stat_usd ELSE 0 END) san,
            sum(CASE WHEN kind='meva' THEN stat_usd ELSE 0 END) mev, count(DISTINCT gtd_no) gtd, count(*) n
            FROM gtd_rows WHERE voided=0 AND substr(report_date,1,7)=? GROUP BY d""", (f"{year}-{month:02d}",))}
    uploads = defaultdict(list)
    for r in con.execute("SELECT id, filename, report_dates, loaded_at, status FROM uploads WHERE status='saved'"):
        for d in json.loads(r["report_dates"]): uploads[d].append({"id": r["id"], "file": r["filename"], "at": r["loaded_at"]})
    import calendar as cal
    out = []
    for day in range(1, cal.monthrange(year, month)[1] + 1):
        d = f"{year}-{month:02d}-{day:02d}"
        state = "opening" if through and d <= through else ("loaded" if d in rows else "empty")
        out.append({"date": d, "state": state, **({k: rows[d][k] for k in ("san", "mev", "gtd", "n")} if d in rows else {}), "uploads": uploads.get(d, [])})
    return {"year": year, "month": month, "opening_date": through, "first_date": first_data_date(con), "days": out}

def day_detail(con, D: str):
    through = opening_date(con)
    amounts = company_amounts(con, D); cmap = companies_map(con); kar = karantin(con, D)
    san_ytd = sum(sum(a["m"]) for (i, k), a in amounts.items() if k == "sanoat")
    mev_ytd = sum(v["ytd"] for v in kar.values())
    san_day = sum(a["day"] for (i, k), a in amounts.items() if k == "sanoat")
    mev_day = sum(v["day"] for v in kar.values())
    month = int(D[5:7])
    comp_rows = []
    for r in con.execute("""SELECT inn, count(DISTINCT gtd_no) gtd, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind='sanoat' AND report_date=? GROUP BY inn ORDER BY s DESC""", (D,)):
        c = cmap.get(r["inn"], {}); a = amounts.get((r["inn"], "sanoat"), {"m": [0] * 12})
        comp_rows.append({"inn": r["inn"], "name": c.get("name"), "district_code": c.get("district_code"), "gtd": r["gtd"], "day": round(r["s"], 3),
                          "month": round(a["m"][month - 1], 3), "ytd": round(sum(a["m"]), 3), "new": c.get("new_date") == D})
    meva_rows = [dict(r) for r in con.execute("""SELECT report_date, inn, exporter, district_code, grown_district, product, country, netto, stat_usd FROM gtd_rows
                                                 WHERE voided=0 AND kind='meva' AND report_date=? ORDER BY stat_usd DESC""", (D,))]
    memo_rows = [dict(r) for r in con.execute("SELECT inn, exporter, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind='memo' AND report_date=? GROUP BY inn", (D,))]
    ups = [dict(r) for r in con.execute("SELECT * FROM uploads WHERE report_dates LIKE ? ORDER BY id DESC", (f'%"{D}"%',))]
    svod_meva = con.execute("""SELECT coalesce(sum(CASE WHEN report_date=? THEN stat_usd END),0), coalesce(sum(stat_usd),0) FROM gtd_rows
                               WHERE voided=0 AND kind='meva_x' AND report_date>? AND report_date<=?""", (D, through or "", D)).fetchone()
    return {"date": D, "opening_date": through, "is_opening": bool(through and D < through), "exact": (not through) or D >= through,
            "svod_meva": {"day": round(svod_meva[0], 3), "ytd": round(svod_meva[1], 3)},
            "day": {"sanoat": round(san_day, 3), "meva": round(mev_day, 3), "total": round(san_day + mev_day, 3)},
            "ytd": {"sanoat": round(san_ytd, 3), "meva": round(mev_ytd, 3), "total": round(san_ytd + mev_ytd, 3)},
            "companies": comp_rows, "meva_rows": meva_rows, "memo_rows": memo_rows, "uploads": ups}


MSHORT = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]

def export_nomma(con, q: dict, out):
    """«Номма-ном» жадвали — саҳифадаги филтрлар билан (туман, тармоқ, қидирув) Excel'га. Расмий маълумотнома кўринишида."""
    import openpyxl
    d = dashboard(con, q.get("date"))
    if d.get("empty"): raise ValueError("База бўш")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Номма-ном"
    _nomma_sheet(ws, d, q.get("district"), q.get("t"), q.get("q"))
    wb.save(out); return out

def _nomma_sheet(ws, d: dict, district: str | None, t: str | None, q: str | None):
    """Битта варақга «номма-ном» жадвалини ёзади (export_nomma ва свод+туманлар китоби учун умумий)."""
    import re
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    dc = (district or "").strip(); t = t or "sanoat"; t = "" if t == "all" else t; qs = (q or "").strip().lower()
    words = [w for w in re.split(r"\s+", qs) if w]
    def match(c): hay = f"{c['inn']} {c['name']}".lower(); return all(w in hay for w in words)
    cm = d["month"]; pk = t or "all"
    D = d["date"]; nxt = (dt.date.fromisoformat(D) + dt.timedelta(days=1)).strftime("%d.%m.%Y")
    thin = Side(style="thin", color="9AA5A0"); bd = Border(thin, thin, thin, thin)
    H = Font(bold=True, size=10); HF = PatternFill("solid", fgColor="D9E7E2")
    TF = PatternFill("solid", fgColor="EEF4F1"); DF = PatternFill("solid", fgColor="0F6E63"); B = Font(bold=True)
    KF = PatternFill("solid", fgColor="F3F8EC")   # карантин рўйхати қаторлари
    C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    nd = dt.date.fromisoformat(D)   # маълумот куни (сарлавҳадаги санадан бир кун олдин)
    day_lbl = f"{nd.day} {['янв.', 'фев.', 'март', 'апр.', 'май', 'июнь', 'июль', 'авг.', 'сент.', 'окт.', 'нояб.', 'дек.'][nd.month - 1]}"   # «1 кунда» ўрнига ҳисобот санаси
    hdr = ["№", "ИНН", "Корхона номи"] + MSHORT[:cm] + [day_lbl, "Жами", "Улуши, %", "Тармоқ"]
    widths = [5, 12, 46] + [9.5] * cm + [9.5, 11, 9, 24]
    NC = len(hdr); last = L(NC); JC = 3 + cm + 2; PC = JC + 1   # Жами ва Улуши устунлари
    blocks = []
    for dd in d["districts"]:
        if dc and dd["code"] != dc: continue
        cs = [c for c in d["companies"] if c["d"] == dd["code"] and (not t or c["t"] == t) and match(c)]
        kar_only = t in ("", "meva") and d.get("meva_mode", "karantin") == "karantin" and (dd.get("meva") or {}).get("ytd") and not words
        if cs or kar_only: blocks.append((dd, sorted(cs, key=lambda c: -c["total"])))
    tname = {"sanoat": "саноат маҳсулотлари", "meva": "мева-сабзавот", "": ""}.get(t, t)
    where = blocks[0][0]["name"] + "даги" if len(blocks) == 1 and dc else "Бухоро вилоятидаги"
    ws.merge_cells(f"A1:{last}1"); ws["A1"] = f"{where} экспортёр корхоналар тўғрисида маълумот" + (f" ({tname})" if tname and t else "")
    ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = C; ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last}2"); ws["A2"] = f"{nxt} йил"; ws["A2"].font = Font(bold=True, size=11); ws["A2"].alignment = C
    ws.merge_cells(f"{L(NC - 3)}3:{last}3"); ws.cell(3, NC - 3, "минг АҚШ доллари").font = Font(italic=True, size=10); ws.cell(3, NC - 3).alignment = Alignment(horizontal="right")
    if qs: ws["A3"] = f"Қидирув: «{qs}»"; ws["A3"].font = Font(italic=True, color="666666")
    HR = 4
    for i, (h, w) in enumerate(zip(hdr, widths), 1):
        c = ws.cell(HR, i, h); c.font = H; c.fill = HF; c.alignment = C; c.border = bd; ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[HR].height = 28
    def put(r, vals, bold=False, fill=None, italic=False):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = bd
            if bold or italic: c.font = Font(bold=bold, italic=italic)
            if fill: c.fill = fill
            if 4 <= i <= JC: c.number_format = "#,##0.00" if i == JC - 1 else "#,##0.0"
            if i == PC: c.number_format = "0.0%"; c.alignment = Alignment(horizontal="center")
            if i in (1, 2): c.alignment = Alignment(horizontal="center")
    row = HR + 1; n = 0
    for dd, cs in blocks:
        if len(blocks) > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NC)
            c = ws.cell(row, 1, dd["name"]); c.font = Font(bold=True, color="FFFFFF", size=11); c.fill = DF; c.alignment = C
            for i in range(1, NC + 1): ws.cell(row, i).border = bd
            ws.row_dimensions[row].height = 22; row += 1
        sums = [0.0] * cm; sd = st = 0.0
        for c in cs:
            for i, v in enumerate(c["m"]): sums[i] += v or 0
            sd += c["day"]; st += c["total"]
        kar = (dd.get("meva") or {}) if d.get("meva_mode", "karantin") == "karantin" else {}
        den = st + (kar.get("ytd") or 0 if t in ("", "meva") else 0)
        for k, c in enumerate(cs, 1):
            n += 1
            put(row, [n, c["inn"], c["name"] + (" (мева)" if c["t"] == "meva" else "")] + [v or None for v in c["m"]] + [c["day"] or None, c["total"], (c["total"] / den) if den else None, c["tarmoq"]]); row += 1
        if t in ("", "meva") and kar.get("ytd"):   # карантин рўйхати бўйича мева-сабзавот — Сводга мос
            det = (d.get("karantin_detail") or {}).get(dd["code"]) or {}
            kc = det.get("companies") or []; op = det.get("opening")
            mv = [0.0] * cm
            if kc or op:   # рўйхат: ГТД экспортёрлари (етиштирилган жой — шу туман) + қолдиқ қатори
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NC)
                c = ws.cell(row, 1, "Мева-сабзавот — карантин рўйхати (Бухорода етиштирилган, экспортёридан қатъи назар)")
                c.font = Font(bold=True, italic=True, color="0F6E63"); c.fill = KF; c.alignment = Alignment(horizontal="left", vertical="center")
                for i in range(1, NC + 1): ws.cell(row, i).border = bd
                row += 1
                for k, a in enumerate(kc, 1):
                    nm = a["name"] + ("" if a.get("bukhara") else " (бошқа вилоят экспортёри)")
                    put(row, [k, a["inn"], nm] + [v or None for v in a["m"]] + [a["day"] or None, a["total"], (a["total"] / den) if den else None, "Мева-сабзавот"], italic=True, fill=KF); row += 1
                    for i, v in enumerate(a["m"]): mv[i] += v or 0
                if op and (op["ytd"] or op["day"]):
                    om = [None] * cm
                    if op.get("month") and 1 <= op["month"] <= cm and op["month_amt"]: om[op["month"] - 1] = op["month_amt"]; mv[op["month"] - 1] += op["month_amt"]
                    thr = dt.date.fromisoformat(op["through"]).strftime("%d.%m.%Y") if op.get("through") else ""
                    put(row, ["", "", f"Қолдиқ — {thr} ҳолатига (карантин свод, корхоналар кесимисиз)"] + om + [op["day"] or None, op["ytd"], (op["ytd"] / den) if den else None, "Мева-сабзавот"], italic=True, fill=KF); row += 1
            else:
                mv[cm - 1] = kar.get("month") or 0
            put(row, ["", "", "Мева-сабзавот жами (карантин рўйхати)"] + [v or None for v in mv] + [kar.get("day") or None, kar["ytd"], kar["ytd"] / den if den else None, "Мева-сабзавот"], bold=True, italic=True, fill=KF); row += 1
            for i, v in enumerate(mv): sums[i] += v or 0
            sd += kar.get("day") or 0; st += kar["ytd"]
        put(row, ["", "", f"{dd['name']} жами"] + sums + [sd, st, 1 if st else None, ""], bold=True, fill=TF); row += 1
        p = (dd.get("plan") or {}).get(pk) or {}
        yp, pp = p.get("year"), p.get("period")
        put(row, ["", "", f"Йиллик режа · бажарилиши, %"] + [None] * (cm + 1) + [yp, (st / yp) if yp else None, ""], italic=True, fill=TF); row += 1
        put(row, ["", "", f"Янв–{MSHORT[cm - 1]} режа · бажарилиши, %"] + [None] * (cm + 1) + [pp, (st / pp) if pp else None, ""], italic=True, fill=TF); row += 1
        row += 1  # тuманлар орасида бўш қатор
    ws.freeze_panes = f"D{HR + 1}"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HR}:{HR}"

def _svod_sheet_plain(ws, con, d: dict):
    """(захира — шаблон бўлмаганда) «СВОД (карантин)» — туманлар кесимида прогноз/амалда (ўз жадвалингиздаги СВОД (карантин) варағи услубида, формуласиз)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    D = d["date"]; year = d["year"]; cm = d["month"]
    nxt = (dt.date.fromisoformat(D) + dt.timedelta(days=1)).strftime("%d.%m.%Y")
    mon = MONTH_NAMES[cm - 1].lower()
    prev_lbl = (d.get("prev_col_label") or d.get("prev_year_label") or f"{year - 1} йил шу кун ҳолатида амалдаги экспорт").strip()
    cmp_lbl = (d.get("prev_cmp_label") or f"{year - 1} йилнинг шу кунига нисбатан").strip()
    med = Side(style="medium"); thin = Side(style="thin")
    bd = Border(thin, thin, thin, thin); C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LA = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    NC = 16; last = L(NC)
    ws.merge_cells(f"A1:{last}1"); ws["A1"] = f"Бухоро вилоятида {year} йилда маҳсулот экспорт қилиш прогнозларининг бажарилиши тўғрисида тезкор маълумот"
    ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = C; ws.row_dimensions[1].height = 40
    ws.merge_cells(f"A2:{last}2"); ws["A2"] = f"{nxt} йил ҳолатига"; ws["A2"].font = Font(bold=True, size=11); ws["A2"].alignment = C
    ws.cell(3, NC, "(минг АҚШ долл.)").font = Font(italic=True, size=10); ws.cell(3, NC).alignment = Alignment(horizontal="right")
    # сарлавҳа: 4–5 қаторлар
    top = [("A", "№"), ("B", "Шаҳар ва туманлар номи"), ("C", f"{year} йил прогнози жами"), ("D", f"{year} йил январь–{mon} прогнози"),
           ("E:H", f"Жами январь–{mon} ойларида"), ("I:M", f"Шундан {mon} ойида"), ("N", prev_lbl), ("O:P", cmp_lbl)]
    nd = dt.date.fromisoformat(D)   # маълумот куни — «1 кунда» ўрнига сана (номма-ном варағидагидек)
    day_lbl = f"{nd.day} {['янв.', 'фев.', 'март', 'апр.', 'май', 'июнь', 'июль', 'авг.', 'сент.', 'окт.', 'нояб.', 'дек.'][nd.month - 1]}"
    sub = {"E": "прогноз", "F": "амалда", "G": "фарқи\n+ −", "H": "%", "I": "прогноз", "J": "амалда", "K": day_lbl, "L": "фарқи\n+ −", "M": "%", "O": "фарқи\n+ −", "P": "%"}
    for rng, txt in top:
        a, b = (rng.split(":") + [None])[:2]
        if b: ws.merge_cells(f"{a}4:{b}4")
        else: ws.merge_cells(f"{a}4:{a}5")
        ws[f"{a}4"] = txt
    for col, txt in sub.items(): ws[f"{col}5"] = txt
    HF = PatternFill("solid", fgColor="D9E7E2")
    for r in (4, 5):
        for i in range(1, NC + 1):
            c = ws.cell(r, i); c.font = Font(bold=True, size=10); c.fill = HF; c.alignment = C; c.border = bd
    ws.row_dimensions[4].height = 48; ws.row_dimensions[5].height = 30
    for col, w in zip("ABCDEFGHIJKLMNOP", [5, 34, 13, 13, 12, 12, 11, 8, 12, 12, 10, 11, 8, 14, 11, 8]): ws.column_dimensions[col].width = w
    TF = PatternFill("solid", fgColor="EEF4F1"); RF = PatternFill("solid", fgColor="DCEBE6")
    cp = {"year": 3, "m9": 4, "pp": 5, "pf": 6, "pd": 7, "pc": 8, "mp": 9, "mf": 10, "day": 11, "md": 12, "mc": 13, "prev": 14, "yd": 15, "yc": 16}
    row = _svod_write(ws, d, 6, cp)
    for r in range(6, row):
        kind = "region" if r == 6 else ("sub" if r in (7, 8) or (r - 9) % 3 else "total")
        bold = kind in ("region", "total"); fill = RF if kind == "region" else (TF if kind == "total" else None)
        for i in range(1, NC + 1):
            c = ws.cell(r, i); c.border = bd
            c.font = Font(bold=bold, size=11 if bold else 10, italic=(kind == "sub"))
            if fill: c.fill = fill
            if i == 1: c.alignment = Alignment(horizontal="center", vertical="center")
            elif i == 2: c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=0 if bold else 2)
            else: c.alignment = Alignment(horizontal="right", vertical="center")
            if i in (8, 13, 16): c.number_format = "0.0%"
            elif i == 11: c.number_format = "#,##0.00"
            elif i >= 3: c.number_format = "#,##0.0;−#,##0.0"
        ws.row_dimensions[r].height = 22 if bold else 18
    # ташқи чегара — қалин
    for r in range(4, row):
        for i in (1, NC):
            c = ws.cell(r, i); b = c.border
            c.border = Border(left=med if i == 1 else b.left, right=med if i == NC else b.right, top=b.top, bottom=b.bottom)
    for i in range(1, NC + 1):
        c = ws.cell(row - 1, i); b = c.border; c.border = Border(left=b.left, right=b.right, top=b.top, bottom=med)
        c = ws.cell(4, i); b = c.border; c.border = Border(left=b.left, right=b.right, top=med, bottom=b.bottom)
    src = (d.get("prev_year_source") or {}).get("text")
    mtxt = "Бухоро корхоналари экспорт қилган (етиштирилган жойидан қатъи назар)" if d.get("meva_mode") == "exporter" else "карантин рўйхати (Бухорода етиштирилган)"
    ws.cell(row + 1, 2, f"Изоҳ: {src + '.' if src else 'ўтган йил кўрсаткичи ' + str(d.get('prev_year_days')) + ' кунга мутаносиб ҳисобланган.'} Мева-сабзавот — {mtxt} бўйича.").font = Font(italic=True, size=9, color="666666")
    ws.freeze_panes = "C6"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:5"

SVOD_KEYS_VAL = ("year", "m9", "pp", "pf", "mp", "mf", "day", "prev")   # фақат туманнинг саноат/мева қаторида рақам, қолгани формула

def _svod_write(ws, d: dict, first: int, cp: dict, sub_name_col: int | None = None):
    """СВОД варағининг маълумот қисми: рақамлар фақат ҳар туманнинг «саноат» ва «мева-сабзавот» қаторларида;
    туман жами, вилоят жами/саноат/мева, фарқ ва фоизлар — варақ ичидаги формулалар."""
    from openpyxl.utils import get_column_letter as L
    col = {k: L(v) for k, v in cp.items()}
    def raw(r, f, p):
        p = p or {}
        vals = {"year": p.get("year"), "m9": p.get("period"), "pp": p.get("period"), "pf": f["ytd"], "mp": p.get("month"),
                "mf": f["mon"], "day": f["day"], "prev": p.get("prev")}
        for k in SVOD_KEYS_VAL:
            v = vals[k]; ws.cell(r, cp[k]).value = round(v, 6) if isinstance(v, (int, float)) else (0 if v is None and k != "day" else v)
    def derived(r):
        c = lambda k: f"{col[k]}{r}"
        ws.cell(r, cp["pd"]).value = f"={c('pf')}-{c('pp')}"
        ws.cell(r, cp["pc"]).value = f'=IF({c("pp")}=0,"",{c("pf")}/{c("pp")})'
        ws.cell(r, cp["md"]).value = f"={c('mf')}-{c('mp')}"
        ws.cell(r, cp["mc"]).value = f'=IF({c("mp")}=0,"",{c("mf")}/{c("mp")})'
        ws.cell(r, cp["yd"]).value = f"={c('pf')}-{c('prev')}"
        ws.cell(r, cp["yc"]).value = f'=IF({c("prev")}=0,"",{c("pf")}/{c("prev")})'
    def summ(r, rows):
        for k in SVOD_KEYS_VAL: ws.cell(r, cp[k]).value = "=" + "+".join(f"{col[k]}{x}" for x in rows)
    n = len(d["districts"]); R0 = first
    ws.cell(R0, 1).value = n; ws.cell(R0, 2).value = "Бухоро вилояти жами"
    ws.cell(R0 + 1, 2).value = "Саноат маҳсулотлари жами"; ws.cell(R0 + 2, 2).value = "Мева-сабзавотлар жами"
    row = R0 + 3; san_rows, mev_rows = [], []
    for i, x in enumerate(d["districts"], 1):
        ws.cell(row, 1).value = i; ws.cell(row, 2).value = x["name"]
        ws.cell(row + 1, 2).value = "Саноат маҳсулотлари"; ws.cell(row + 2, 2).value = "Мева-сабзавотлар"
        if sub_name_col: ws.cell(row, sub_name_col).value = None; ws.cell(row + 1, sub_name_col).value = x["name"]; ws.cell(row + 2, sub_name_col).value = x["name"]
        raw(row + 1, _svod_agg(d, x["code"], "sanoat"), x["plan"].get("sanoat"))
        raw(row + 2, _svod_agg(d, x["code"], "meva"), x["plan"].get("meva"))
        summ(row, [row + 1, row + 2])
        for rr in (row, row + 1, row + 2): derived(rr)
        san_rows.append(row + 1); mev_rows.append(row + 2); row += 3
    summ(R0 + 1, san_rows); summ(R0 + 2, mev_rows); summ(R0, [R0 + 1, R0 + 2])
    rp = d.get("region_plan") or {}   # вилоят режаси — расмий рақам (туманлар йиғиндисидан яхлитлашда фарқ қилиши мумкин)
    for rr, k in ((R0 + 1, "sanoat"), (R0 + 2, "meva")):
        p = rp.get(k) or {}
        for key, pk in (("year", "year"), ("m9", "period"), ("pp", "period"), ("mp", "month")):
            if isinstance(p.get(pk), (int, float)): ws.cell(rr, cp[key]).value = round(p[pk], 6)
    for rr in (R0, R0 + 1, R0 + 2): derived(rr)
    return row


def _svod_agg(d: dict, code, kind):
    cm = d["month"]
    cs = [c for c in d["companies"] if c["t"] == "sanoat" and (not code or c["d"] == code)]
    ds = [x for x in d["districts"] if not code or x["code"] == code]
    s = {"ytd": sum(c["total"] for c in cs), "mon": sum((c["m"][cm - 1] if len(c["m"]) >= cm else 0) or 0 for c in cs), "day": sum(c["day"] for c in cs)}
    m = {"ytd": sum(x["meva"]["ytd"] for x in ds), "mon": sum(x["meva"]["month"] for x in ds), "day": sum(x["meva"]["day"] for x in ds)}
    if kind == "sanoat": return s
    if kind == "meva": return m
    return {k: s[k] + m[k] for k in s}

def _svod_sheet(ws, con, d: dict):
    """«СВОД (карантин)» — ўз жадвалингиздаги (шаблон) варақнинг айнан кўриниши: услуб, чегара, кенглик, баландлик, бирлаштирилган
    катаклар шаблондан нусхаланади; рақамлар фақат туманнинг саноат/мева қаторларида, жамилар, фарқ ва фоизлар — варақ ичидаги формулалар (24.09.2026). Шаблон бўлмаса — оддий кўриниш."""
    from copy import copy
    from openpyxl.utils import get_column_letter as L
    from . import tpl, xl
    tp = tpl.template_path(con)
    if not tp.exists(): return _svod_sheet_plain(ws, con, d)
    import openpyxl
    src_wb = openpyxl.load_workbook(tp)
    kname = xl.find_sheet(src_wb.sheetnames, tpl.KARANTIN_SVOD)
    if not kname: return _svod_sheet_plain(ws, con, d)
    src = src_wb[kname]
    try: lay = tpl.svod_layout(src)
    except Exception: return _svod_sheet_plain(ws, con, d)
    hr, sub = lay["header_row"], lay["sub_row"]
    NC = 17   # A..Q
    # маълумот қаторлари: биринчи устунда «13» (вилоят) дан бошлаб 3 + 13×3 қатор
    first = next((r for r in range(sub + 1, sub + 6) if isinstance(src.cell(r, 1).value, (int, float))), sub + 1)
    n_dist = len(d["districts"]); last = first + 3 + 3 * n_dist   # охирги маълумот қатори + 1 (бўш қатор)
    # ---- нусха: услуб, қиймат (A–C матнлари), баландлик, кенглик, бирлаштириш
    for r in range(1, last + 1):
        if src.row_dimensions[r].height: ws.row_dimensions[r].height = src.row_dimensions[r].height
        for c in range(1, NC + 1):
            sc = src.cell(r, c); tc = ws.cell(r, c)
            if sc.has_style:
                tc.font = copy(sc.font); tc.fill = copy(sc.fill); tc.border = copy(sc.border)
                tc.alignment = copy(sc.alignment); tc.number_format = sc.number_format; tc.protection = copy(sc.protection)
            v = sc.value
            if isinstance(v, str) and v.startswith("="): v = None   # формулалар ёзилмайди
            if r < first or c <= 3: tc.value = v                     # сарлавҳа ва A–C матнлари шаблондан
    for col, dim in src.column_dimensions.items():
        if dim.min and dim.min > NC: continue
        nd_ = ws.column_dimensions[col]; nd_.width = dim.width; nd_.hidden = dim.hidden
        if dim.min and dim.max: nd_.min, nd_.max = dim.min, min(dim.max, NC)
    for mr in src.merged_cells.ranges:
        if mr.max_col <= NC and mr.max_row <= last: ws.merge_cells(str(mr))
    try:   # шартли форматлаш (қизил % ва ҳ.к.) ҳам шаблондан
        for cf in src.conditional_formatting:
            for rule in cf.rules: ws.conditional_formatting.add(str(cf.sqref), copy(rule))
    except Exception: pass
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines; ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    ws.page_setup.orientation = src.page_setup.orientation; ws.page_setup.paperSize = src.page_setup.paperSize
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = copy(src.page_margins); ws.print_area = f"A1:{L(NC)}{last}"; ws.print_title_rows = f"{hr}:{sub}"
    # ---- сарлавҳалар (сана ва ой)
    D = d["date"]; year = d["year"]; cm = d["month"]; mon = MONTH_NAMES[cm - 1].lower()
    nxt = (dt.date.fromisoformat(D) + dt.timedelta(days=1)).strftime("%d.%m.%Y")
    nd = dt.date.fromisoformat(D)
    day_lbl = f"{nd.day} {['янв.', 'фев.', 'март', 'апр.', 'май', 'июнь', 'июль', 'авг.', 'сент.', 'окт.', 'нояб.', 'дек.'][nd.month - 1]}"
    top = {c.column: c.value for c in src[hr] if isinstance(c.value, str)}
    for col, v in top.items():
        lv = v.lower()
        if lv.startswith("жами январ"): ws.cell(hr, col).value = f"Жами январь-{mon} ойларида"
        elif lv.startswith("шундан"): ws.cell(hr, col).value = f"Шундан {mon} ойида"
        elif "ойлик прогноз" in lv: ws.cell(hr, col).value = f"{year} йил \n{cm} ойлик прогноз"
        elif "ҳолатида" in lv and (d.get("prev_col_label") or d.get("prev_year_label")): ws.cell(hr, col).value = d.get("prev_col_label") or d["prev_year_label"]
        elif "нисбатан" in lv and d.get("prev_cmp_label"): ws.cell(hr, col).value = d["prev_cmp_label"]
    for c in src[sub]:
        if xl._norm(c.value).startswith("1 кунда"): ws.cell(sub, c.column).value = day_lbl
    # сана қатори: шаблонда =+номманом!A2 бўлган катак
    for r in range(1, hr):
        for c in range(1, NC + 1):
            v = src.cell(r, c).value
            if isinstance(v, str) and v.startswith("=") and "a2" in v.lower(): ws.cell(r, c).value = f"{nxt} йил"
    # ---- устунлар (сарлавҳа бўйича топилади; топилмаса шаблон тартиби D..Q)
    cp = {"year": 4, "m9": 5, "pp": 6, "pf": 7, "pd": 8, "pc": 9, "mp": 10, "mf": 11, "day": 12, "md": 13, "mc": 14, "prev": 15, "yd": 16, "yc": 17}
    if lay.get("period_fact"): cp.update({"pp": lay["period_fact"] - 1, "pf": lay["period_fact"], "pd": lay["period_fact"] + 1, "pc": lay["period_fact"] + 2})
    if lay.get("month_fact"): cp.update({"mp": lay["month_fact"] - 1, "mf": lay["month_fact"], "day": lay["month_fact"] + 1, "md": lay["month_fact"] + 2, "mc": lay["month_fact"] + 3})
    _svod_write(ws, d, first, cp, sub_name_col=3)
    ws.freeze_panes = ws.cell(sub + 1, 4)

def export_svod_nomma(con, q: dict, out):
    """Битта китоб: 1-варақ «СВОД (карантин)» + ҳар туман учун алоҳида «номма-ном» варағи (саноат + мева, туман номи билан)."""
    import re, openpyxl
    from . import svodx
    fs = fset_param(q.get("fs")) if (q.get("prev") or "settings") == "settings" else None
    from . import plans as _pl
    _pl.require_cover(con, q.get("date") or last_data_date(con), plan_param(q.get("plan")))   # мос режа йўқ — юклаб бўлмайди
    d = dashboard(con, q.get("date"), plan_param(q.get("plan")), fs)
    if d.get("empty"): raise ValueError("База бўш")
    svodx.apply(con, d, q.get("prev") or "settings", q.get("meva") or "karantin", q.get("sp"), fs)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "СВОД (карантин)" if d["meva_mode"] == "karantin" else "СВОД"
    _svod_sheet(ws, con, d)
    used = {ws.title}
    for x in d["districts"]:
        name = re.sub(r'[\\/*?:\[\]]', " ", x["name"]).strip()[:31] or x["code"]
        while name in used: name = name[:29] + " 2"
        used.add(name)
        _nomma_sheet(wb.create_sheet(name), d, x["code"], "all", "")
    wb.save(out); return out
