"""Excel on the user's own workbook (Sozlamalar -> «2. O'z jadvalingiz», .xlsx).

Rule (11.09.2026): the site fills only номманом; every other sheet is computed by the workbook's own formulas.
Exactly like the manual routine, the site also writes:
  * СВОД (карантин) – meva-sabzavot rows: YTD "+X", month "+X", «1 кунда»; karantin month sheet: GTD rows
  * СВОД … млн      – БНПЗ: Қоровулбозор sanoat YTD "+X"
  * СВОД sheets     – day term of the prev-year formula  =O7/273*(151+30+31+31+6)
Nothing else is changed: formulas, external links, styles stay as the user left them.
"""
from __future__ import annotations
import re, copy, json, datetime as dt
from pathlib import Path
from collections import defaultdict
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer, Token
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from . import db
from .xl import _norm, inn_str, find_sheet, sheet_key

NOMMA = "номманом"
KARANTIN_SVOD = "СВОД (карантин)"
GOVERNOR_SVOD = "СВОД млн"
BNPZ_DISTRICT = "қоровулбозор тумани"
YELLOW = "FFFFFF00"
SITE_AUTHOR = "сайт"
MONTH_STEMS = ["январ", "феврал", "март", "апрел", "май", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр"]
MONTH_UZ = ["yanvar", "fevral", "mart", "aprel", "may", "iyun", "iyul", "avgust", "sentyabr", "oktyabr", "noyabr", "dekabr"]

def template_path(con=None) -> Path:
    """Every applied template is kept as data/templates/<label>_<sha8>.xlsx and the DB points at it,
    so restoring a DB backup also brings back the matching template."""
    store = db.get_setting(con, "template_store") if con is not None else None
    if store and (db.BASE / store).exists(): return db.BASE / store
    return db.BASE / "template.xlsx"

def dmy(iso: str) -> str:
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}" if iso else "—"

def clean(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

def sanitize(src: Path, dst: Path):
    """Excel sometimes stores \\x1a inside external-link XML; openpyxl cannot parse it."""
    import zipfile
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith((".xml", ".rels")): data = data.replace(b"\x1a", b"")
            zout.writestr(item, data)

EXCEL_BUSY = ("0x800AC472", "0x80010001", "-2146777998", "-2147418111", "EXCEL_BUSY")

_XL_PS = r"""$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
try { [Console]::OutputEncoding = [System.Text.Encoding]::UTF8 } catch {}
$busy = @(-2147418111, -2146777998)   # RPC_E_CALL_REJECTED, VBA_E_IGNORE: Excel is busy (dialog, loading, update)
function Invoke-Xl([scriptblock]$Block, [int]$Tries = 60) {
  for ($i = 0; $i -lt $Tries; $i++) {
    try { return (& $Block) }
    catch {
      $e = $_.Exception; $isBusy = $false
      while ($e) { if ($busy -contains $e.HResult) { $isBusy = $true }; $e = $e.InnerException }
      if (-not $isBusy) { throw }
      Start-Sleep -Milliseconds 500
    }
  }
  throw 'EXCEL_BUSY'
}
$xl = $null; $wb = $null; $xlPid = 0; $code = 0
try {
  $xl = __CREATE__
  try {
    Add-Type -Namespace XlConv -Name Win -MemberDefinition '[DllImport("user32.dll")] public static extern int GetWindowThreadProcessId(System.IntPtr h, out int p);'
    $hwnd = Invoke-Xl { $xl.Hwnd }
    [void][XlConv.Win]::GetWindowThreadProcessId([System.IntPtr][int64]$hwnd, [ref]$xlPid)
  } catch {}
  Invoke-Xl { $xl.Visible = $false }
  Invoke-Xl { $xl.DisplayAlerts = $false }
  Invoke-Xl { $xl.AskToUpdateLinks = $false }
  Invoke-Xl { $xl.EnableEvents = $false }
  $m = [System.Reflection.Missing]::Value
  # UpdateLinks 0, ReadOnly, no password prompt (''), ignore read-only recommendation, no "file in use" notify
  $wb = Invoke-Xl { $xl.Workbooks.Open('__SRC__', 0, $true, $m, '', '', $true, $m, $m, $false, $false) }
  Invoke-Xl { $wb.SaveAs('__DST__', 51) }
  Invoke-Xl { $wb.Close($false) }
  Write-Output 'OK'
} catch {
  $code = 1
  $e = $_.Exception; $msg = $e.Message; $hr = ''
  while ($e) { if ($e.HResult -and $e.HResult -ne -2146233087) { $hr = ('0x{0:X8}' -f $e.HResult) }; if ($e.InnerException) { $msg = $e.InnerException.Message }; $e = $e.InnerException }
  Write-Output ('ERR: ' + $hr + ' ' + ($msg -replace '\s+', ' '))
} finally {
  if ($xl) {
    try { Invoke-Xl { $xl.Quit() } 10 } catch {}
    try { [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($xl) } catch {}
  }
  if ($xlPid) {   # only the Excel started here; the user's own Excel windows are not touched
    try { $pr = Get-Process -Id $xlPid -ErrorAction Stop; if (-not $pr.WaitForExit(8000)) { Stop-Process -Id $xlPid -Force } } catch {}
  }
}
exit $code
"""

def excel_error(stdout: bytes | str, stderr: bytes | str) -> str:
    """Readable reason from the conversion script (its ERR: line, else PowerShell's CLIXML error text)."""
    dec = lambda b: b.decode("utf-8", "ignore") if isinstance(b, (bytes, bytearray)) else (b or "")
    out, err = dec(stdout), dec(stderr)
    line = next((l[4:].strip() for l in out.splitlines() if l.startswith("ERR:")), "")
    if not line and err:
        parts = re.findall(r'<S S="Error">(.*?)</S>', err, re.S)
        line = re.sub(r"_x000[DA]_", " ", " ".join(parts) if parts else err)
        line = re.sub(r"\s+", " ", line).strip()
    if any(k in line for k in EXCEL_BUSY):
        return ("Excel band yoki javob bermayapti (ichida ochiq oyna/xabar bo'lishi mumkin). Barcha Excel oynalarini yoping, "
                "Диспетчер задач da qolib ketgan EXCEL.EXE bo'lsa yakunlang; Excelni bir marta qo'lda ochib, chiqqan xabarlarni "
                "(yangilanish, faollashtirish) yoping va yana «Tekshirish» ni bosing")
    if "80040154" in line or "ComObject" in line:
        return "Bu kompyuterda Microsoft Excel topilmadi"
    return (line[:300] or "Excel faylni saqlamadi")

def xls_to_xlsx(src: Path, dst: Path, timeout: int = 240) -> None:
    """Convert the user's .xls with the Excel installed on this Windows computer (formulas, links and styles stay Excel-native)."""
    import os, subprocess, base64
    if os.name != "nt": raise RuntimeError("Excel faqat Windows kompyuterida mavjud")
    src, dst = Path(src).resolve(), Path(dst).resolve()
    dst.unlink(missing_ok=True)
    q = lambda p: str(p).replace("'", "''")
    ps = _XL_PS.replace("__CREATE__", "New-Object -ComObject Excel.Application").replace("__SRC__", q(src)).replace("__DST__", q(dst))
    enc = base64.b64encode(ps.encode("utf-16-le")).decode()
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-EncodedCommand", enc],
                           capture_output=True, timeout=timeout, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired:
        raise RuntimeError(excel_error("ERR: EXCEL_BUSY", ""))
    if not dst.exists() or dst.stat().st_size == 0:
        raise RuntimeError(excel_error(r.stdout, r.stderr))

def load(path: Path):
    return openpyxl.load_workbook(str(path), keep_links=True)

# ------------------------------------------------------------------ formula references
_CELL = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")
_ADDR = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+(:\$?[A-Za-z]{1,3}\$?\d+)?$")
_ROWS = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")

def _split_ref(ref: str):
    """'[1]Лист2'!$A$5 -> (ext='1', sheet='Лист2', addr='$A$5'); A1 -> (None, None, 'A1')."""
    if "!" not in ref: return None, None, ref
    sheet, addr = ref.rsplit("!", 1)
    s = sheet[1:-1].replace("''", "'") if sheet.startswith("'") and sheet.endswith("'") else sheet
    m = re.match(r"^\[(\d+)\](.*)$", s)
    return (m.group(1), m.group(2), addr) if m else (None, s, addr)

def _map_rows(addr: str, fn) -> str:
    m = _ROWS.match(addr)
    if m: return f"{m.group(1)}{fn(int(m.group(2)), m.group(1))}:{m.group(3)}{fn(int(m.group(4)), m.group(3))}"
    if not _ADDR.match(addr): return addr
    return _CELL.sub(lambda x: f"{x.group(1)}{x.group(2)}{x.group(3)}{fn(int(x.group(4)), x.group(3))}", addr)

def _rewrite(formula: str, host: str, target: str | None, fn, external=False) -> str:
    """Apply fn(row, dollar) to every row number referenced on sheet `target` (None = any sheet incl. external)."""
    if not isinstance(formula, str) or not formula.startswith("="): return formula
    items = Tokenizer(formula).items
    out, changed = [], False
    for t in items:
        v = t.value
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ext, sh, addr = _split_ref(v)
            hit = target is None or (ext is None and ((sh is None and host == target) or sh == target))
            if hit and (ext is None or external):
                na = _map_rows(addr, fn)
                if na != addr: v = v[: len(v) - len(addr)] + na; changed = True
        out.append(v)
    return "=" + "".join(out) if changed else formula

def shift_formula(formula, host, target, at, n=1):
    return _rewrite(formula, host, target, lambda r, d: r + n if r >= at else r)

def copy_down(formula, host, delta):
    """Excel copy/paste of a formula `delta` rows lower: relative rows move, $rows stay (external refs too)."""
    return _rewrite(formula, host, None, lambda r, d: r if d else r + delta, external=True)

def insert_row(wb, ws, at: int):
    """Insert one row at `at` in ws and keep every reference in the workbook pointing at the same cells."""
    title = ws.title
    ws.insert_rows(at, 1)
    from openpyxl.worksheet.formula import ArrayFormula
    for s in wb.worksheets:
        for row in s.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and (s.title == title or title in v):
                    nv = shift_formula(v, s.title, title, at)
                    if nv != v: c.value = nv
                elif isinstance(v, ArrayFormula) and v.text:
                    v.text = shift_formula(v.text, s.title, title, at)
                    if s.title == title and v.ref: v.ref = _map_rows(v.ref, lambda r, d: r + 1 if r >= at else r)
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= at: mr.shift(0, 1)
        elif mr.max_row >= at: mr.expand(down=1)
    dims = ws.row_dimensions
    moved = sorted([k for k in list(dims.keys()) if k >= at], reverse=True)
    for k in moved:
        d = dims[k]; del dims[k]; d.index = k + 1; dims[k + 1] = d
    if ws.auto_filter.ref: ws.auto_filter.ref = _map_rows(ws.auto_filter.ref, lambda r, d: r + 1 if r >= at else r)
    pa = ws.print_area
    if pa:
        parts = []
        for p in (pa if isinstance(pa, list) else str(pa).split(",")):
            addr = p.split("!")[-1]
            parts.append(_map_rows(addr, lambda r, d: r + 1 if r >= at else r))
        ws.print_area = parts
    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(" ".join(_map_rows(str(r), lambda r, d: r + 1 if r >= at else r) for r in dv.sqref.ranges))

# ------------------------------------------------------------------ layout of the user's workbook
def _month_of(text) -> int | None:
    s = _norm(text)
    return next((i + 1 for i, st in enumerate(MONTH_STEMS) if st in s), None)

def _fill(c):
    try:
        return c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type == "solid" else None
    except Exception:
        return None

def parse_label(v) -> str | None:
    if isinstance(v, (dt.date, dt.datetime)): return (v.date() if isinstance(v, dt.datetime) else v).isoformat()
    m = re.search(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", str(v or ""))
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat() if m else None

def nomma_layout(wb) -> dict:
    if NOMMA not in wb.sheetnames: raise ValueError("Jadvalda «номманом» varag'i yo'q")
    ws = wb[NOMMA]
    hr = next((r for r in range(1, 12) if any(_norm(c.value) == "инн" for c in ws[r])), None)
    if not hr: raise ValueError("номманом: «ИНН» sarlavhasi topilmadi")
    hdr = {c.column: _norm(c.value) for c in ws[hr] if c.value not in (None, "")}
    def col(pred, what):
        c = next((k for k, v in hdr.items() if pred(v)), None)
        if c is None: raise ValueError(f"номманом: «{what}» ustuni topilmadi")
        return c
    c_inn = col(lambda v: v == "инн", "ИНН"); c_name = col(lambda v: v.startswith("корхона номи"), "Корхона номи")
    c_hud = col(lambda v: v.startswith("ҳудуди"), "Ҳудуди"); c_total = col(lambda v: v.startswith("экспорти ҳажми жами"), "Экспорти ҳажми жами")
    c_day = col(lambda v: v.startswith("1 кунда"), "1 кунда"); c_tar = col(lambda v: v.startswith("тармоқ"), "Тармоқ")
    c_prod = col(lambda v: v.startswith("маҳсулоти"), "Маҳсулоти")
    months = {k: _month_of(v) for k, v in hdr.items() if "ойида" in v and _month_of(v)}
    left = [k for k in months if k < c_day]
    if not left: raise ValueError("номманом: joriy oy ustuni («шундан … ойида») topilmadi")
    c_month = max(left); cur_month = months[c_month]
    # district / category columns: taken from the SUMIFS of the karantin svod, fallback by content
    c_dist = c_cat = None
    kname = find_sheet(wb.sheetnames, KARANTIN_SVOD)
    if kname:
        rx = re.compile(r"SUMIFS\(номманом!\$?([A-Z]+):\$?[A-Z]+,номманом!\$?([A-Z]+):\$?[A-Z]+,[^,]+,номманом!\$?([A-Z]+):\$?[A-Z]+,")
        for row in wb[kname].iter_rows(min_row=1, max_row=60):
            for c in row:
                m = rx.search(c.value) if isinstance(c.value, str) else None
                if m:
                    from openpyxl.utils import column_index_from_string as CI
                    c_dist, c_cat = CI(m.group(2)), CI(m.group(3)); break
            if c_dist: break
    label_row = next((r for r in range(1, hr) if parse_label(ws.cell(r, 1).value)), None)
    label = parse_label(ws.cell(label_row, 1).value) if label_row else None
    region = {}
    blocks, companies, tail, dup = [], defaultdict(dict), {}, []
    jami = None; cur = None; yellow_cols = None; cat_text = {}
    maxr = ws.max_row
    sum_rx = re.compile(r"SUM\(\$?[A-Z]+\$?(\d+):\$?[A-Z]+\$?(\d+)\)")
    for r in range(hr + 1, maxr + 1):
        inn = inn_str(ws.cell(r, c_inn).value); name = ws.cell(r, c_name).value
        nname = _norm(name)
        if not inn:
            if nname.startswith("вилоят бўйича жами"): region["total"] = r; continue
            if nname.startswith("саноат") and "total" in region and "sanoat" not in region: region["sanoat"] = r; continue
            if nname.startswith("мева") and "total" in region and "meva" not in region: region["meva"] = r; continue
            if nname == "жами": jami = r; cur = None; continue
            f = ws.cell(r, c_total).value
            m = sum_rx.search(f) if isinstance(f, str) else None
            if m and nname and jami is None:
                cur = {"name": str(name).strip(), "key": nname, "header": r, "first": int(m.group(1)), "last": int(m.group(2)), "rows": []}
                blocks.append(cur)
            continue
        if jami is not None:
            tail[inn] = r; continue
        if cur is None or not (cur["first"] <= r <= cur["last"]):
            continue
        cat_v = ws.cell(r, c_cat).value if c_cat else None
        cat = "meva" if _norm(cat_v).startswith("мева") else "sanoat"
        if cat_v: cat_text.setdefault(cat, cat_v)
        if cat in companies[inn]: dup.append((inn, r))
        companies[inn].setdefault(cat, r)
        yel = _fill(ws.cell(r, c_name)) == YELLOW
        if yel and yellow_cols is None:
            yellow_cols = [c.column for c in ws[r] if _fill(c) == YELLOW]
        cur["rows"].append({"row": r, "inn": inn, "cat": cat, "yellow": yel, "name": name})
    if not blocks: raise ValueError("номманом: tuman bloklari (SUM formulali sarlavha qatorlari) topilmadi")
    return {"sheet": NOMMA, "header_row": hr, "label_row": label_row, "label": label,
            "cols": {"inn": c_inn, "name": c_name, "hudud": c_hud, "total": c_total, "day": c_day, "month": c_month, "tarmoq": c_tar, "product": c_prod,
                     "district": c_dist, "category": c_cat, "months": sorted(months)},
            "cur_month": cur_month, "region": region, "blocks": blocks, "companies": dict(companies), "tail": tail, "jami": jami,
            "yellow_cols": yellow_cols, "cat_text": cat_text, "dups": dup}

def svod_layout(ws) -> dict:
    """Karantin / governor svod: header «прогнози», sub-header with «амалда» x2, period and month groups."""
    hr = next((r for r in range(1, 9) if any(isinstance(c.value, str) and "прогнози" in c.value.lower() for c in ws[r])), None)
    if not hr: raise ValueError(f"«{ws.title}»: «прогнози» sarlavhasi topilmadi")
    sub = next((r for r in range(hr + 1, hr + 6) if sum(1 for c in ws[r] if _norm(c.value).startswith("амалда")) >= 2), None)
    if not sub: raise ValueError(f"«{ws.title}»: «амалда» ustunlari topilmadi")
    top = {c.column: _norm(c.value) for c in ws[hr] if c.value}
    subs = {c.column: _norm(c.value) for c in ws[sub] if c.value}
    c_period = next((k for k, v in top.items() if v.startswith("жами январ")), None)
    c_month = next((k for k, v in top.items() if re.match(r"шундан .* ойида", v)), None)
    if not c_period or not c_month: raise ValueError(f"«{ws.title}»: «Жами январь…» / «Шундан … ойида» topilmadi")
    after = lambda start, label: next((k for k in sorted(subs) if k >= start and subs[k].startswith(label)), None)
    return {"header_row": hr, "sub_row": sub, "period_fact": after(c_period, "амалда"), "month_fact": after(c_month, "амалда"),
            "day_fact": after(c_month, "1 кунда"), "month": _month_of(top[c_month]), "month_label": ws.cell(hr, c_month).value}

_PREV = re.compile(r"/(\d+)\*\((\s*\d+(?:\s*\+\s*\d+)*\s*)\)")
def _doy(iso): return dt.date.fromisoformat(iso).timetuple().tm_yday

def prev_year_cells(wb, through: str):
    """Cells with /N*(a+b+...) whose day count equals the template date (maintained by the user)."""
    out, stale = [], 0
    for sh in wb.worksheets:
        if not sheet_key(sh.title).startswith("свод"): continue
        for row in sh.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and "*(" in v:
                    m = _PREV.search(v)
                    if not m: continue
                    terms = [int(x) for x in m.group(2).split("+")]
                    if sum(terms) == _doy(through): out.append((sh.title, c.coordinate))
                    else: stale += 1
    return out, stale

def karantin_month_sheet(wb, month: int):
    stem = MONTH_STEMS[month - 1]
    return next((s for s in wb.sheetnames if "карантин" in s.lower() and stem in s.lower() and not sheet_key(s).startswith("свод")), None)

def inspect(path: Path) -> dict:
    """Everything the check screen shows + what export needs. Raises ValueError with a user message."""
    wb = load(path)
    lay = nomma_layout(wb)
    if not lay["label"]: raise ValueError("номманом A2 katagida sana topilmadi (masalan «07.09.2026 йил»)")
    label = lay["label"]; through = (dt.date.fromisoformat(label) - dt.timedelta(days=1)).isoformat()
    tm = int(through[5:7]); last_of_month = (dt.date.fromisoformat(through) + dt.timedelta(days=1)).day == 1
    if not (lay["cur_month"] == tm or (last_of_month and lay["cur_month"] == tm % 12 + 1)):
        raise ValueError(f"Jadval sanasi {dmy(label)} ({dmy(through)} gacha ma'lumot), lekin номманом joriy oy ustuni «{MONTH_UZ[lay['cur_month']-1]}». Sana va oy mos emas.")
    kname = find_sheet(wb.sheetnames, KARANTIN_SVOD)
    if not kname: raise ValueError("«СВОД (карантин) …» varag'i topilmadi")
    kl = svod_layout(wb[kname])
    kws = wb[kname]
    meva_rows = {}
    for r in range(kl["sub_row"] + 1, kws.max_row + 1):
        b = _norm(kws.cell(r, 2).value); cn = _norm(kws.cell(r, 3).value)
        if b.startswith("мева") and cn: meva_rows[cn] = r
    gname = find_sheet(wb.sheetnames, GOVERNOR_SVOD); bn = None
    if gname:
        gl = svod_layout(wb[gname]); gws = wb[gname]
        row = next((r for r in range(gl["sub_row"] + 1, gws.max_row + 1) if _norm(gws.cell(r, 2).value).startswith("саноат") and _norm(gws.cell(r, 3).value) == BNPZ_DISTRICT), None)
        if row and gl["period_fact"]: bn = {"sheet": gname, "cell": f"{L(gl['period_fact'])}{row}", "row": row, "col": gl["period_fact"]}
    prev, stale = prev_year_cells(wb, through)
    n_companies = sum(len(b["rows"]) for b in lay["blocks"])
    yellow = sorted({x["inn"] for b in lay["blocks"] for x in b["rows"] if x["yellow"]})
    # formulas must survive tokenizer round-trip, otherwise row insertion is unsafe
    bad = 0
    for s in wb.worksheets:
        for row in s.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    try:
                        if "=" + "".join(t.value for t in Tokenizer(v).items) != v: bad += 1
                    except Exception: bad += 1
    lost = []
    import zipfile
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        if any(n.startswith("xl/charts/") for n in names): lost.append("diagramma")
        if any(n.startswith("xl/media/") for n in names): lost.append("rasm")
        # comments the site wrote itself («Бухоро эмас — сайт бўйича», author «сайт») survive openpyxl; other comments are reported
        if any("comments" in n for n in names) and any(c.comment is not None and c.comment.author != SITE_AUTHOR for s in wb.worksheets for row in s.iter_rows() for c in row):
            lost.append("izoh (примечание)")
        if any(re.search(rb"<xdr:(sp|cxnSp|grpSp)\b", z.read(n)) for n in names if n.startswith("xl/drawings/drawing") and n.endswith(".xml")): lost.append("shakl (фигура)")
    return {"label": label, "through": through, "cur_month": lay["cur_month"], "lost_objects": lost, "karantin_sheet": kname, "karantin_meva_rows": len(meva_rows),
            "karantin_month": kl["month"], "karantin_month_sheet": karantin_month_sheet(wb, lay["cur_month"]),
            "governor": bn, "prev_year_cells": len(prev), "prev_year_stale": stale, "external_links": len(wb._external_links),
            "companies": n_companies, "blocks": len(lay["blocks"]), "yellow": yellow, "tail": sorted(lay["tail"]), "dups": [d[0] for d in lay["dups"]],
            "formula_bad": bad, "sheets": wb.sheetnames,
            "rows": [{"inn": x["inn"], "cat": x["cat"], "row": x["row"], "yellow": x["yellow"], "district": b["name"]} for b in lay["blocks"] for x in b["rows"]]}

# ------------------------------------------------------------------ writing
def _num_text(x: float) -> str:
    s = f"{round(x, 3):.3f}".rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s

def _existing_text(v) -> str:
    if isinstance(v, bool): raise ValueError
    if isinstance(v, int): return str(v)
    if isinstance(v, float): return repr(v) if v != int(v) else str(int(v))
    raise ValueError

def append_terms(cell, amounts) -> bool:
    """Manual style: first amount as a number, then  =a+b+c ."""
    terms = [_num_text(a) for a in amounts if round(a, 3) != 0]
    if not terms: return False
    v = cell.value
    if v is None or (isinstance(v, str) and not v.strip()):
        nv = float(terms[0]) if len(terms) == 1 else "=" + "+".join(terms)
    elif isinstance(v, str) and v.startswith("="):
        nv = v + "+" + "+".join(terms)
    elif isinstance(v, (int, float)):
        nv = "=" + _existing_text(v) + "+" + "+".join(terms)
    else:
        raise ValueError(f"{cell.parent.title}!{cell.coordinate}: katakda matn bor («{v}») — summa qo'shib bo'lmaydi")
    if isinstance(nv, str): nv = nv.replace("+-", "-")
    cell.value = nv
    return True

def _is_sum_formula(v):
    return isinstance(v, str) and v.startswith("=") and ("SUM(" in v.upper() or "!" in v)

def data_days(con, T: str, D: str) -> list[str]:
    days = set()
    for (rd,) in con.execute("SELECT report_dates FROM uploads WHERE status='saved'"):
        for d in json.loads(rd or "[]"):
            if T < d <= D: days.add(d)
    return sorted(days)

_ARITH = re.compile(r"^=\+?[\d\s.+\-*/()]+$")
def _cell_number(v) -> float | None:
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, str) and _ARITH.match(v):
        try: return float(eval(v[1:], {"__builtins__": {}}, {}))
        except Exception: return None
    return None

def company_totals(path: Path) -> dict:
    """(inn, category) -> year-to-date total of a номманом row: cached F if Excel saved it, else the sum of the month cells."""
    wb = load(path); lay = nomma_layout(wb); ws = wb[NOMMA]
    try:
        wv = openpyxl.load_workbook(str(path), data_only=True, keep_links=False)[NOMMA]
    except Exception:
        wv = None
    cols = lay["cols"]; out = {}
    for b in lay["blocks"]:
        for x in b["rows"]:
            cached = wv.cell(x["row"], cols["total"]).value if wv is not None else None
            if isinstance(cached, (int, float)) and not isinstance(cached, bool):
                tot = float(cached)
            else:
                tot = sum((_cell_number(ws.cell(x["row"], c).value) or 0.0) for c in cols["months"])
            key = (x["inn"], x["cat"]); out[key] = out.get(key, 0.0) + tot
            out.setdefault("_names", {})[x["inn"]] = x["name"]
    return out

def compare_templates(con, old_path: Path, old_T: str, new_path: Path, new_T: str) -> dict:
    """Previous template + GTD saved on the site for (old_T, new_T]  vs  the new workbook, company by company."""
    old, new = company_totals(old_path), company_totals(new_path)
    names = {**old.pop("_names", {}), **new.pop("_names", {})}
    meva_rows = {i for (i, c) in old if c == "meva"}
    from .daily import is_bukhara
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, in_base, bukhara, source FROM companies")}
    added = defaultdict(float)
    for r in con.execute("""SELECT inn, kind, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND report_date>? AND report_date<=?
                            AND kind IN ('sanoat','meva','meva_x') GROUP BY inn, kind""", (old_T, new_T)):
        if r["kind"] == "sanoat":
            key = (r["inn"], "sanoat") if (r["inn"], "sanoat") in old or (r["inn"], "meva") not in old else (r["inn"], "meva")
            added[key] += r["s"]
        elif r["kind"] == "meva_x" or is_bukhara(comp.get(r["inn"]), r["inn"], meva_rows):
            added[(r["inn"], "meva")] += r["s"]
    days = data_days(con, old_T, new_T)
    all_days = [(dt.date.fromisoformat(old_T) + dt.timedelta(days=i)).isoformat() for i in range(1, (dt.date.fromisoformat(new_T) - dt.date.fromisoformat(old_T)).days + 1)]
    missing = [d for d in all_days if d not in days and dt.date.fromisoformat(d).weekday() < 6]
    mism, removed = [], []
    for key in set(old) | set(new) | set(added):
        if key not in new:
            if old.get(key, 0) or added.get(key, 0): removed.append({"inn": key[0], "name": names.get(key[0]), "cat": key[1], "old": round(old.get(key, 0), 3)})
            continue
        exp = old.get(key, 0.0) + added.get(key, 0.0); act = new[key]
        if abs(act - exp) > 0.01:
            mism.append({"inn": key[0], "name": names.get(key[0]), "cat": key[1], "expected": round(exp, 3), "actual": round(act, 3), "diff": round(act - exp, 3)})
    mism.sort(key=lambda x: -abs(x["diff"]))
    return {"compared": len(new), "matched": len(new) - len(mism), "mismatch": mism[:25], "mismatch_count": len(mism),
            "mismatch_sum": round(sum(x["diff"] for x in mism), 3), "removed": removed, "missing_days": missing,
            "old_total": round(sum(old.values()), 3), "new_total": round(sum(new.values()), 3), "site_added": round(sum(added.values()), 3)}

def build(con, D: str, out_path: Path, template: Path | None = None) -> dict:
    template = Path(template) if template else template_path(con)
    T = db.get_setting(con, "template_through")
    if not T or not template.exists():
        raise ValueError("Excel shabloni yo'q. Sozlamalar → «2. O'z jadvalingiz» ga oxirgi jadvalingizni .xlsx qilib yuklang.")
    if D < T:
        raise ValueError(f"Shablon jadvalingiz {dmy(T)} gacha ma'lumotni o'z ichiga oladi. {dmy(D)} uchun Excel saytdan chiqmaydi — o'sha kungi faylingizdan foydalaning.")
    wb = load(template)
    lay = nomma_layout(wb)
    ws = wb[NOMMA]
    if int(D[5:7]) != lay["cur_month"] or D[:4] != T[:4]:
        raise ValueError(f"{dmy(D)} — {MONTH_UZ[int(D[5:7])-1]} oyi, jadvalingiz esa {MONTH_UZ[lay['cur_month']-1]} oyi uchun tayyorlangan (номманом Q ustuni). "
                         f"Yangi oy: jadvalni {MONTH_UZ[int(D[5:7])-1]} oyiga tayyorlab (Q ustunini ko'chirish, sarlavhalar, reja) Sozlamalar → «2. O'z jadvalingiz» ga yuklang.")
    warnings = []
    base_through = db.get_setting(con, "opening_through_date")
    if base_through and T < base_through:
        warnings.append(f"Shablon ({dmy(T)}) bojxona bazasidan ({dmy(base_through)}) eski — bazadagi kunlarning GTD lari saytda saqlanmagan bo'lishi mumkin. Yangiroq jadvalingizni yuklang.")
    days = data_days(con, T, D); last_day = days[-1] if days else None
    rows = [dict(r) for r in con.execute("""SELECT g.inn, g.kind, g.report_date d, g.district_code dc, sum(g.stat_usd) s, max(g.exporter) exporter
        FROM gtd_rows g WHERE g.voided=0 AND g.report_date>? AND g.report_date<=? GROUP BY g.inn, g.kind, g.report_date, g.district_code""", (T, D))]
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT * FROM companies")}
    dname = {r["code"]: r["name_uz"] for r in con.execute("SELECT code, name_uz FROM districts")}
    # «Бухоро эмас — ҳисобга киритилмасин» (companies.bukhara=0): nothing of theirs goes to Excel either — the same
    # read-time flag the dashboard uses (report.not_bukhara_inns), so switching it in «Korxonalar» acts everywhere
    not_buk = {i for (i,) in con.execute("SELECT inn FROM companies WHERE bukhara=0")}
    from .report import bukhara_windows, in_window
    win = bukhara_windows(con)      # limited Bukhara period: only the days inside it go to Excel
    rows = [r for r in rows if r["kind"] != "not_bukhara" and (r["inn"] not in not_buk or r["kind"] == "bnpz") and (r["kind"] == "bnpz" or in_window(win, r["inn"], r["d"]))]
    from .daily import is_bukhara
    meva_tpl = {i for i, v in lay["companies"].items() if "meva" in v}
    buk_meva = lambda r: r["kind"] == "meva_x" or (r["kind"] == "meva" and is_bukhara(comp.get(r["inn"]), r["inn"], meva_tpl))
    # GTD of the days the template already covers (base_through < d <= T) are normally inside the template's current-month
    # column. A company that is NOT in the template's номманом could not have been written there, so its rows of those
    # days are added to its new row here ("early" rows). They never touch the karantin svod — that part is in the template.
    if base_through and base_through < T:
        early = [dict(r) for r in con.execute("""SELECT g.inn, g.kind, g.report_date d, g.district_code dc, sum(g.stat_usd) s, max(g.exporter) exporter
            FROM gtd_rows g WHERE g.voided=0 AND g.report_date>? AND g.report_date<=? AND g.kind IN ('sanoat','meva','meva_x')
            GROUP BY g.inn, g.kind, g.report_date, g.district_code""", (base_through, T))]
        for r in early:
            if r["inn"] in lay["tail"] or r["inn"] in not_buk or not in_window(win, r["inn"], r["d"]): continue
            own = lay["companies"].get(r["inn"], {})
            if (r["kind"] == "sanoat" and not own) or (r["kind"] in ("meva", "meva_x") and buk_meva(r) and "meva" not in own):
                r["early"] = True; rows.append(r)
        # tail rows (БНПЗ 202080378, memo SDK …): their days in (base_through, T] are not in the template either — the
        # template's F column is the customs base through base_through, the later days come only from the site (O2:
        # 05.09 БНПЗ 162,42 never reached the Excel). They go where their later days go: tail F column and «СВОД 450 млн».
        early_tail = [dict(r) for r in con.execute("""SELECT g.inn, g.kind, g.report_date d, g.district_code dc, sum(g.stat_usd) s, max(g.exporter) exporter
            FROM gtd_rows g WHERE g.voided=0 AND g.report_date>? AND g.report_date<=? AND g.kind IN ('bnpz','memo')
            GROUP BY g.inn, g.kind, g.report_date, g.district_code ORDER BY g.report_date""", (base_through, T))]
        for r in early_tail:
            if r["inn"] not in lay["tail"]: continue
            if r["kind"] != "bnpz" and (r["inn"] in not_buk or not in_window(win, r["inn"], r["d"])): continue
            r["early"] = True; rows.append(r)
            nm = (comp.get(r["inn"]) or {}).get("name") or r["exporter"] or r["inn"]
            nm = "БНПЗ" if r["kind"] == "bnpz" else nm
            amt = f"{r['s']:,.2f}".replace(",", " ").replace(".", ",")
            warnings.append(f"{nm} {dmy(r['d'])}: {amt} — шаблонда йўқ эди, қўшилди")
    # ---- 0. «Баҳсли корхоналар» (dispute.py): template Jan..base-month vs customs base, decided per company profile
    from . import dispute
    from .report import window_base_months
    year_i = int(D[:4]); bm = dispute.base_month(con) if base_through and base_through[:4] == D[:4] else 0
    disp = dispute.disputes(con, wb, lay) if bm else []
    eff = dispute.effective_months(con, year_i) if bm else {}
    def site_months(inn):      # the site's opening of one company (opening_balances; limited Bukhara period re-summed by date)
        return window_base_months(con, inn, win[inn], year_i) if inn in win else eff.get(inn, [0.0] * 12)
    # (d) in the base, not in the template, «Бухоро»: a new row with the base months (even without GTD after T)
    base_add = {r["inn"]: site_months(r["inn"]) for r in disp if r["decision"] == "add_row" and r["registered"] and r["inn"] not in lay["tail"]}
    base_add = {i: m for i, m in base_add.items() if any(round(v, 3) for v in m[:bm])}
    # ---- 1. new номманом rows at the end of their district block (before anything is written)
    #   sanoat: a new exporter; meva: a Bukhara exporter's meva-sabzavot (general svod) without a «Мева-сабзавотлар» row yet
    new_inns, new_meva = {i: sum(m[:bm]) for i, m in base_add.items()}, {}
    for r in rows:
        if r["inn"] in lay["tail"]: continue
        if r["kind"] == "sanoat" and r["inn"] not in lay["companies"]:
            new_inns[r["inn"]] = new_inns.get(r["inn"], 0.0) + r["s"]
        elif r["kind"] in ("meva", "meva_x") and buk_meva(r) and "meva" not in lay["companies"].get(r["inn"], {}):
            new_meva[r["inn"]] = new_meva.get(r["inn"], 0.0) + r["s"]
    added, no_district = [], []
    bykey = {b["key"]: b for b in lay["blocks"]}
    for cat, pool in (("sanoat", new_inns), ("meva", new_meva)):
        for inn in sorted(pool, key=lambda i: (comp.get(i, {}).get("district_code") or "", -pool[i])):
            c = comp.get(inn) or {}
            blk = bykey.get(_norm(dname.get(c.get("district_code"), "")))
            if not blk: no_district.append(f"{c.get('name') or inn} ({inn})"); continue
            added.append((inn, blk["key"], blk["name"], cat))
    if no_district:
        raise ValueError("Yangi korxonaning tumani aniqlanmagan — Korxonalar menyusida tumanini belgilang: " + "; ".join(no_district))
    meva_tar = next((ws.cell(x["row"], lay["cols"]["tarmoq"]).value for b in lay["blocks"] for x in b["rows"] if x["cat"] == "meva"), None) or "Мева-сабзавот маҳсулотлари"
    for inn, key, _, cat in added:
        lay = nomma_layout(wb)          # positions move after every insertion
        blk = next(b for b in lay["blocks"] if b["key"] == key)
        proto = blk["rows"][-1]["row"] if blk["rows"] else next(x["row"] for b in lay["blocks"] for x in b["rows"])
        at = blk["last"] + 1
        insert_row(wb, ws, at)
        c = comp.get(inn) or {}
        if cat == "meva":
            pr = con.execute("""SELECT product FROM gtd_rows WHERE voided=0 AND inn=? AND kind IN ('meva','meva_x') AND report_date>? AND report_date<=?
                                ORDER BY stat_usd DESC LIMIT 1""", (inn, T, D)).fetchone()
            c = {**c, "tarmoq": meva_tar, "product": pr[0] if pr else c.get("product")}
        _fill_new_row(ws, lay, blk, proto, at, inn, c, con, category=cat)
    lay = nomma_layout(wb)
    cols = lay["cols"]
    # ---- 1b. disputed companies: (b) rewrite Jan..base-month cells to the base, (c) clear «Бухоро эмас», (d) fill the new base rows
    disp_info = {"rewritten": [], "cleared": [], "added_base": []}
    if bm:
        from openpyxl.comments import Comment
        mcol = {c.column: _month_of(c.value) for c in ws[lay["header_row"]] if c.column in cols["months"]}
        base_cols = {col: mo for col, mo in mcol.items() if mo and mo <= bm}
        money = lambda v: f"{v:,.2f}".replace(",", " ").replace(".", ",")
        def put_months(row, months):
            for col, mo in base_cols.items():
                v = round(months[mo - 1], 3)
                ws.cell(row, col).value = v if v else None
        for inn, months in base_add.items():
            row = lay["companies"].get(inn, {}).get("sanoat")
            if not row: continue
            put_months(row, months); nm = (comp.get(inn) or {}).get("name") or inn
            disp_info["added_base"].append({"inn": inn, "name": nm, "sum": round(sum(months[:bm]), 3)})
            warnings.append(f"{nm}: шаблонда йўқ эди — божхона базаси бўйича янги қатор ({money(sum(months[:bm]))}, янв–{MONTH_STEMS[bm-1]})")
        for r in disp:
            own = lay["companies"].get(r["inn"], {})
            if r["decision"] == "customs" and r["in_template"] and own.get("sanoat"):
                months = site_months(r["inn"]); put_months(own["sanoat"], months)
                disp_info["rewritten"].append({"inn": r["inn"], "name": r["name"], "template": r["template_total"], "base": round(sum(months[:bm]), 3)})
                warnings.append(f"{r['name']}: шаблон {money(r['template_total'])} → божхона базаси {money(sum(months[:bm]))} (янв–{MONTH_STEMS[bm-1]} қайта ёзилди)")
        # (c) «Бухоро эмас»: every номманом row of the company loses its month cells (E..P and the current month) — nowhere on the site either
        for inn, own in lay["companies"].items():
            if inn not in not_buk: continue
            tot = 0.0
            for cat, row in own.items():
                for col in cols["months"]:
                    tot += _cell_number(ws.cell(row, col).value) or 0.0
                    ws.cell(row, col).value = None
                ws.cell(row, cols["name"]).comment = Comment("Бухоро эмас — сайт бўйича", SITE_AUTHOR)
            nm = (comp.get(inn) or {}).get("name") or inn
            disp_info["cleared"].append({"inn": inn, "name": nm, "sum": round(tot, 3)})
            warnings.append(f"{nm}: «Бухоро эмас» — ой катаклари тозаланди ({money(tot)})")
    # ---- 2. amounts per номманом row
    per_row = defaultdict(lambda: defaultdict(float)); tail_rows = defaultdict(lambda: defaultdict(float))
    meva = defaultdict(lambda: defaultdict(float)); bnpz = defaultdict(float); unrouted = []
    for r in rows:
        inn, kind, d, s = r["inn"], r["kind"], r["d"], r["s"]
        rw = lay["companies"].get(inn, {})
        if kind == "bnpz":
            bnpz[d] += s
            if inn in lay["tail"]: tail_rows[lay["tail"][inn]][d] += s
        elif kind == "memo":
            if inn in lay["tail"]: tail_rows[lay["tail"][inn]][d] += s
            elif rw: per_row[rw.get("sanoat") or rw.get("meva")][d] += s
            else: unrouted.append((inn, kind, s))
        elif kind == "sanoat":
            target = rw.get("sanoat") or rw.get("meva")
            if target: per_row[target][d] += s
            elif inn in lay["tail"]: tail_rows[lay["tail"][inn]][d] += s
            else: unrouted.append((inn, kind, s))
        elif kind in ("meva", "meva_x"):
            if kind == "meva" and not r.get("early"): meva[r["dc"]][d] += s          # karantin svod: grown in Bukhara
            if buk_meva(r):                                    # general svod: Bukhara exporter's meva-sabzavot row
                if "meva" in rw: per_row[rw["meva"]][d] += s
                else: unrouted.append((inn, kind, s))
    if unrouted:
        warnings.append("номманом ga yozilmadi: " + "; ".join(f"{i} {k} {s:.2f}" for i, k, s in unrouted))
    # ---- 3. номманом: current month "+X", 1 кунда
    c_month, c_day, c_total = cols["month"], cols["day"], cols["total"]
    written = 0.0
    for r, per in per_row.items():
        amts = [per[d] for d in sorted(per)]
        if append_terms(ws.cell(r, c_month), amts): written += sum(amts)
    for r, per in tail_rows.items():
        append_terms(ws.cell(r, c_total), [per[d] for d in sorted(per)])
    if last_day:
        for b in lay["blocks"]:
            for x in b["rows"]:
                cell = ws.cell(x["row"], c_day)
                if cell.value is not None and not _is_sum_formula(cell.value): cell.value = None
        for r, per in per_row.items():
            if round(per.get(last_day, 0), 3): ws.cell(r, c_day).value = round(per[last_day], 3)
    ws.cell(lay["label_row"], 1).value = f"{dt.date.fromisoformat(D) + dt.timedelta(days=1):%d.%m.%Y} йил"
    # ---- 4. karantin svod: meva-sabzavot
    kname = find_sheet(wb.sheetnames, KARANTIN_SVOD); kws = wb[kname]; kl = svod_layout(kws)
    mrows = {}
    for r in range(kl["sub_row"] + 1, kws.max_row + 1):
        b = _norm(kws.cell(r, 2).value); cn = _norm(kws.cell(r, 3).value)
        if b.startswith("мева") and cn: mrows[cn] = r
    meva_total = 0.0
    for dc, per in meva.items():
        r = mrows.get(_norm(dname.get(dc, "")))
        if not r:
            warnings.append(f"Karantin svodda tuman topilmadi: {dname.get(dc, dc)} — {sum(per.values()):.2f} ming $"); continue
        amts = [per[d] for d in sorted(per)]
        append_terms(kws.cell(r, kl["period_fact"]), amts)
        append_terms(kws.cell(r, kl["month_fact"]), amts)
        meva_total += sum(amts)
    if last_day and kl["day_fact"]:
        for r in mrows.values():
            cell = kws.cell(r, kl["day_fact"])
            if cell.value is not None and not _is_sum_formula(cell.value): cell.value = None
        for dc, per in meva.items():
            r = mrows.get(_norm(dname.get(dc, "")))
            if r and round(per.get(last_day, 0), 3): kws.cell(r, kl["day_fact"]).value = round(per[last_day], 3)
    msheet = karantin_month_sheet(wb, lay["cur_month"])
    n_month_rows = _append_karantin_rows(wb[msheet], con, T, D) if msheet else 0
    if not msheet and meva:
        warnings.append(f"«{MONTH_UZ[lay['cur_month']-1].capitalize()} карантин» varag'i yo'q — GTD qatorlari qo'yilmadi (svoddagi summa yozildi)")
    # ---- 5. БНПЗ -> «СВОД … млн», Қоровулбозор sanoat YTD
    bn_info = None
    if bnpz:
        gname = find_sheet(wb.sheetnames, GOVERNOR_SVOD)
        if gname:
            gws = wb[gname]; gl = svod_layout(gws)
            row = next((r for r in range(gl["sub_row"] + 1, gws.max_row + 1) if _norm(gws.cell(r, 2).value).startswith("саноат") and _norm(gws.cell(r, 3).value) == BNPZ_DISTRICT), None)
            if row:
                append_terms(gws.cell(row, gl["period_fact"]), [bnpz[d] for d in sorted(bnpz)])
                bn_info = {"sheet": gname, "cell": f"{L(gl['period_fact'])}{row}", "sum": round(sum(bnpz.values()), 3)}
        if not bn_info: warnings.append("«СВОД … млн» da Қоровулбозор sanoat qatori topilmadi — БНПЗ qo'shilmadi")
    # ---- 6. prev-year day term
    before = _doy(D[:8] + "01") - 1; day = int(D[8:10]); n_prev = 0
    for title, coord in prev_year_cells(wb, T)[0]:
        c = wb[title][coord]
        def rep(m):
            terms = [int(x) for x in m.group(2).split("+")]
            acc = 0
            for k in range(len(terms) + 1):
                if acc == before: return f"/{m.group(1)}*({'+'.join(str(t) for t in terms[:k] + [day])})"
                if k < len(terms): acc += terms[k]
            return m.group(0)
        nv = _PREV.sub(rep, c.value, count=1)
        if nv != c.value: c.value = nv; n_prev += 1
    wb.calculation.fullCalcOnLoad = True
    tmp = Path(str(out_path) + ".tmp")
    wb.save(tmp)
    try:
        tmp.replace(out_path)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise ValueError(f"«{Path(out_path).name}» fayli Excelda ochiq — uni yoping va qayta bosing.")
    return {"file": str(out_path), "template_through": T, "days": days, "last_day": last_day, "sanoat_written": round(written, 3),
            "new_rows": [{"inn": i, "name": (comp.get(i) or {}).get("name"), "district": nm, "category": cat} for i, _, nm, cat in added],
            "meva": round(meva_total, 3), "karantin_month_rows": n_month_rows, "bnpz": bn_info, "prev_year_cells": n_prev, "warnings": warnings,
            "disputes": disp_info}

def _fill_new_row(ws, lay, blk, proto: int, at: int, inn: str, c: dict, con, category: str = "sanoat"):
    cols = lay["cols"]
    maxc = ws.max_column
    skip_formula = set(cols["months"]) | {cols["day"], cols["inn"], cols["name"], cols["hudud"], cols["tarmoq"], cols["product"], cols["district"], cols["category"]}
    for ci in range(1, maxc + 1):
        src = ws.cell(proto, ci); dst = ws.cell(at, ci)
        if src.has_style: dst._style = copy.copy(src._style)
        v = src.value
        if isinstance(v, str) and v.startswith("=") and ci not in skip_formula:
            dst.value = copy_down(v, ws.title, at - proto)
    if ws.row_dimensions[proto].height: ws.row_dimensions[at].height = ws.row_dimensions[proto].height
    same_block = blk["rows"] and proto == blk["rows"][-1]["row"]
    ws.cell(at, cols["inn"]).value = int(inn) if inn.isdigit() and len(inn) <= 15 else inn
    ws.cell(at, cols["name"]).value = clean(c.get("name") or inn)
    ws.cell(at, cols["hudud"]).value = ws.cell(proto, cols["hudud"]).value if same_block else clean(c.get("hudud"))
    ws.cell(at, cols["tarmoq"]).value = clean(c.get("tarmoq"))
    ws.cell(at, cols["product"]).value = clean(c.get("product"))
    if cols["district"]: ws.cell(at, cols["district"]).value = ws.cell(proto, cols["district"]).value if same_block else blk["name"]
    if cols["category"]: ws.cell(at, cols["category"]).value = lay["cat_text"].get(category, "Мева-сабзавотлар" if category == "meva" else "Саноат маҳсулотлари ")
    if not isinstance(ws.cell(at, 1).value, str):
        ws.cell(at, 1).value = f"=+A{at-1}+1"
    for ci in (lay["yellow_cols"] or range(1, cols["tarmoq"] + 1)):
        ws.cell(at, ci).fill = PatternFill("solid", fgColor=YELLOW)
    # block total now ends at the new row
    for cell in ws[blk["header"]]:
        v = cell.value
        if isinstance(v, str) and v.startswith("="):
            nv = re.sub(r"(\$?[A-Z]{1,3}\$?)" + str(blk["first"]) + r":(\$?[A-Z]{1,3}\$?)" + str(at - 1) + r"(?!\d)",
                        lambda m: f"{m.group(1)}{blk['first']}:{m.group(2)}{at}", v)
            if nv != v: cell.value = nv
    # numbering of the next block continues from the new row
    rx = re.compile(r"^=\+?A\$?" + str(at - 1) + r"\+1$")
    for r in range(at + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and rx.match(v):
            ws.cell(r, 1).value = f"=+A{at}+1"; break

def _append_karantin_rows(ws, con, T: str, D: str) -> int:
    hdr = {_norm(str(c.value).replace(",", ".")): c.column for c in ws[1] if c.value}
    c_gtd = next((col for k, col in hdr.items() if k.startswith("код таможенного поста")), None)
    c_item = next((col for k, col in hdr.items() if k.startswith("номер товара")), None)
    last = max((r for r in range(2, ws.max_row + 1) if any(ws.cell(r, c).value not in (None, "") for c in hdr.values())), default=1)
    have = set()
    if c_gtd and c_item:
        for r in range(2, last + 1):
            have.add((str(ws.cell(r, c_gtd).value or "").strip(), inn_str(ws.cell(r, c_item).value)))
    rows = con.execute("SELECT raw, gtd_no, item_no FROM gtd_rows WHERE voided=0 AND kind='meva' AND report_date>? AND report_date<=? ORDER BY report_date, id", (T, D)).fetchall()
    n = 0
    for rec in rows:
        if (str(rec["gtd_no"] or "").strip(), inn_str(rec["item_no"])) in have: continue
        raw = {_norm(str(k).replace(",", ".")): v for k, v in json.loads(rec["raw"]).items()}
        r = last + 1 + n
        for k, col in hdr.items():
            proto = ws.cell(last, col) if last > 1 else None
            if proto is not None and proto.has_style: ws.cell(r, col)._style = copy.copy(proto._style)
            if proto is not None and isinstance(proto.value, str) and proto.value.startswith("=") and k not in raw:
                ws.cell(r, col).value = copy_down(proto.value, ws.title, r - last); continue
            if k in raw:
                v = raw[k]
                if isinstance(v, str) and re.fullmatch(r"-?\d+(\.\d+)?", v.strip()) and any(t in k for t in ("количество", "нетто", "стоимость", "курс", "инн", "окпо")):
                    v = float(v) if "." in v else int(v)
                ws.cell(r, col).value = clean(v)
        n += 1
    return n
