"""Тугаган йиллар натижалари (N ойлик) — Sozlamalar.

ФАҚАТ ТУГАГАН ЙИЛЛАР учун. Иш тартиби оддий: йил ва давр (N ойлик) танланади → Excel шаблон
юклаб олинади → тўлдирилади → қайта юкланади. Ҳар катак — шу йилнинг январь–N ой якуни
(амалда, минг $), туман ва тур (саноат / мева-сабзавот) кесимида.

Жорий йил бу ерда йўқ: жорий йил РЕЖАСИ «Режа» панелида алоҳида киритилади, жорий йил
АМАЛДАГИ кўрсаткичи эса кунлик божхона маълумотидан ўзи шаклланади.

Қоида (Жафар, 21–22.09.2026): «ўтган йил шу даврга» солиштириш керак бўлган ҳар қандай жадвал
аввало шу натижаларга қарайди — керакли ой (N ойлик) бўйича. Ой охирида аниқ рақам олинади;
оралиқ кун учун икки якорь орасида кунлар бўйича чизиқли интерполяция; якорь бўлмаса — эски
«Ўтган йил базаси» (prognoz.prev_year_base, кунга мутаносиб) ишлайди (захира).

Вилоят қатори — туманлар йиғиндиси (алоҳида киритилмайди). kind: sanoat | meva; 'all' = иккиси йиғиндиси.

29.09.2026 (Жафар): «Амалда» натижалари НОМЛИ ТЎПЛАМЛАР (fact_sets) бўлиб сақланади — бир йил учун бир нечта
вариант (масалан «Вазирлик», «Ҳокимият»): ким сўрашига қараб турли рақам билан ишланади. Ҳар йилнинг битта тўплами
«асосий» (main=1) — сайт ҳамма жойда стандарт шундан олади; Dashboard/Свод ва Excel модалида бошқасини танлаш мумкин
(`set_id`). Давр (N ойлик) энди тўплам ичида ягона: UNIQUE(set_id, months).
"""
from __future__ import annotations
import datetime as dt, calendar, re, threading
from collections import defaultdict
from . import db, report

KINDS = ("sanoat", "meva")
KIND_UZ = {"sanoat": "Саноат маҳсулотлари", "meva": "Мева-сабзавотлар"}
MONTHS_GEN = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
FIELDS = ("fact",)          # фақат «амалда» — прогноз ва «ўтган йил» устунлари олиб ташланди (22.09.2026)

SCHEMA = """
CREATE TABLE IF NOT EXISTS fact_sets (                        -- «Амалда» тўплами: йил + номи (29.09.2026)
  id INTEGER PRIMARY KEY AUTOINCREMENT, year INTEGER NOT NULL, title TEXT NOT NULL, note TEXT,
  main INTEGER NOT NULL DEFAULT 0,                             -- 1 = шу йилнинг асосий тўплами (стандарт)
  created_at TEXT DEFAULT (datetime('now','localtime')), updated_at TEXT
);
CREATE TABLE IF NOT EXISTS period_results (
  id INTEGER PRIMARY KEY AUTOINCREMENT, set_id INTEGER, year INTEGER NOT NULL, months INTEGER NOT NULL,   -- 8 = январь–август
  title TEXT, note TEXT, source TEXT, created_at TEXT DEFAULT (datetime('now','localtime')), updated_at TEXT,
  UNIQUE(set_id, months)
);
CREATE TABLE IF NOT EXISTS period_values (
  period_id INTEGER NOT NULL REFERENCES period_results(id) ON DELETE CASCADE,
  district_code TEXT NOT NULL, kind TEXT NOT NULL,          -- sanoat | meva
  fact REAL,
  PRIMARY KEY (period_id, district_code, kind)
);
"""

_MIG_LOCK = threading.Lock()
MAIN_TITLE = "Асосий"

def ensure(con):
    con.executescript(SCHEMA)
    _migrate_past_only(con)
    _migrate_sets(con)
    _fix_orphans(con)

def _migrate_sets(con):
    """Бир марталик (29.09.2026): period_results га set_id; эски UNIQUE(year, months) → UNIQUE(set_id, months).
    Мавжуд ҳар йил учун «Асосий» тўплам (main=1) яратилиб, ўша йил даврлари унга бириктирилади — рақамлар ўзгармайди.
    SQLite да UNIQUE ни ўзгартириб бўлмайди — жадвал қайта қурилади (foreign_keys OFF, id лар сақланади)."""
    if "set_id" in {r[1] for r in con.execute("PRAGMA table_info(period_results)")}: return
    with _MIG_LOCK:
        if "set_id" in {r[1] for r in con.execute("PRAGMA table_info(period_results)")}: return
        try: db.backup_if_stale("fact_sets", minutes=0)
        except Exception: pass
        con.commit()
        con.execute("PRAGMA foreign_keys=OFF")
        try:
            con.executescript(f"""
            BEGIN;
            CREATE TABLE period_results_v2 (
              id INTEGER PRIMARY KEY AUTOINCREMENT, set_id INTEGER, year INTEGER NOT NULL, months INTEGER NOT NULL,
              title TEXT, note TEXT, source TEXT, created_at TEXT DEFAULT (datetime('now','localtime')), updated_at TEXT,
              UNIQUE(set_id, months));
            INSERT INTO fact_sets(year, title, note, main)
              SELECT DISTINCT year, '{MAIN_TITLE}', 'аввалги «Тугаган йиллар натижалари»', 1 FROM period_results
              WHERE year NOT IN (SELECT year FROM fact_sets WHERE main=1);
            INSERT INTO period_results_v2(id, set_id, year, months, title, note, source, created_at, updated_at)
              SELECT p.id, (SELECT s.id FROM fact_sets s WHERE s.year=p.year AND s.main=1 ORDER BY s.id LIMIT 1),
                     p.year, p.months, p.title, p.note, p.source, p.created_at, p.updated_at FROM period_results p;
            DROP TABLE period_results;
            ALTER TABLE period_results_v2 RENAME TO period_results;
            INSERT INTO audit_log(action, details) VALUES('fact_sets_migrated', (SELECT '{{"sets": ' || count(*) || '}}' FROM fact_sets));
            COMMIT;""")
        except Exception:
            try: con.execute("ROLLBACK")
            except Exception: pass
            raise
        finally:
            con.execute("PRAGMA foreign_keys=ON")

def _fix_orphans(con):
    """set_id сиз давр (эски код ёки миграция ёзган) — ўша йилнинг асосий тўпламига."""
    ys = [r[0] for r in con.execute("SELECT DISTINCT year FROM period_results WHERE set_id IS NULL")]
    if not ys: return
    with con:
        for y in ys:
            sid = main_set(con, y, create=True)
            for r in con.execute("SELECT id, months FROM period_results WHERE set_id IS NULL AND year=?", (y,)).fetchall():
                if con.execute("SELECT 1 FROM period_results WHERE set_id=? AND months=?", (sid, r["months"])).fetchone():
                    con.execute("DELETE FROM period_values WHERE period_id=?", (r["id"],)); con.execute("DELETE FROM period_results WHERE id=?", (r["id"],))
                else:
                    con.execute("UPDATE period_results SET set_id=? WHERE id=?", (sid, r["id"]))

def _migrate_past_only(con):
    """Бир марталик: эски тузилмада (plan / fact / prev_fact) жорий ва келгуси йил даврлари бор эди.
    Уларнинг «ўтган йил шу давр амалда» (prev_fact) рақамлари ўша ўтган йилнинг ўз даврига (fact)
    кўчирилади, сўнг жорий/келгуси йил даврлари ўчирилади — энди фақат тугаган йиллар сақланади."""
    if db.get_setting(con, "periods_past_only") == "1": return
    cols = {r[1] for r in con.execute("PRAGMA table_info(period_values)")}
    cy = _cur_year(con)
    rows = con.execute("SELECT id, year, months FROM period_results WHERE year>=?", (cy,)).fetchall()
    if rows:
        try: db.backup_if_stale("periods_past_only", minutes=0)
        except Exception: pass
    moved = 0
    with con:
        for r in rows:
            py, m = r["year"] - 1, r["months"]
            if "prev_fact" in cols:
                vals = con.execute("SELECT district_code, kind, prev_fact FROM period_values WHERE period_id=? AND prev_fact IS NOT NULL", (r["id"],)).fetchall()
                if vals:
                    t = con.execute("SELECT id FROM period_results WHERE year=? AND months=?", (py, m)).fetchone()
                    pid = t["id"] if t else con.execute("INSERT INTO period_results(year,months,title,source) VALUES(?,?,?,?)",
                                                        (py, m, label(py, m), "migrate")).lastrowid
                    for v in vals:
                        con.execute("INSERT OR IGNORE INTO period_values(period_id,district_code,kind,fact) VALUES(?,?,?,NULL)", (pid, v["district_code"], v["kind"]))
                        con.execute("UPDATE period_values SET fact=COALESCE(fact,?) WHERE period_id=? AND district_code=? AND kind=?",
                                    (v["prev_fact"], pid, v["district_code"], v["kind"]))
                        moved += 1
            con.execute("DELETE FROM period_values WHERE period_id=?", (r["id"],))
            con.execute("DELETE FROM period_results WHERE id=?", (r["id"],))
        if rows: db.log(con, "periods_past_only", removed=len(rows), moved=moved)
        db.set_setting(con, "periods_past_only", "1")

def _cur_year(con) -> int:
    try: return int(db.get_setting(con, "year") or dt.date.today().year)
    except (TypeError, ValueError): return dt.date.today().year

def max_year(con) -> int:
    """Киритиш мумкин бўлган энг катта йил — тугаган охирги йил."""
    return _cur_year(con) - 1

def label(year: int, months: int) -> str:
    m = int(months)
    if m >= 12: return f"{year} йил якуни (12 ойлик)"
    return f"{year} йил {m} ойлик (январь–{MONTHS_GEN[m - 1]})" if m > 1 else f"{year} йил 1 ойлик (январь)"

def _doy_end(year: int, months: int) -> int:
    """N-ой охирги кунининг йил бошидан кун рақами."""
    return dt.date(year, months, calendar.monthrange(year, months)[1]).timetuple().tm_yday

# ------------------------------------------------------------------ o'qish
def _vals(con, pid: int | None) -> dict:
    out = defaultdict(lambda: {k: {f: None for f in FIELDS} for k in KINDS})
    if not pid: return out
    for r in con.execute("SELECT district_code, kind, fact FROM period_values WHERE period_id=?", (pid,)):
        if r["kind"] in KINDS: out[r["district_code"]][r["kind"]] = {"fact": r["fact"]}
    return out

def _totals(vals: dict, codes: list[str]) -> dict:
    t = {k: {f: 0.0 for f in FIELDS} for k in (*KINDS, "all")}
    for c in codes:
        for k in KINDS:
            for f in FIELDS:
                v = vals[c][k][f] or 0.0; t[k][f] += v; t["all"][f] += v
    return {k: {f: round(v, 3) for f, v in d.items()} for k, d in t.items()}

def main_set(con, year: int, create: bool = False):
    """Йилнинг асосий тўплами id (йўқ бўлса None ёки create=True да яратилади)."""
    r = con.execute("SELECT id FROM fact_sets WHERE year=? AND main=1 ORDER BY id LIMIT 1", (int(year),)).fetchone()
    if r: return r["id"]
    r = con.execute("SELECT id FROM fact_sets WHERE year=? ORDER BY id LIMIT 1", (int(year),)).fetchone()   # асосий белгиси йўқ — биринчиси
    if r:
        if create: con.execute("UPDATE fact_sets SET main=1 WHERE id=?", (r["id"],))
        return r["id"]
    if not create: return None
    return con.execute("INSERT INTO fact_sets(year, title, main) VALUES(?,?,1)", (int(year), MAIN_TITLE)).lastrowid

def get_set(con, set_id) -> dict | None:
    if not set_id: return None
    r = con.execute("SELECT * FROM fact_sets WHERE id=?", (int(set_id),)).fetchone()
    return dict(r) if r else None

def _sid(con, year: int, set_id=None):
    """set_id берилса — ўша (йўқ бўлса хато), берилмаса — йилнинг асосий тўплами (бўлмаса None)."""
    if set_id not in (None, "", 0, "0"):
        st = get_set(con, set_id)
        if not st: raise ValueError("«Амалда» тўплами топилмади (ўчирилган бўлиши мумкин) — саҳифани янгиланг")
        return st["id"]
    return main_set(con, year)

def find(con, year: int, months: int, set_id=None):
    sid = _sid(con, year, set_id)
    if not sid: return None
    r = con.execute("SELECT id FROM period_results WHERE set_id=? AND months=?", (sid, int(months))).fetchone()
    return r["id"] if r else None

def list_periods(con, set_id=None, main_only: bool = False) -> list[dict]:
    """Даврлар (ҳар бири тўплами билан). set_id — фақат шу тўплам; main_only — фақат асосий тўпламлар (аввалги хулқ)."""
    ensure(con)
    codes = [d["code"] for d in report.districts(con)]
    out = []
    w, a = ("WHERE p.set_id=?", (int(set_id),)) if set_id else (("WHERE s.main=1", ()) if main_only else ("", ()))
    for r in con.execute(f"""SELECT p.*, s.title set_title, s.main set_main FROM period_results p LEFT JOIN fact_sets s ON s.id=p.set_id
                             {w} ORDER BY p.year DESC, s.main DESC, s.id, p.months DESC""", a):
        d = dict(r); vals = _vals(con, d["id"]); tot = _totals(vals, codes)
        d["label"] = label(d["year"], d["months"]); d["totals"] = tot["all"]
        d["by_kind"] = {k: tot[k]["fact"] for k in KINDS}
        d["filled"] = sum(1 for c in codes for k in KINDS if vals[c][k]["fact"] is not None)
        out.append(d)
    return out

def grid(con, year: int, months: int, set_id=None) -> dict:
    """Йил/ой бўйича тўр (тўплам ичида). Давр ҳали яратилмаган бўлса — бўш тўр қайтарилади."""
    ensure(con)
    st = get_set(con, set_id) if set_id else None
    year, months = (int(st["year"]) if st else int(year)), int(months)
    pid = find(con, year, months, set_id)
    p = dict(con.execute("SELECT * FROM period_results WHERE id=?", (pid,)).fetchone()) if pid else {"id": None, "year": year, "months": months, "source": None, "updated_at": None}
    p["set"] = st or get_set(con, main_set(con, year))
    vals = _vals(con, pid)
    ds = report.districts(con); codes = [d["code"] for d in ds]
    rows = [{"district_code": d["code"], "district": d["name_uz"], "kind": k,
             "fact": (round(vals[d["code"]][k]["fact"], 3) if vals[d["code"]][k]["fact"] is not None else None)}
            for d in ds for k in KINDS]
    p["label"] = label(year, months)
    return {"period": p, "rows": rows, "totals": _totals(vals, codes), "fields": list(FIELDS), "max_year": max_year(con)}

# ------------------------------------------------------------------ hisobotlar uchun qidiruv
def _sum_kind(v: dict, kind: str, field: str = "fact"):
    """kind: sanoat | meva | all. None — қатор киритилмаган."""
    if kind == "all":
        a, b = v["sanoat"][field], v["meva"][field]
        return None if a is None and b is None else (a or 0.0) + (b or 0.0)
    return v[kind][field]

def _value(con, pid: int, code: str, kind: str, field: str = "fact"):
    vals = _vals(con, pid)
    if code == report.REGION:
        codes = [d["code"] for d in report.districts(con)]
        xs = [_sum_kind(vals[c], kind, field) for c in codes]
        return None if all(x is None for x in xs) else sum(x or 0.0 for x in xs)
    return _sum_kind(vals[code], kind, field) if code in vals else None

def anchors(con, year: int, code: str, kind: str, set_id=None) -> list[tuple[int, float, str]]:
    """Ўтган йил (year-1) жамғарма якорлари: [(кун, қиймат, изоҳ)] — ўша йил асосий тўплами (ёки танланган set_id) даврлари."""
    ensure(con)
    st = get_set(con, set_id) if set_id else None
    py = int(st["year"]) if st else year - 1; pts = {}
    sid = st["id"] if st else main_set(con, py)
    if not sid: return []
    for r in con.execute("SELECT id, months FROM period_results WHERE set_id=? ORDER BY months", (sid,)):
        v = _value(con, r["id"], code, kind)
        if v is not None: pts[r["months"]] = (v, label(py, r["months"]))
    return [(_doy_end(py, m), v, lbl) for m, (v, lbl) in sorted(pts.items())]

def has_any(con, year: int) -> bool:
    ensure(con)
    return bool(con.execute("SELECT 1 FROM period_results WHERE year=? LIMIT 1", (year - 1,)).fetchone())

def interpolate(points: list[tuple[int, float]], day: int) -> tuple[float, str]:
    """(0,0) дан бошланадиган кесма-чизиқли жамғарма: якорь орасида кунлар бўйича, охиргисидан кейин ўртача кунлик суръат."""
    pts = sorted({0: 0.0, **{d: v for d, v in points}}.items())
    if day <= 0 or len(pts) == 1: return 0.0, "base"
    for (d0, v0), (d1, v1) in zip(pts, pts[1:]):
        if day == d1: return v1, "exact"
        if d0 < day < d1: return v0 + (v1 - v0) * (day - d0) / (d1 - d0), "interp"
    dl, vl = pts[-1]
    return vl / dl * day, "extrap"

def with_base(pts: list[tuple[int, float]], base: tuple[int, float] | None) -> list[tuple[int, float]]:
    """Эски «Ўтган йил базаси» якорини қўшади — фақат шу кунда расмий N ойлик натижа бўлмаса.
    25.09.2026: база (273 кун) 9 ойлик натижани (ҳам 273-кун) устидан ёзиб юборарди."""
    if not base: return list(pts)
    if any(d == base[0] for d, _ in pts): return list(pts)
    return list(pts) + [base]

def prev_cum(con, code: str, kind: str, D: str, base: tuple[float, int] | None = None, set_id=None) -> tuple[float, str] | None:
    """Ўтган йил D кунигача жамғарма (тугаган йил натижалари + эски база якори). Якорь бўлмаса None.
    set_id — танланган «амалда» тўплами: фақат шу тўплам якорлари (эски база аралаштирилмайди)."""
    y = int(D[:4]); pts = [(d, v) for d, v, _ in anchors(con, y, code, kind, set_id)]
    if not pts: return None
    if not set_id: pts = with_base(pts, (int(base[1] or 273), float(base[0])) if base and base[0] else None)
    day = report.prev_year_days(con, D)
    return interpolate(pts, day)

def prev_info(con, D: str, set_id=None) -> dict:
    """UI/Excel учун манба изоҳи: {mode, text, set}. mode: none | exact | interp | extrap."""
    y = int(D[:4]); pts = anchors(con, y, report.REGION, "all", set_id)
    st = get_set(con, set_id) if set_id else get_set(con, main_set(con, y - 1))
    if not pts: return {"mode": "none", "text": "", "set": st}
    day = report.prev_year_days(con, D)
    base = None if set_id else con.execute("SELECT prev_year_base, prev_year_base_days FROM prognoz WHERE year=? AND district_code=? AND kind='all'", (y, report.REGION)).fetchone()
    allp = with_base([(d, v) for d, v, _ in pts], (int(base["prev_year_base_days"] or 273), float(base["prev_year_base"])) if base and base["prev_year_base"] else None)
    base_used = any(d == int(base["prev_year_base_days"] or 273) for d, _ in allp) and not any(d == int(base["prev_year_base_days"] or 273) for d, _, _ in pts) if base and base["prev_year_base"] else False
    _, mode = interpolate(allp, day)
    names = {d: l for d, _, l in pts}
    if base_used:
        names.setdefault(int(base["prev_year_base_days"] or 273), db.get_setting(con, "prev_year_base_label") or f"ўтган йил базаси ({base['prev_year_base_days'] or 273} кун)")
    srt = sorted(allp)
    if mode == "exact":
        text = f"Ўтган йил — Sozlamalar → Тугаган йиллар натижалари: {names.get(day, '')} (аниқ рақам)"
    elif mode == "interp":
        # 0-кун (йил боши) якори srt да йўқ — биринчи якордан олдинги кун учун lo = 0
        lo = max((d for d, _ in srt if d < day), default=0); hi = min((d for d, _ in srt if d > day), default=day)
        if lo == 0:
            text = f"Ўтган йил — Тугаган йиллар натижалари: йил бошидан «{names.get(hi, '')}» гача {day} кунга интерполяция"
        else:
            text = f"Ўтган йил — Тугаган йиллар натижалари: «{names.get(lo, '')}» ва «{names.get(hi, '')}» орасида {day - lo} кунга интерполяция"
    else:
        lo = srt[-1][0]
        text = f"Ўтган йил — Тугаган йиллар натижалари: «{names.get(lo, '')}» дан ўртача кунлик суръат билан {day} кунга"
    if st: text = text.replace("Тугаган йиллар натижалари", f"Амалда «{st['title']}»", 1)
    return {"mode": mode, "text": text, "set": st}

def full_prev(con, year: int, code: str, kind: str, set_id=None):
    """Ўтган йил жами: (year-1, 12 ойлик) якуни (асосий тўплам ёки set_id)."""
    ensure(con)
    pid = find(con, year - 1, 12, set_id)
    return _value(con, pid, code, kind) if pid else None

# ------------------------------------------------------------------ yozish
def _num(v, what: str):
    if v in (None, ""): return None
    try: x = float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))
    except ValueError: raise ValueError(f"{what}: рақам эмас — «{v}»")
    if x != x or x < 0: raise ValueError(f"{what}: манфий бўлмасин")
    return round(x, 3)

def _check(con, year: int, months: int) -> tuple[int, int]:
    year, months = int(year), int(months)
    if not 2000 <= year <= 2100: raise ValueError("Йил нотўғри")
    if not 1 <= months <= 12: raise ValueError("Давр 1–12 ой оралиғида бўлсин")
    my = max_year(con)
    if year > my: raise ValueError(f"Бу бўлим фақат тугаган йиллар учун — {my} ва ундан олдинги йилни танланг (жорий йил режаси «Режа» панелида, амалдаги кўрсаткич кунлик маълумотдан шаклланади)")
    return year, months

def ensure_period(con, year: int, months: int, set_id=None) -> int:
    """Давр бор бўлса id, бўлмаса яратиб id қайтаради (тўплам: set_id ёки йилнинг асосийси — йўқ бўлса яратилади)."""
    ensure(con)
    st = get_set(con, set_id) if set_id else None
    if set_id and not st: raise ValueError("«Амалда» тўплами топилмади — саҳифани янгиланг")
    year, months = _check(con, int(st["year"]) if st else year, months)
    pid = find(con, year, months, set_id)
    if pid: return pid
    with con:
        sid = st["id"] if st else main_set(con, year, create=True)
        pid = con.execute("INSERT INTO period_results(set_id,year,months,title,source) VALUES(?,?,?,?,?)",
                          (sid, year, months, label(year, months), "manual")).lastrowid
        db.log(con, "period_created", id=pid, year=year, months=months, set_id=sid)
    return pid

def save(con, year: int, months: int, rows: list[dict], set_id=None) -> dict:
    """rows: [{district_code, kind, fact}] — бўш катак = None (киритилмаган)."""
    pid = ensure_period(con, year, months, set_id)
    known = {d["code"] for d in report.districts(con)}
    db.backup_if_stale("period")
    n = 0
    with con:
        for r in rows:
            c, k = str(r.get("district_code") or ""), r.get("kind")
            if c not in known or k not in KINDS: continue
            con.execute("INSERT OR REPLACE INTO period_values(period_id,district_code,kind,fact) VALUES(?,?,?,?)",
                        (pid, c, k, _num(r.get("fact"), f"{c} · {k}"))); n += 1
        con.execute("UPDATE period_results SET updated_at=datetime('now','localtime') WHERE id=?", (pid,))
        con.execute("UPDATE fact_sets SET updated_at=datetime('now','localtime') WHERE id=(SELECT set_id FROM period_results WHERE id=?)", (pid,))
        db.log(con, "period_saved", id=pid, rows=n)
    return {"ok": True, "id": pid, "rows": n}

def delete(con, year: int, months: int, set_id=None):
    ensure(con)
    pid = find(con, int(year), int(months), set_id)
    if not pid: return {"ok": True}
    with con:
        con.execute("DELETE FROM period_values WHERE period_id=?", (pid,))
        con.execute("DELETE FROM period_results WHERE id=?", (pid,))
        db.log(con, "period_deleted", id=pid, year=year, months=months)
    return {"ok": True}

# ------------------------------------------------------------------ «Амалда» тўпламлари (29.09.2026)
def sets(con) -> list[dict]:
    """Тўпламлар рўйхати (карталар учун): ҳар бири даврлари (N ойлик), тўлдирилганлиги ва жамилари билан."""
    ensure(con)
    codes = [d["code"] for d in report.districts(con)]
    out = []
    for r in con.execute("SELECT * FROM fact_sets ORDER BY year DESC, main DESC, id"):
        st = dict(r); st["periods"] = []
        for p in con.execute("SELECT id, months, source, updated_at FROM period_results WHERE set_id=? ORDER BY months", (st["id"],)):
            vals = _vals(con, p["id"]); tot = _totals(vals, codes)
            filled = sum(1 for c in codes for k in KINDS if vals[c][k]["fact"] is not None)
            st["periods"].append({"id": p["id"], "months": p["months"], "label": label(st["year"], p["months"]), "source": p["source"],
                                  "updated_at": p["updated_at"], "filled": filled, "cells": len(codes) * len(KINDS),
                                  "sanoat": tot["sanoat"]["fact"], "meva": tot["meva"]["fact"], "all": tot["all"]["fact"]})
        out.append(st)
    return out

def year_sets(con, year: int) -> list[dict]:
    """Бир йил тўпламлари (танлаш рўйхатлари учун): id, title, main, даврлар (тўлдирилганлари)."""
    return [s for s in sets(con) if s["year"] == int(year)]

def _check_title(con, year: int, title: str, skip_id=None) -> str:
    t = re.sub(r"\s+", " ", str(title or "")).strip()
    if not t: raise ValueError("Тўплам номини киритинг (масалан: Вазирлик, Ҳокимият)")
    if len(t) > 80: raise ValueError("Ном жуда узун (80 белгигача)")
    dup = any(r["title"].casefold() == t.casefold() for r in con.execute("SELECT title FROM fact_sets WHERE year=? AND id<>?", (int(year), int(skip_id or 0))))
    if dup: raise ValueError(f"{year} йилда «{t}» номли тўплам бор — бошқа ном беринг")
    return t

def create_set(con, year: int, title: str, note: str = "", copy_from=None) -> dict:
    """Янги тўплам. Йилнинг биринчи тўплами — асосий. copy_from — бошқа тўпламнинг ҳамма даврлари нусхаси (фақат шу йилда)."""
    ensure(con)
    year = int(year)
    if not 2000 <= year <= 2100: raise ValueError("Йил нотўғри")
    my = max_year(con)
    if year > my: raise ValueError(f"«Амалда» — фақат тугаган йиллар учун: {my} ва ундан олдинги йилни танланг (жорий йил амалдаги кўрсаткичи кунлик маълумотдан шаклланади)")
    t = _check_title(con, year, title)
    src = get_set(con, copy_from) if copy_from else None
    if copy_from and not src: raise ValueError("Нусха олинадиган тўплам топилмади")
    if src and int(src["year"]) != year: raise ValueError("Нусха фақат шу йил тўпламидан олинади")
    db.backup_if_stale("fact_set")
    with con:
        is_main = 0 if con.execute("SELECT 1 FROM fact_sets WHERE year=? LIMIT 1", (year,)).fetchone() else 1
        sid = con.execute("INSERT INTO fact_sets(year, title, note, main) VALUES(?,?,?,?)", (year, t, (note or "").strip(), is_main)).lastrowid
        n = 0
        if src:
            for p in con.execute("SELECT * FROM period_results WHERE set_id=?", (src["id"],)).fetchall():
                pid = con.execute("INSERT INTO period_results(set_id,year,months,title,note,source) VALUES(?,?,?,?,?,?)",
                                  (sid, year, p["months"], p["title"], p["note"], f"copy:{src['id']}")).lastrowid
                con.execute("INSERT INTO period_values(period_id,district_code,kind,fact) SELECT ?,district_code,kind,fact FROM period_values WHERE period_id=?", (pid, p["id"]))
                n += 1
        db.log(con, "fact_set_created", id=sid, year=year, title=t, copy_from=copy_from, periods=n)
    return {"ok": True, "id": sid, "main": bool(is_main), "periods": n}

def update_set(con, set_id: int, title: str | None = None, note: str | None = None, main: bool | None = None) -> dict:
    ensure(con)
    st = get_set(con, set_id)
    if not st: raise ValueError("Тўплам топилмади")
    with con:
        if title is not None:
            t = _check_title(con, st["year"], title, st["id"])
            con.execute("UPDATE fact_sets SET title=?, updated_at=datetime('now','localtime') WHERE id=?", (t, st["id"]))
        if note is not None:
            con.execute("UPDATE fact_sets SET note=?, updated_at=datetime('now','localtime') WHERE id=?", ((note or "").strip(), st["id"]))
        if main:
            con.execute("UPDATE fact_sets SET main=0 WHERE year=?", (st["year"],))
            con.execute("UPDATE fact_sets SET main=1 WHERE id=?", (st["id"],))
        db.log(con, "fact_set_updated", id=st["id"], title=title, main=bool(main))
    return {"ok": True}

def delete_set(con, set_id: int) -> dict:
    """Тўплам ва унинг ҳамма даврлари ўчади. Асосийси ўчса — шу йилнинг энг эски бошқа тўплами асосий бўлади."""
    ensure(con)
    st = get_set(con, set_id)
    if not st: return {"ok": True}
    db.backup_if_stale("fact_set", minutes=0)
    with con:
        for p in con.execute("SELECT id FROM period_results WHERE set_id=?", (st["id"],)).fetchall():
            con.execute("DELETE FROM period_values WHERE period_id=?", (p["id"],))
        con.execute("DELETE FROM period_results WHERE set_id=?", (st["id"],))
        con.execute("DELETE FROM fact_sets WHERE id=?", (st["id"],))
        if st["main"]:
            r = con.execute("SELECT id FROM fact_sets WHERE year=? ORDER BY id LIMIT 1", (st["year"],)).fetchone()
            if r: con.execute("UPDATE fact_sets SET main=1 WHERE id=?", (r["id"],))
        db.log(con, "fact_set_deleted", id=st["id"], year=st["year"], title=st["title"])
    return {"ok": True}

# ------------------------------------------------------------------ Excel: shablon va yuklash
def file_name(year: int, months: int) -> str:
    return f"{year} йил {int(months)} ойлик экспорт якуни.xlsx"

def template_xlsx(con, year: int, months: int, out, set_id=None):
    """Битта рақам устунли шаблон: Т/р · Туман/шаҳарлар · Код · «{Y} йил январь–{ой} амалда».
    Ҳар туман — 3 қатор (жами формула, саноат, мева-сабзавот). Тўлдириб қайта юклаш мумкин."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="999999"); bd = Border(thin, thin, thin, thin)
    st = get_set(con, set_id) if set_id else None
    year, months = (int(st["year"]) if st else int(year)), int(months); mon = MONTHS_GEN[months - 1]
    pid = find(con, year, months, set_id)
    vals = _vals(con, pid)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "03_Бух"
    ws.merge_cells("A1:D1")
    ws["A1"] = f"Бухоро вилоятининг {year} йил январь–{mon} ойларидаги экспорт якуни тўғрисида МАЪЛУМОТ" + (f" («{st['title']}»)" if st else "")
    ws["A1"].font = Font(name="Times New Roman", size=13, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", wrap_text=True); ws.row_dimensions[1].height = 36
    ws["D2"] = "минг долл."; ws["D2"].alignment = Alignment(horizontal="right")
    hdr = [("A", "Т/р"), ("B", "Туман/шаҳарлар"), ("C", "Код"), ("D", f"{year} йил январь–{mon} амалда")]
    for col, h in hdr:
        c = ws[f"{col}3"]; c.value = h
        c.font = Font(name="Times New Roman", bold=True, size=11, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F6E63")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = bd
    ws.row_dimensions[3].height = 44
    for col, w in (("A", 6), ("B", 32), ("C", 10), ("D", 24)): ws.column_dimensions[col].width = w
    ds = report.districts(con); r = 4
    reg = r; r += 3; first = r
    for i, d in enumerate(ds, 1):
        ws.cell(r, 1, i); ws.cell(r, 2, d["name_uz"]).font = Font(bold=True); ws.cell(r, 3, d["code"])
        c = ws.cell(r, 4, f"=D{r + 1}+D{r + 2}"); c.number_format = "#,##0.000"; c.font = Font(bold=True)
        for k, off in (("sanoat", 1), ("meva", 2)):
            ws.cell(r + off, 2, "саноат маҳсулотлари" if k == "sanoat" else "мева-сабзавотлар").alignment = Alignment(indent=2)
            ws.cell(r + off, 3, d["code"])
            v = vals[d["code"]][k]["fact"]
            cc = ws.cell(r + off, 4, round(v, 3) if v is not None else None); cc.number_format = "#,##0.000"
        for rr in range(r, r + 3):
            for j in range(1, 5): ws.cell(rr, j).border = bd
        r += 3
    last = r - 1
    ws.cell(reg, 2, "Бухоро вилояти").font = Font(bold=True); ws.cell(reg, 3, report.REGION)
    ws.cell(reg + 1, 2, "саноат маҳсулотлари").alignment = Alignment(indent=2)
    ws.cell(reg + 2, 2, "мева-сабзавотлар").alignment = Alignment(indent=2)
    L = get_column_letter(4)
    ws.cell(reg, 4, f"={L}{reg + 1}+{L}{reg + 2}").number_format = "#,##0.000"; ws.cell(reg, 4).font = Font(bold=True)
    ws.cell(reg + 1, 4, f"=SUMIF($B${first}:$B${last},\"саноат*\",{L}{first}:{L}{last})").number_format = "#,##0.000"
    ws.cell(reg + 2, 4, f"=SUMIF($B${first}:$B${last},\"мева*\",{L}{first}:{L}{last})").number_format = "#,##0.000"
    for rr in range(reg, reg + 3):
        for j in range(1, 5): ws.cell(rr, j).border = bd; ws.cell(rr, j).fill = PatternFill("solid", fgColor="EAF3F1")
    ws.cell(last + 2, 2, "Изоҳ: фақат «саноат маҳсулотлари» ва «мева-сабзавотлар» қаторларини тўлдиринг (минг $). "
                         "«Код» устунини ўзгартирманг; туман ва вилоят қаторлари — формула. Бўш катак = киритилмаган.").font = Font(italic=True, size=9, color="666666")
    ws.freeze_panes = "D4"
    wb.save(out); return out

_SUFFIX = re.compile(r"(тумани|туман|шаҳри|шахри|шаҳар|шахар|ш\.|т\.)")
def _dkey(name: str) -> str:
    """Туман номи калити: кирилл ҳарфлари асосга, унлилар ташланади, «шаҳри/тумани» тури сақланади
    («Қоравулбозор тумани» ≈ «Қоровулбозор тумани», «Ромитон» ≈ «Ромитан»)."""
    s = str(name or "").lower().replace("ў", "у").replace("қ", "к").replace("ғ", "г").replace("ҳ", "х").replace("ё", "е").replace("ъ", "").replace("ь", "")
    typ = "ш" if re.search(r"шах|ш\.", s) else "т"
    s = _SUFFIX.sub(" ", s)
    base = re.sub(r"[^а-я]", "", s.split()[0] if s.split() else "")
    return re.sub(r"[аеиоуыэюя]", "", base) + typ

_MONTHS = ("январ", "феврал", "март", "апрел", "май", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр")
def _file_months(h: str) -> int | None:
    """Сарлавҳадан давр: «8 ойда» → 8, «январь–сентябрь» → 9 (энг охирги ой номи). Аниқланмаса None."""
    h = str(h or "").lower()
    m = re.search(r"(\d{1,2})\s*ой", h)
    if m and 1 <= int(m.group(1)) <= 12: return int(m.group(1))
    found = [(h.rfind(s), i + 1) for i, s in enumerate(_MONTHS) if re.search(rf"(?<![а-яўқғҳ]){s}", h)]
    return max(found)[1] if found else None

def import_xlsx(con, year: int, months: int, path, set_id=None) -> dict:
    """Сайт шаблонини ёки вазирлик «03_Бух»/«СВОД» варағини ўқийди.
    Рақам устуни: сарлавҳасида танланган ЙИЛ ва «амалда» бўлган устун (масалан «2025 йил 8 ойда амалда»);
    топилмаса — шаблондаги ягона «амалда» устуни.
    Туман қаторидан кейинги «саноат …» / «мева …» қаторлари тур сифатида олинади."""
    from .xl import read_sheet, sheet_names, _norm, num
    st = get_set(con, set_id) if set_id else None
    if set_id and not st: raise ValueError("«Амалда» тўплами топилмади — саҳифани янгиланг")
    year, months = _check(con, int(st["year"]) if st else year, months)
    rows, sheet = None, None
    for nm in sorted(sheet_names(path), key=lambda n: (0 if "бух" in n.lower() else 1)):
        rr = read_sheet(path, nm)
        if any(any("туман" in _norm(c) for c in row) and any(("амалда" in _norm(c)) or ("код" == _norm(c)) or ("экспорт" in _norm(c)) for c in row) for row in rr[:14]):
            rows, sheet = rr, nm; break
    if rows is None: raise ValueError("Туманлар жадвали топилмади: сарлавҳада «Туман» ва «Амалда» (ёки «Код») бўлиши керак")
    hi = next(i for i, row in enumerate(rows[:14]) if any("туман" in _norm(c) for c in row) and any(("амалда" in _norm(c)) or ("код" == _norm(c)) or ("экспорт" in _norm(c)) for c in row))
    top = [_norm(c) for c in rows[hi]]
    sub = [_norm(c) for c in rows[hi + 1]] if hi + 1 < len(rows) else []
    filled, cur = [], ""                       # merged сарлавҳа: юқори қатор бўш бўлса чапдагиси давом этади
    for c in top:
        cur = c or cur; filled.append(cur)
    head = [f"{t} {s}".strip() for t, s in zip(filled, sub + [""] * max(0, len(filled) - len(sub)))]
    c_name = next((i for i, c in enumerate(top) if "туман" in c), None)
    c_code = next((i for i, c in enumerate(top) if c == "код"), None)
    ys = str(year)
    c_val = next((i for i, h in enumerate(head) if ys in h and "амалда" in h), None)
    if c_val is None: c_val = next((i for i, h in enumerate(head) if ys in h and "экспорт" in h), None)
    if c_val is None:
        cand = [i for i, h in enumerate(head) if "амалда" in h]
        if len(cand) == 1: c_val = cand[0]
    if c_val is None:
        raise ValueError(f"«{year} йил … амалда» устуни топилмади. Шаблонни юклаб олиб тўлдиринг ёки расмий жадвалда шу йил устуни борлигини текширинг")
    fm = _file_months(head[c_val])
    if fm and fm != months:
        raise ValueError(f"Файл {fm} ойлик («{head[c_val]}»), экранда эса {months} ойлик танланган. Даврни {fm} ойлик қилиб танланг ёки {months} ойлик шаблонни юкланг")
    dk = {_dkey(d["name_uz"]): d["code"] for d in report.districts(con)}; known = set(dk.values())
    start = hi + (2 if sub and any(sub) and not any("туман" in c for c in sub) else 1)
    out, cur_code, bad = {}, None, []
    for row in rows[start:]:
        nm = str(row[c_name] or "").strip() if c_name is not None and c_name < len(row) else ""
        code = str(row[c_code] or "").strip().split(".")[0] if c_code is not None and c_code < len(row) else ""
        low = nm.lower()
        kind = "sanoat" if low.startswith("саноат") else ("meva" if low.startswith("мева") else None)
        if kind is None:
            if "вилоят" in low: cur_code = report.REGION; continue
            if code in known: cur_code = code
            elif nm:
                cur_code = dk.get(_dkey(nm))
                if cur_code is None: bad.append(nm)
            else: cur_code = None
            continue
        if not cur_code or cur_code == report.REGION: continue
        v = row[c_val] if c_val < len(row) else None
        out[(cur_code, kind)] = {"fact": num(v) if v not in (None, "") else None}
    if not out: raise ValueError("Туман қаторлари топилмади («саноат маҳсулотлари» / «мева-сабзавотлар» қаторлари керак)")
    if all(v["fact"] is None for v in out.values()):
        raise ValueError(f"Файлда рақам йўқ — «{head[c_val]}» устуни бўш. Туманлар бўйича рақамларни киритиб, қайта юкланг")
    res = save(con, year, months, [{"district_code": c, "kind": k, **v} for (c, k), v in out.items()], set_id)
    with con: con.execute("UPDATE period_results SET source=?, updated_at=datetime('now','localtime') WHERE id=?", (f"excel:{sheet}", res["id"]))
    return {"rows": len(out), "sheet": sheet, "column": head[c_val], "unknown": sorted(set(bad)), **res}
