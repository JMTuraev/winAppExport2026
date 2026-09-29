"""Davlat bo'yicha tashqi savdo (eksport + import) — Экспорт географияси sahifasidagi Excel tugmasi, davlat tanlanganda.
Davr — modal oynada: har yil (masalan 2024, 2025, 2026) uchun 1 yanvardan tanlangan kungacha (`ends`), standart — oxirgi yopilgan oy.
Fayl boshida bojxona shaklidagi 2 jadval (ЭК/ИМ давлат-товар: соҳа → маҳсулот, har yil netto + qiymat, farq oxirgi 2 yil),
keyin Умумий, Экспорт, Импорт, номма-ном, Соҳа × туман (oxirgi yil vs undan oldingi tanlangan yil).
Eksport — «bojxona uslubi»: bojxona bazasi davri (items) — eksportyor bo'yicha (sanoat + meva-sabzavot, faqat Buxoro korxonalari),
kunlik GTD davri — sanoat/meva_x eksportyor bo'yicha, meva — karantin qoidasi (Buxoroda yetishtirilgan, tuman — yetishtirilgan joy).
Shu sababli bojxona 8 oylik jadvali bilan teng; «Экспорт географияси» sahifasidagi davlat qatoridan baza davri mevasi qadar farq qiladi.
Import: bojxona bazasi (items regime='ИМ'), faqat Buxoro korxonalari («Бухоро эмас» — yo'q, Buxoro davri hisobga olinadi),
davlat — «страна отправления» (declarations.country); kunlik GTD'da import yo'q — oxirgi baza sanasidan keyingi kunlar kesiladi."""
from __future__ import annotations
import calendar
import datetime as dt
from collections import defaultdict
from . import db, report, geo

OTHER = "Туман аниқланмаган"
NOSECT = "Соҳа аниқланмаган"
NOTARM = "Тармоқ белгиланмаган"
MEVA = "Мева-сабзавот маҳсулотлари"      # meva-sabzavot qatorlari uchun yagona tarmoq nomi (korxonalar reyestridagi nom)
MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]


def _is_meva_tarm(t) -> bool:
    return (t or "").strip().startswith("Мева-сабзавот")


def _md(y: int, md: str) -> str:
    """y yilning md (MM-DD) kuni; 29-fevral kabisa bo'lmagan yilda 28 ga tushadi."""
    m, d = int(md[:2]), int(md[3:5])
    return f"{y}-{m:02d}-{min(d, calendar.monthrange(y, m)[1]):02d}"


def _month_end(s: str) -> bool:
    return int(s[8:10]) == calendar.monthrange(int(s[:4]), int(s[5:7]))[1]


def _per_words(e: str) -> str:
    """01.01 — e davri so'z bilan: «январь-август» (oy oxiri) yoki «1 январь — 28 сентябрь» (kunlik)."""
    m = int(e[5:7])
    if _month_end(e): return MONTHS[0] if m == 1 else f"{MONTHS[0]}-{MONTHS[m - 1]}"
    return f"1 {MONTHS[0]} — {int(e[8:10])} {MONTHS[m - 1]}"


def meta(con) -> dict:
    """Modal uchun: oxirgi sana, import oxirgi sanasi, oxirgi yopilgan oy va yillar."""
    last = report.last_data_date(con) or dt.date.today().isoformat()
    imp_last = con.execute("SELECT max(rdate) FROM items WHERE regime='ИМ'").fetchone()[0] or ""
    first = con.execute("SELECT min(rdate) FROM items").fetchone()[0] or ""
    y, m = int(last[:4]), int(last[5:7])
    if not _month_end(last):
        m -= 1
        if m == 0: y, m = y - 1, 12
    return {"last_date": last, "imp_last": imp_last, "first_date": first, "closed_year": y, "closed_month": m,
            "closed_end": f"{y}-{m:02d}-{calendar.monthrange(y, m)[1]:02d}", "years": [y - 2, y - 1, y]}


# ------------------------------------------------------------------ records
class _Ctx:
    def __init__(self, con, q):
        from . import dispute
        self.con = con
        self.targets = geo._vals(q, "country")
        self.dists, self.tarmoqs, self.kinds = geo._vals(q, "district"), geo._vals(q, "tarmoq"), geo._vals(q, "kind")
        self.with_memo = str(q.get("memo") or "1") != "0"
        self.through = db.get_setting(con, "opening_through_date") or ""
        self.comp = {r["inn"]: dict(r) for r in con.execute(
            "SELECT inn, name, district_code, tarmoq, kind, excluded, bukhara, bukhara_from, bukhara_to FROM companies")}
        self.tset = dispute.template_inns(con)
        self.hs = {r["hs10"]: dict(r) for r in con.execute("SELECT hs10, tovar1, tovar2, sprav1 FROM hs_codes")}
        self.dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
        self.dorder = [d["code"] for d in report.districts(con)]

    def in_bukhara(self, c, d):
        if c is None or c.get("bukhara") == 0: return False
        return (c.get("bukhara_from") or "") <= d <= (c.get("bukhara_to") or "9999")


def _rec(ctx, inn, name, dc, tarmoq, d, value, netto, hs10, tovar1, tovar2, pname, gtd, excl):
    return {"inn": inn, "name": name, "dc": dc or "", "tarmoq": (tarmoq or "").strip(), "date": d, "value": value or 0.0,
            "netto": netto or 0.0, "hs4": (hs10 or "")[:4], "tovar1": tovar1 or NOSECT, "tovar2": tovar2 or "—",
            "pname": pname, "gtd": gtd, "excl": bool(excl)}


def export_recs(ctx: _Ctx, f: str, t: str) -> list:
    """Eksport qatorlari — «bojxona uslubi» (modul izohiga qarang). Barcha yillar uchun bir xil qoida:
    items (bojxona bazasi) — sanoat va meva-sabzavot eksportyor bo'yicha (Buxoro korxonasi, Buxoro davri);
    gtd_rows (baza sanasidan keyingi kunlik GTD) — sanoat va meva_x eksportyor bo'yicha, meva — karantin qoidasi.
    Alohida hisob (БНПЗ/SDK) qatorlari excl=True bilan qaytadi."""
    from .companies import country_uz, has_base_detail
    from . import dispute
    con, out = ctx.con, []
    meva_filter = any(_is_meva_tarm(x) for x in ctx.tarmoqs)

    def kind_ok(k):
        return k in ("sanoat", "meva", "meva_x", "bnpz", "memo")      # bnpz/memo — excl belgisi bilan, keyin with_memo bo'yicha

    def keep(c, d, k):
        if not ctx.in_bukhara(c, d): return False
        if ctx.dists and c.get("district_code") not in ctx.dists: return False
        is_meva = k in ("meva", "meva_x")
        if ctx.tarmoqs and not ((c.get("tarmoq") or "").strip() in ctx.tarmoqs or (is_meva and meva_filter)): return False
        if ctx.kinds:
            if is_meva:
                if "meva" not in ctx.kinds: return False
            elif c.get("kind") not in ctx.kinds: return False
        return True

    def tarm_of(c, k):
        t = (c.get("tarmoq") or "").strip()
        return MEVA if k in ("meva", "meva_x") or _is_meva_tarm(t) else t

    if has_base_detail(con):
        src = list(con.execute("""SELECT i.inn, i.rdate, i.gtd, d.country, i.value, i.netto, i.hs10, i.tovar1, i.tovar2, i.sprav1, i.kind
                                  FROM items i JOIN declarations d ON d.decl_id=i.decl_id
                                  WHERE i.regime='ЭК' AND i.rdate>=? AND i.rdate<=?""", (f, t)))
        tpl = [{**x, "hs10": "", "sprav1": None} for x in (dispute.template_items(con, int(t[:4]), f, t) if ctx.tset else [])]
        for r in src + tpl:
            if not kind_ok(r["kind"]): continue
            if r["inn"] in ctx.tset and r["kind"] == "sanoat" and not str(r["gtd"] or "").startswith("tpl:"): continue
            cn = country_uz(r["country"]) if r["country"] else dispute.TPL_T1
            if cn not in ctx.targets: continue
            c = ctx.comp.get(r["inn"])
            if not keep(c, r["rdate"], r["kind"]): continue
            excl = bool(c.get("excluded")) or r["kind"] in ("bnpz", "memo")
            out.append(_rec(ctx, r["inn"], c.get("name"), c.get("district_code"), tarm_of(c, r["kind"]), r["rdate"], r["value"], r["netto"],
                            r["hs10"], r["tovar1"], r["tovar2"], r["sprav1"] or r["tovar2"] or r["tovar1"], r["gtd"], excl))
    th = ctx.through
    for r in con.execute("SELECT inn, exporter, district_code dc, report_date, gtd_no, country, stat_usd, netto, tnved, product, kind "
                         "FROM gtd_rows WHERE voided=0 AND report_date>=? AND report_date<=?" + (" AND report_date>?" if th else ""),
                         (f, t, *((th,) if th else ()))):
        if r["kind"] in geo.KIND_SKIP or country_uz(r["country"]) not in ctx.targets: continue
        h = ctx.hs.get(str(r["tnved"] or "")) or {}
        if r["kind"] == "meva":           # karantin qoidasi: yetishtirilgan tuman bo'yicha, eksportyor reyestrda bo'lishi shart emas
            if ctx.kinds and "meva" not in ctx.kinds: continue
            if ctx.dists and (r["dc"] or "?") not in ctx.dists: continue
            if ctx.tarmoqs and not meva_filter: continue
            c = ctx.comp.get(r["inn"]) or {}
            nm, dc, tarm, excl = c.get("name") or r["exporter"], r["dc"], MEVA, False
        else:
            if not kind_ok(r["kind"]): continue
            c = ctx.comp.get(r["inn"])
            if not keep(c, r["report_date"], r["kind"]): continue
            nm, dc, tarm = c.get("name"), c.get("district_code"), tarm_of(c, r["kind"])
            excl = bool(c.get("excluded")) or r["kind"] in ("bnpz", "memo")
        out.append(_rec(ctx, r["inn"], nm, dc, tarm, r["report_date"], r["stat_usd"], (r["netto"] or 0) / 1000, str(r["tnved"] or ""),
                        h.get("tovar1"), h.get("tovar2"), h.get("sprav1") or h.get("tovar2") or (r["product"] or "")[:60] or None, r["gtd_no"], excl))
    return out


def import_recs(ctx: _Ctx, f: str, t: str) -> list:
    """Import qatorlari — bojxona bazasi. Faqat tuman filtri qo'llanadi (tarmoq/tur — eksport tushunchalari)."""
    from .companies import country_uz
    out = []
    for r in ctx.con.execute("""SELECT i.inn, i.rdate, i.gtd, d.country, i.value, i.netto, i.hs10, i.tovar1, i.tovar2, i.sprav1
                                FROM items i JOIN declarations d ON d.decl_id=i.decl_id
                                WHERE i.regime='ИМ' AND i.rdate>=? AND i.rdate<=?""", (f, t)):
        if country_uz(r["country"]) not in ctx.targets: continue
        c = ctx.comp.get(r["inn"])
        if not ctx.in_bukhara(c, r["rdate"]): continue
        if ctx.dists and c.get("district_code") not in ctx.dists: continue
        out.append(_rec(ctx, r["inn"], c.get("name"), c.get("district_code"), c.get("tarmoq"), r["rdate"], r["value"], r["netto"],
                        r["hs10"], r["tovar1"], r["tovar2"], r["sprav1"] or r["tovar2"] or r["tovar1"], r["gtd"], c.get("excluded")))
    return out


def _split(ctx, recs):
    """(hisobdagi, alohida hisobdagi) — «Алоҳида ҳисоб» belgisi o'chiq bo'lsa БНПЗ/SDK chetga olinadi."""
    if ctx.with_memo: return recs, []
    return [r for r in recs if not r["excl"]], [r for r in recs if r["excl"]]


def _sum(recs): return sum(r["value"] for r in recs)


def _group(cur, prev, key):
    g = defaultdict(lambda: {"cur": 0.0, "prev": 0.0, "inns": set(), "pinns": set(), "gtd": set(), "netto": 0.0})
    for r in cur:
        o = g[key(r)]; o["cur"] += r["value"]; o["inns"].add(r["inn"]); o["gtd"].add(r["gtd"]); o["netto"] += r["netto"]
    for r in prev:
        o = g[key(r)]; o["prev"] += r["value"]; o["pinns"].add(r["inn"])
    return g


# ------------------------------------------------------------------ data
def periods(con, q: dict) -> list:
    """Tanlangan yillar davri: [(yil, 'YYYY-01-01', oxirgi kun)], yil bo'yicha tartiblangan.
    q['ends'] = '2024-08-31,2025-08-31,2026-08-31' (modal); bo'lmasa — 'to' (yoki oxirgi sana) kuni bo'yicha 3 yil (eski chaqiruv)."""
    ends = {}
    for s in str(q.get("ends") or "").split(","):
        s = s.strip()
        if not s: continue
        try: d = dt.date.fromisoformat(s[:10])
        except ValueError: raise ValueError(f"Нотўғри сана: {s}")
        ends[d.year] = d.isoformat()
    if not ends:
        _, t = geo._period(con, q); y = int(t[:4])
        ends = {yy: _md(yy, t[5:10]) for yy in (y - 2, y - 1, y)}
    if len(ends) > 6: raise ValueError("Кўпи билан 6 та йил танланади")
    return [(y, f"{y}-01-01", ends[y]) for y in sorted(ends)]


def build(con, q: dict) -> dict:
    ctx = _Ctx(con, q)
    if not ctx.targets: raise ValueError("Давлат танланмаган")
    P = periods(con, q)
    # import ma'lumoti faqat bojxona bazasi sanasigacha: oxirgi yil davri undan o'tsa — shu kunlik davrdagi yillar ham baza kuniga kesiladi
    imp_last = con.execute("SELECT max(rdate) FROM items WHERE regime='ИМ'").fetchone()[0] or ""
    last_end = P[-1][2]
    cut = bool(imp_last) and last_end > imp_last
    cache = {}

    def ex(f, t):
        if ("e", f, t) not in cache: cache[("e", f, t)] = _split(ctx, export_recs(ctx, f, t))
        return cache[("e", f, t)]

    def im(f, t):
        if t < f: return [], []
        if ("i", f, t) not in cache: cache[("i", f, t)] = _split(ctx, import_recs(ctx, f, t))
        return cache[("i", f, t)]

    Y = []
    for y, f, e in P:
        ie = e
        if cut and e[5:] == last_end[5:]: ie = _md(y, imp_last[5:10])
        if imp_last: ie = min(ie, imp_last)
        (e_c, e_x), (i_c, i_x) = ex(f, e), im(f, ie)
        Y.append({"year": y, "f": f, "t": e, "it": ie, "ex": e_c, "exx": e_x, "im": i_c, "imx": i_x})
    cur = Y[-1]
    prv = Y[-2] if len(Y) > 1 else {"year": None, "f": None, "t": None, "it": None, "ex": [], "exx": [], "im": [], "imx": []}

    # yillar dinamikasi va oylar (oxirgi yil: 01.01 — tanlangan kun, oldingi 2 yil butun)
    years = []
    for yy in (cur["year"] - 2, cur["year"] - 1, cur["year"]):
        a, b = f"{yy}-01-01", (cur["t"] if yy == cur["year"] else f"{yy}-12-31")
        years.append({"year": yy, "to": b, "exp": ex(a, b)[0], "imp": im(a, min(b, imp_last) if imp_last else b)[0]})
    return {"ctx": ctx, "Y": Y, "f": cur["f"], "t": cur["t"], "pf": prv["f"], "pt": prv["t"], "it": cur["it"], "ipt": prv["it"],
            "cy": cur["year"], "py": prv["year"], "imp_last": imp_last, "cut": cut,
            "ex_c": cur["ex"], "ex_p": prv["ex"], "im_c": cur["im"], "im_p": prv["im"],
            "ex_cx": cur["exx"], "ex_px": prv["exx"], "im_cx": cur["imx"], "im_px": prv["imx"],
            "years": years, "title": ", ".join(sorted(ctx.targets))}


# ------------------------------------------------------------------ Excel
def export_xlsx(con, q: dict, out) -> object:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    D = build(con, q); ctx = D["ctx"]
    dmy = lambda s: f"{s[8:10]}.{s[5:7]}.{s[:4]}" if s else "—"
    thin = Side(style="thin", color="C9D1CD"); BD = Border(thin, thin, thin, thin)
    HF, HFont = PatternFill("solid", fgColor="0F6E63"), Font(bold=True, color="FFFFFF", size=10)
    TF, DF, SF = PatternFill("solid", fgColor="D6E9E4"), PatternFill("solid", fgColor="EEF5F3"), PatternFill("solid", fgColor="F7F9F8")
    NUM, PCT, INT = '#,##0.0;[Red]-#,##0.0;"–"', '0.0%;[Red]-0.0%;"–"', '#,##0;-#,##0;"–"'
    WRAP = Alignment(wrap_text=True, vertical="center")
    per_cur = f"{dmy(D['f'])} — {dmy(D['t'])}"
    per_prev = f"{dmy(D['pf'])} — {dmy(D['pt'])}" if D["pf"] else "—"
    iper_cur = f"{dmy(D['f'])} — {dmy(D['it'])}"
    iper_prev = f"{dmy(D['pf'])} — {dmy(D['ipt'])}" if D["pf"] else "—"
    if D["it"] < D["f"]:
        iper_cur = iper_prev = f"маълумот йўқ (божхона базаси {dmy(D['imp_last'])} гача)"
    CY, PY = f"{D['cy']} йил", (f"{D['py']} йил" if D["py"] else "—")
    filt = []
    if ctx.dists: filt.append("туман: " + ", ".join(ctx.dn.get(x, x) for x in sorted(ctx.dists)))
    if ctx.tarmoqs: filt.append("тармоқ: " + ", ".join(sorted(ctx.tarmoqs)) + " (фақат экспорт)")
    if ctx.kinds: filt.append("тур: " + ", ".join({"sanoat": "саноат", "meva": "мева-сабзавот"}.get(k, k) for k in sorted(ctx.kinds)) + " (фақат экспорт)")
    filt_extra = list(filt)
    filt.append("алоҳида ҳисоб (БНПЗ, SDK): " + ("киритилган" if ctx.with_memo else "киритилмаган"))

    wb = openpyxl.Workbook()

    # ---------------- 0. Божхона шаклидаги 2 жадвал: ЭК / ИМ давлат-товар (соҳа → маҳсулот, ҳар йил нетто + қиймат)
    def goods(ws, is_exp):
        Yl, n = D["Y"], len(D["Y"])
        rk, ek = ("ex", "t") if is_exp else ("im", "it")
        BLK = Side(style="thin", color="FF000000"); BB = Border(BLK, BLK, BLK, BLK)
        RED, GREEN = "FFFF0000", "FF00B050"
        CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
        g1 = defaultdict(lambda: [[0.0, 0.0] for _ in range(n)])
        g2 = defaultdict(lambda: [[0.0, 0.0] for _ in range(n)])
        tot = [[0.0, 0.0] for _ in range(n)]
        for i, Yi in enumerate(Yl):
            for x in Yi[rk]:
                k1 = x["tovar1"] or NOSECT
                k2 = x["tovar2"] if x["tovar2"] and x["tovar2"] != "—" else "Бошқалар"
                for acc in (g1[k1][i], g2[(k1, k2)][i], tot[i]):
                    acc[0] += x["netto"] or 0; acc[1] += x["value"] or 0
        sk = lambda a: tuple(-a[i][1] for i in range(n - 1, -1, -1)) + tuple(-a[i][0] for i in range(n - 1, -1, -1))
        yrs, ends = [Yi["year"] for Yi in Yl], [Yi[ek] for Yi in Yl]
        same = len({e[5:] for e in ends}) == 1
        if n == 1: ylab = f"{yrs[0]} йил"
        elif yrs == list(range(yrs[0], yrs[-1] + 1)): ylab = f"{yrs[0]}-{yrs[-1]} йиллар"
        else: ylab = ", ".join(map(str, yrs[:-1])) + f" ва {yrs[-1]} йиллар"
        sub = f"({ylab} {_per_words(ends[0])})" if same else "(" + ", ".join(f"{y} йил {_per_words(e)}" for y, e in zip(yrs, ends)) + ")"
        many = len(ctx.targets) > 1
        what = (f"давлат{'лар' if many else ''}ига қилинган экспорт" if is_exp else f"давлат{'лар' if many else ''}идан импорт қилинган")
        last = 1 + 2 * n + (2 if n >= 2 else 0)
        widths = [62] + [15, 15] * n + ([14, 11] if n >= 2 else [])
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
        c = ws.cell(1, 1, f"{D['title']} {what} товарлар бўйича\nМАЪЛУМОТ")
        c.font = Font(name="Arial", size=14, bold=True); c.alignment = CEN
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last); ws.row_dimensions[1].height = 48
        c = ws.cell(2, 1, sub); c.font = Font(name="Arial", size=12, bold=True, color=RED)
        c = ws.cell(2, last, "минг долл."); c.font = Font(name="Arial", size=11, italic=True); c.alignment = Alignment(horizontal="right")
        ws.row_dimensions[2].height = 18
        hf = Font(name="Arial", size=12, bold=True)

        def hcell(r, col, v):
            c = ws.cell(r, col, v); c.font = hf; c.alignment = CEN; c.border = BB
        hcell(3, 1, "Давлат/Товар"); hcell(4, 1, None); ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
        for i, (y, e) in enumerate(zip(yrs, ends)):
            col = 2 + 2 * i
            hcell(3, col, f"{y} йил" + ("" if same else f"\n({_per_words(e)})")); hcell(3, col + 1, None)
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 1)
            hcell(4, col, "Нетто вазни, тн"); hcell(4, col + 1, "Қиймати, \nминг $")
        if n >= 2:
            fc = 2 + 2 * n
            hcell(3, fc, "Қийматдаги фарқи" + (f"\n({yrs[-1]}/{yrs[-2]})" if n > 2 else "")); hcell(3, fc + 1, None)
            ws.merge_cells(start_row=3, start_column=fc, end_row=3, end_column=fc + 1)
            hcell(4, fc, "(+/-)"); hcell(4, fc + 1, "%")
        ws.row_dimensions[3].height = 22 if same and n <= 2 else 36
        ws.row_dimensions[4].height = 32

        def put(r, label, arr, color, bold, indent):
            f_ = Font(name="Arial", size=12, bold=bold, color=color)
            c = ws.cell(r, 1, label); c.font = f_; c.border = BB
            c.alignment = CEN if indent is None else Alignment(horizontal="left", vertical="center", indent=indent, wrap_text=True)
            for i in range(n):
                for j in range(2):
                    v = arr[i][j]
                    c = ws.cell(r, 2 + 2 * i + j, round(v, 6) if abs(v) > 1e-9 else None)
                    c.number_format = "#,##0.00"; c.font = f_; c.border = BB; c.alignment = CEN
            if n >= 2:
                fc = 2 + 2 * n; cv, pv = L(3 + 2 * (n - 1)), L(3 + 2 * (n - 2))
                both0 = abs(arr[n - 1][1]) <= 1e-9 and abs(arr[n - 2][1]) <= 1e-9
                c = ws.cell(r, fc, None if both0 else f"=+{cv}{r}-{pv}{r}"); c.number_format = "#,##0.0"
                c.font = f_; c.border = BB; c.alignment = CEN
                c = ws.cell(r, fc + 1, None if both0 else f'=IF({pv}{r}=0,IF({cv}{r}=0,"",1),{L(fc)}{r}/{pv}{r})')
                c.number_format = "0%"; c.font = f_; c.border = BB; c.alignment = CEN

        r = 5
        put(r, "Жами", tot, RED, True, None); r += 1
        for k1 in sorted(g1, key=lambda k: (sk(g1[k]), k)):
            put(r, k1, g1[k1], GREEN, True, 1); r += 1
            for k in sorted((k for k in g2 if k[0] == k1), key=lambda k: (sk(g2[k]), k[1])):
                put(r, k[1], g2[k], None, False, 2); r += 1
        if not g1:
            c = ws.cell(r, 1, "маълумот йўқ"); c.font = Font(name="Arial", size=12, italic=True); r += 1
        last_row = r - 1
        notes = []
        if is_exp and ctx.through and any(Yi["t"] > ctx.through for Yi in Yl):
            notes.append(f"{dmy(ctx.through)} гача — божхона базаси, кейинги кунлар — кунлик ГТД (мева-сабзавот — карантин қоидаси бўйича).")
        if not is_exp and D["cut"]:
            notes.append(f"Импорт — божхона базасидан, {dmy(D['imp_last'])} гача (кунлик ГТД да импорт йўқ); "
                         f"шу сабаб импорт даври {_per_words(ends[-1])} билан чекланган.")
        if filt_extra: notes.append("Филтр: " + "; ".join(filt_extra) + ".")
        xs = sum(len(Yi[rk + "x"]) for Yi in Yl)
        if xs and not ctx.with_memo: notes.append("Алоҳида ҳисобдаги корхоналар (БНПЗ, SDK) киритилмаган.")
        for k, t_ in enumerate(notes):
            rr = r + 1 + k
            c = ws.cell(rr, 1, ("Изоҳ: " if k == 0 else "") + t_); c.font = Font(name="Arial", size=10, italic=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=last); ws.row_dimensions[rr].height = 28
        if notes: last_row = r + len(notes)
        ws.freeze_panes = "B5"
        ws.print_title_rows = "3:4"
        ws.print_area = f"A1:{L(last)}{last_row}"
        ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top, ws.page_margins.bottom = 0.45, 0.5
        ws.oddFooter.center.text = "&P / &N"; ws.oddFooter.center.size = 9
        return ws

    goods(wb.active, True); wb.active.title = "ЭК давлат-товар"
    goods(wb.create_sheet("ИМ давлат-товар"), False)

    def setup(ws, widths, landscape=True):
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4

    def title(ws, text, sub, ncol):
        ws.cell(1, 1, text).font = Font(bold=True, size=14, color="0F3D38")
        ws.cell(2, 1, sub).font = Font(italic=True, size=10, color="5A6470")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        ws.row_dimensions[1].height = 22

    def head(ws, r, cols):
        for i, h in enumerate(cols, 1):
            c = ws.cell(r, i, h); c.font = HFont; c.fill = HF; c.border = BD
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 32

    def row(ws, r, vals, fmts, fill=None, bold=False, indent=0):
        for i, (v, fm) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(r, i, v); c.border = BD
            if fm: c.number_format = fm
            if fill: c.fill = fill
            if bold: c.font = Font(bold=True)
            if i == 2 and indent: c.alignment = Alignment(indent=indent, vertical="center")

    def section(ws, r, text, ncol):
        c = ws.cell(r, 1, text); c.font = Font(bold=True, size=12, color="0F6E63")
        return r + 1

    gr = lambda cur, prev: (cur / prev - 1) if prev else None

    # ---------------- 1. Умумий
    ws = wb.create_sheet("Умумий")
    setup(ws, [46, 16, 16, 15, 12], landscape=False)
    title(ws, f"Бухоро вилояти — {D['title']} билан ташқи савдо", "минг АҚШ доллари · манба: божхона базаси ва кунлик ГТД", 5)
    r = 4
    ws.cell(r, 1, f"Давр ({CY}):").font = Font(bold=True); ws.cell(r, 2, f"экспорт {per_cur}; импорт {iper_cur}"); r += 1
    ws.cell(r, 1, f"Солиштириш ({PY}):").font = Font(bold=True); ws.cell(r, 2, f"экспорт {per_prev}; импорт {iper_prev}"); r += 1
    ws.cell(r, 1, "Филтрлар:").font = Font(bold=True); ws.cell(r, 2, "; ".join(filt)); r += 2
    head(ws, r, ["Кўрсаткич", CY, PY, "Фарқи", "Ўсиш, %"]); r += 1
    E, Ep, I, Ip = _sum(D["ex_c"]), _sum(D["ex_p"]), _sum(D["im_c"]), _sum(D["im_p"])
    for lbl, a, b, bold in [("Товар айланмаси (экспорт + импорт)", E + I, Ep + Ip, True), ("  Экспорт", E, Ep, False),
                            ("  Импорт", I, Ip, False), ("Савдо сальдоси (экспорт − импорт)", E - I, Ep - Ip, True)]:
        row(ws, r, [lbl, a, b, a - b, gr(a, b) if lbl != "Савдо сальдоси (экспорт − импорт)" else None], [None, NUM, NUM, NUM, PCT],
            fill=TF if bold else None, bold=bold); r += 1
    if D["it"] < D["t"] and D["it"] >= D["f"]:
        Es = sum(x["value"] for x in D["ex_c"] if x["date"] <= D["it"]); Eps = sum(x["value"] for x in D["ex_p"] if x["date"] <= D["ipt"])
        row(ws, r, [f"Бир хил давр ({dmy(D['f'])[:5]} — {dmy(D['it'])}): товар айланмаси", Es + I, Eps + Ip, Es + I - Eps - Ip, gr(Es + I, Eps + Ip)],
            [None, NUM, NUM, NUM, PCT], fill=SF); ws.cell(r, 1).alignment = WRAP; ws.row_dimensions[r].height = 28; r += 1
        row(ws, r, [f"  шундан экспорт ({dmy(D['f'])[:5]} — {dmy(D['it'])})", Es, Eps, Es - Eps, gr(Es, Eps)], [None, NUM, NUM, NUM, PCT], fill=SF); r += 1
    nE = lambda recs: len({x["inn"] for x in recs}); nG = lambda recs: len({x["gtd"] for x in recs})
    for lbl, a, b in [("Экспортёр корхоналар, сони", nE(D["ex_c"]), nE(D["ex_p"])), ("Импортёр корхоналар, сони", nE(D["im_c"]), nE(D["im_p"])),
                      ("Экспорт ГТД, сони", nG(D["ex_c"]), nG(D["ex_p"])), ("Импорт ГТД, сони", nG(D["im_c"]), nG(D["im_p"]))]:
        row(ws, r, [lbl, a, b, a - b, gr(a, b)], [None, INT, INT, INT, PCT]); r += 1
    if not ctx.with_memo and (D["ex_cx"] or D["im_cx"] or D["ex_px"] or D["im_px"]):
        r += 1; ws.cell(r, 1, "Алоҳида ҳисобдаги корхоналар (БНПЗ, SDK) — юқоридаги жамига киритилмаган:").font = Font(italic=True, bold=True); r += 1
        for lbl, a, b in [("  Экспорт", _sum(D["ex_cx"]), _sum(D["ex_px"])), ("  Импорт", _sum(D["im_cx"]), _sum(D["im_px"]))]:
            if not (a or b): continue
            row(ws, r, [lbl, a, b, a - b, gr(a, b)], [None, NUM, NUM, NUM, PCT], fill=SF); r += 1

    r += 1; r = section(ws, r, "Йиллар кесимида", 5)
    head(ws, r, ["Йил", "Экспорт", "Импорт", "Товар айланмаси", "Сальдо"]); r += 1
    for Y in D["years"]:
        e, i_ = _sum(Y["exp"]), _sum(Y["imp"])
        lab = f"{Y['year']} йил" + (f" (01.01 — {dmy(Y['to'])[:5]})" if Y["year"] == D["years"][-1]["year"] else "")
        row(ws, r, [lab, e, i_, e + i_, e - i_], [None, NUM, NUM, NUM, NUM]); r += 1

    r += 1; y0, y1 = D["years"][-1], D["years"][-2]
    r = section(ws, r, f"Ойлар кесимида ({y0['year']} ва {y1['year']})", 5)
    head(ws, r, ["Ой", f"Экспорт {y0['year']}", f"Экспорт {y1['year']}", f"Импорт {y0['year']}", f"Импорт {y1['year']}"]); r += 1
    OY = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
    def bym(recs):
        m = [0.0] * 12
        for x in recs: m[int(x["date"][5:7]) - 1] += x["value"]
        return m
    m = [bym(y0["exp"]), bym(y1["exp"]), bym(y0["imp"]), bym(y1["imp"])]
    em, im_ = int(D["t"][5:7]), int((min(D["t"], D["imp_last"]) if D["imp_last"] else D["t"])[5:7])
    for k in range(12):
        v = [m[0][k] if k < em else None, m[1][k], m[2][k] if k < im_ else None, m[3][k]]
        row(ws, r, [OY[k]] + v, [None, NUM, NUM, NUM, NUM]); r += 1
    row(ws, r, ["Жами"] + [sum(m[j]) for j in range(4)], [None, NUM, NUM, NUM, NUM], fill=TF, bold=True); r += 2
    notes = [
        "Изоҳлар:",
        f"• Экспорт — божхона базаси ({dmy(ctx.through)} гача) ва ундан кейинги кунлик ГТД. Мева-сабзавот: божхона базаси даврида — "
        "Бухоро экспортёрлари бўйича (божхона маълумоти билан бир хил), кунлик ГТД даврида — карантин қоидаси (Бухорода етиштирилган, "
        "туман — етиштирилган жой). Шу сабаб «Экспорт географияси» саҳифасидаги давлат қаторидан база давридаги мева-сабзавот қадар фарқ қилиши мумкин.",
        f"• Импорт — божхона базасидан (охирги сана {dmy(D['imp_last'])}); кунлик ГТД фақат экспорт. Давлат — жўнатувчи давлат (ГТД 15-устун).",
        f"• Солиштириш ({PY}) — божхона базаси, модалда танланган давр (давлат кесимида Созламалардаги натижалар йўқ).",
        "• Фақат Бухоро корхоналари: «Бухоро эмас» белгиланганлар ва ИНН кўрсатилмаган қаторлар ҳисобга олинмайди, Бухорода ҳисобга олиш даври инобатга олинади.",
    ]
    for n in notes:
        c = ws.cell(r, 1, n); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(bold=(n == "Изоҳлар:"), size=9, color="5A6470")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 14 if n == "Изоҳлар:" else 40
        r += 1
    ws.freeze_panes = None

    # ---------------- 2. Экспорт / Импорт (tahlil varaqlari)
    def dname(dc): return ctx.dn.get(dc, OTHER)
    def dsort(keys): return sorted(keys, key=lambda k: (ctx.dorder.index(k) if k in ctx.dorder else 999, k))

    def analysis(name, cur, prev, per, pper, is_exp):
        ws = wb.create_sheet(name)
        setup(ws, [6, 52, 15, 15, 14, 11, 10, 12, 10, 12])
        word = "экспорти" if is_exp else "импорти"
        title(ws, f"Бухоро вилоятининг {D['title']} билан {word}",
              f"{CY}: {per} · {PY}: {pper} · минг АҚШ доллари", 10)
        tot, ptot = _sum(cur), _sum(prev)
        cols = ["№", "", CY, PY, "Фарқи", "Ўсиш, %", "Улуши, %", "Корхоналар", "ГТД", "Нетто, т"]
        fm = [INT, None, NUM, NUM, NUM, PCT, PCT, INT, INT, NUM]
        r = 4

        def block(r, heading, first_col, g, order, sub=None):
            r = section(ws, r, heading, 10); head(ws, r, cols[:1] + [first_col] + cols[2:]); r += 1
            for n, k in enumerate(order, 1):
                o = g[k]
                row(ws, r, [n, k, o["cur"], o["prev"], o["cur"] - o["prev"], gr(o["cur"], o["prev"]),
                            o["cur"] / tot if tot else None, len(o["inns"]), len(o["gtd"]), o["netto"]], fm, fill=DF if sub else None, bold=bool(sub))
                r += 1
                if sub:
                    for s, so in sub(k):
                        row(ws, r, [None, s, so["cur"], so["prev"], so["cur"] - so["prev"], gr(so["cur"], so["prev"]),
                                    so["cur"] / tot if tot else None, len(so["inns"]), len(so["gtd"]), so["netto"]], fm, indent=2)
                        r += 1
            row(ws, r, [None, "Жами", tot, ptot, tot - ptot, gr(tot, ptot), 1 if tot else None,
                        len({x["inn"] for x in cur}), len({x["gtd"] for x in cur}), sum(x["netto"] for x in cur)], fm, fill=TF, bold=True)
            return r + 2

        # tumanlar
        g = _group(cur, prev, lambda x: dname(x["dc"]))
        order = [dname(k) for k in dsort({x["dc"] for x in cur + prev})]
        order = list(dict.fromkeys(order))
        r = block(r, "1. Туманлар кесимида", "Туман", g, order)
        # sohalar (tovar1 -> tovar2)
        g1 = _group(cur, prev, lambda x: x["tovar1"])
        g2 = _group(cur, prev, lambda x: (x["tovar1"], x["tovar2"]))
        o1 = sorted(g1, key=lambda k: (-g1[k]["cur"], -g1[k]["prev"]))
        sub = lambda k: sorted(((k2[1], v) for k2, v in g2.items() if k2[0] == k), key=lambda z: (-z[1]["cur"], -z[1]["prev"]))
        r = block(r, "2. Соҳалар (товар гуруҳлари) кесимида", "Соҳа / маҳсулот гуруҳи", g1, o1, sub=sub)
        if is_exp:
            g3 = _group(cur, prev, lambda x: x["tarmoq"] or NOTARM)
            r = block(r, "3. Корхона тармоқлари кесимида", "Тармоқ", g3, sorted(g3, key=lambda k: (-g3[k]["cur"], -g3[k]["prev"])))
        # mahsulotlar HS4
        nm = {}
        for x in sorted(cur + prev, key=lambda x: -x["value"]):
            nm.setdefault(x["hs4"] or "—", x["pname"])
        g4 = _group(cur, prev, lambda x: x["hs4"] or "—")
        o4 = sorted(g4, key=lambda k: (-g4[k]["cur"], -g4[k]["prev"]))
        r = section(ws, r, f"{4 if is_exp else 3}. Маҳсулотлар (ТН ВЭД 4 белги)", 10)
        head(ws, r, ["№", "Маҳсулот", CY, PY, "Фарқи", "Ўсиш, %", "Улуши, %", "Корхоналар", "ТН ВЭД", "Нетто, т"]); r += 1
        for n, k in enumerate(o4, 1):
            o = g4[k]
            row(ws, r, [n, nm.get(k) or "—", o["cur"], o["prev"], o["cur"] - o["prev"], gr(o["cur"], o["prev"]),
                        o["cur"] / tot if tot else None, len(o["inns"]), k, o["netto"]], fm[:8] + ["@", NUM]); r += 1
        row(ws, r, [None, "Жами", tot, ptot, tot - ptot, gr(tot, ptot), 1 if tot else None, len({x["inn"] for x in cur}), None,
                    sum(x["netto"] for x in cur)], fm, fill=TF, bold=True); r += 2
        ws.freeze_panes = "C4"
        return ws

    # ---------------- 3. Номма-ном
    def nomma(name, cur, prev, per, pper, is_exp):
        ws = wb.create_sheet(name)
        cols = ["№", "№ туман", "ИНН", "Корхона номи", "Тармоқ" if is_exp else "Асосий соҳа", "Асосий маҳсулот", CY,
                PY, "Фарқи", "Ўсиш, %", "Туман ичида, %", "ГТД", "Нетто, т", "Охирги сана"]
        setup(ws, [6, 7, 16, 40, 22, 34, 14, 14, 13, 10, 11, 8, 11, 12])
        word = "экспортёр" if is_exp else "импортёр"
        title(ws, f"{D['title']} — {word} корхоналар номма-ном", f"{CY}: {per} · {PY}: {pper} · минг АҚШ доллари", len(cols))
        r = 4; head(ws, r, cols); r += 1
        g = defaultdict(lambda: {"cur": 0.0, "prev": 0.0, "gtd": set(), "netto": 0.0, "last": None, "name": None, "dc": None,
                                 "tarmoq": None, "p": defaultdict(float), "s": defaultdict(float)})
        # kalit (ИНН, туман): karantin mevasi yetishtirilgan tuman bo'yicha — bitta eksportyor bir necha tumanda turishi mumkin
        for x in cur:
            o = g[(x["inn"], x["dc"] or "")]; o["cur"] += x["value"]; o["gtd"].add(x["gtd"]); o["netto"] += x["netto"]
            o["last"] = max(filter(None, [o["last"], x["date"]])); o["p"][x["pname"] or "—"] += x["value"]; o["s"][x["tovar1"]] += x["value"]
            o["name"] = o["name"] or x["name"]; o["dc"] = o["dc"] or x["dc"]; o["tarmoq"] = o["tarmoq"] or x["tarmoq"]
        for x in prev:
            o = g[(x["inn"], x["dc"] or "")]; o["prev"] += x["value"]
            o["name"] = o["name"] or x["name"]; o["dc"] = o["dc"] or x["dc"]; o["tarmoq"] = o["tarmoq"] or x["tarmoq"]
            o.setdefault("pp", defaultdict(float))[x["pname"] or "—"] += x["value"]
            o.setdefault("ps", defaultdict(float))[x["tovar1"]] += x["value"]
        by_d = defaultdict(list)
        for (inn, dc_), o in g.items(): by_d[dc_].append((inn, o))
        n = 0; tot, ptot = _sum(cur), _sum(prev)
        for dc in dsort(by_d):
            lst = sorted(by_d[dc], key=lambda z: (-z[1]["cur"], -z[1]["prev"], z[1]["name"] or ""))
            dc_cur, dc_prev = sum(o["cur"] for _, o in lst), sum(o["prev"] for _, o in lst)
            row(ws, r, [None, None, None, f"{dname(dc)} ({len([1 for _, o in lst if o['cur']])} та)" if any(o["cur"] for _, o in lst) else f"{dname(dc)} (жорий даврда йўқ)", None, None, dc_cur, dc_prev,
                        dc_cur - dc_prev, gr(dc_cur, dc_prev), dc_cur / tot if tot else None,
                        len({gg for _, o in lst for gg in o["gtd"]}), sum(o["netto"] for _, o in lst), None],
                [None] * 6 + [NUM, NUM, NUM, PCT, PCT, INT, NUM, None], fill=DF, bold=True)
            r += 1
            for k, (inn, o) in enumerate(lst, 1):
                n += 1
                prod_src = o["p"] or o.get("pp") or {}
                prod = max(prod_src, key=prod_src.get) if prod_src else None
                sect_src = o["s"] or o.get("ps") or {}
                sect = max(sect_src, key=sect_src.get) if sect_src else None
                row(ws, r, [n, k, inn, o["name"] or inn, (o["tarmoq"] or "") if is_exp else sect, prod, o["cur"], o["prev"],
                            o["cur"] - o["prev"], gr(o["cur"], o["prev"]), o["cur"] / dc_cur if dc_cur else None, len(o["gtd"]),
                            o["netto"], dmy(o["last"]) if o["last"] else None],
                    [INT, INT, "@", None, None, None, NUM, NUM, NUM, PCT, PCT, INT, NUM, None])
                ws.cell(r, 4).alignment = WRAP; ws.cell(r, 6).alignment = WRAP
                r += 1
        row(ws, r, [None, None, None, f"Жами: {len({k[0] for k, o in g.items() if o['cur']})} та корхона", None, None, tot, ptot, tot - ptot,
                    gr(tot, ptot), 1 if tot else None, len({x["gtd"] for x in cur}), sum(x["netto"] for x in cur), None],
            [None] * 6 + [NUM, NUM, NUM, PCT, PCT, INT, NUM, None], fill=TF, bold=True)
        ws.freeze_panes = "E5"; ws.auto_filter.ref = f"A4:{L(len(cols))}{r - 1}"
        ws.print_title_rows = "4:4"
        return ws

    def matrix(ws, r, heading, cur):
        """Соҳалар (қатор) × туманлар (устун), жорий давр."""
        g1 = defaultdict(float); mx = defaultdict(float)
        for x in cur: g1[x["tovar1"]] += x["value"]; mx[(x["tovar1"], x["dc"])] += x["value"]
        secs = sorted((k for k in g1 if g1[k]), key=lambda k: -g1[k]); dks = dsort({x["dc"] for x in cur})
        r = section(ws, r, heading, 3)
        if not secs:
            ws.cell(r, 2, "маълумот йўқ"); return r + 2
        hdr = ["№", "Соҳа", "Жами"] + [dname(dc).replace(" тумани", " т.").replace(" шаҳри", " ш.") for dc in dks]
        for i, h in enumerate(hdr, 1):
            c = ws.cell(r, i, h); c.font = HFont; c.fill = HF; c.border = BD
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 34; r += 1
        for n, s_ in enumerate(secs, 1):
            row(ws, r, [n, s_, g1[s_]] + [mx[(s_, dc)] for dc in dks], [INT, None] + [NUM] * (len(dks) + 1))
            ws.cell(r, 2).alignment = WRAP; r += 1
        row(ws, r, [None, "Жами", sum(g1.values())] + [sum(mx[(s_, dc)] for s_ in secs) for dc in dks], [None, None] + [NUM] * (len(dks) + 1), fill=TF, bold=True)
        return r + 2

    analysis("Экспорт", D["ex_c"], D["ex_p"], per_cur, per_prev, True)
    analysis("Импорт", D["im_c"], D["im_p"], iper_cur, iper_prev, False)
    nomma("Экспорт номма-ном", D["ex_c"], D["ex_p"], per_cur, per_prev, True)
    nomma("Импорт номма-ном", D["im_c"], D["im_p"], iper_cur, iper_prev, False)
    ws = wb.create_sheet("Соҳа × туман")
    setup(ws, [5, 38, 13] + [11] * 14)
    title(ws, f"{D['title']} — соҳалар ва туманлар кесими", f"{CY} · минг АҚШ доллари", 12)
    r = matrix(ws, 4, f"Экспорт ({per_cur})", D["ex_c"])
    matrix(ws, r, f"Импорт ({iper_cur})", D["im_c"])
    wb.save(out); return out


def file_name(q: dict) -> str:
    import re
    cs = sorted(geo._vals(q, "country"))
    nm = ", ".join(cs) if len(cs) <= 3 else f"{', '.join(cs[:3])} ва б."
    nm = re.sub(r'[\\/:*?"<>|]', "_", nm)
    return f"{nm} — экспорт ва импорт {dt.date.today():%d.%m.%Y}.xlsx"
