"""Ҳисоботлар → «Экспорт қилган корхоналар рўйхати» (26.09.2026, Jafar talabi).

Tanlangan davrda (joriy yil, oylar oralig'i — standart: yanvardan oxirgi ma'lumot kunigacha) eksport qilgan korxonalar —
yirik summadan kichigiga: Корхона номи · Тармоқ · Туман · Экспорт ҳажми.
Hisob «Номма-ном» sahifasi bilan bir xil — report.company_amounts (bojxona bazasi + kunlik GTD, «Бухоро эмас» chiqariladi,
Buxoro davri). Meva-sabzavot — karantin qoidasi (Buxoroda yetishtirilgan) bo'yicha tumanlar kesimida alohida bo'lim
(26.09.2026 tuzatish: jami Dashboard «Жами» = саноат + карантин мева bilan teng bo'lishi shart; yanv–avg karantin faqat tuman jami).
БНПЗ va SDK — faqat «Алоҳида ҳисоб» belgilansa.
"""
from __future__ import annotations
import datetime as dt
from collections import defaultdict
from pathlib import Path
from . import db, report

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
TYPES = {"all": "", "sanoat": "саноат маҳсулотлари", "meva": "мева-сабзавот маҳсулотлари"}


def _period(con, q) -> dict:
    last = report.last_data_date(con) or dt.date.today().isoformat()
    ld = dt.date.fromisoformat(last)
    m2 = max(1, min(ld.month, int(q.get("m2") or ld.month)))
    m1 = max(1, min(m2, int(q.get("m1") or 1)))
    partial = m2 == ld.month and (ld + dt.timedelta(days=1)).month == ld.month   # oy hali to'lmagan
    to_day = ld.day if partial else (dt.date(ld.year + (m2 == 12), m2 % 12 + 1, 1) - dt.timedelta(days=1)).day
    months = m2 - m1 + 1
    name = MONTHS[m1 - 1] if m1 == m2 else f"{MONTHS[m1 - 1]}–{MONTHS[m2 - 1]}"
    return {"year": ld.year, "m1": m1, "m2": m2, "months": months, "last": last, "partial": partial,
            "from": f"{ld.year}-{m1:02d}-01", "to": f"{ld.year}-{m2:02d}-{to_day:02d}",
            "label": f"{ld.year} йил {name}" + (f" ({months} ой)" if months > 1 else ""),
            "range": f"{ld.year} йил 1 {MONTHS[m1 - 1]} – {to_day} {MONTHS[m2 - 1]}",
            "phrase": f"{ld.year} йил {name} " + ("ойида" if m1 == m2 else "ойларида")}


def _memo_amounts(con, year: int, D: str) -> dict:
    """{inn: [12]} — БНПЗ (bnpz) va SDK (memo): bojxona bazasi oylari + keyingi kunlik GTD."""
    through = report.opening_date(con) or ""
    skip = report.not_bukhara_inns(con)
    out = defaultdict(lambda: [0.0] * 12)
    for r in con.execute("SELECT inn, month, amount FROM opening_balances WHERE year=? AND kind IN ('bnpz','memo')", (year,)):
        if r["inn"] not in skip: out[r["inn"]][r["month"] - 1] += r["amount"] or 0
    for r in con.execute("""SELECT inn, substr(report_date,6,2) m, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind IN ('bnpz','memo')
                            AND report_date>? AND report_date<=? AND substr(report_date,1,4)=? GROUP BY inn, m""", (through, D, str(year))):
        if r["inn"] not in skip: out[r["inn"]][int(r["m"]) - 1] += r["s"] or 0
    return out


def _karantin(con, per: dict, dsel: set) -> tuple[list, str]:
    """Meva-sabzavot — karantin qoidasi (Buxoroda yetishtirilgan, eksportyordan qat'i nazar), dashboard/Свод bilan bir xil:
    karantin_opening (shablon svodi, faqat tuman jami — korxona kesimi yo'q) + keyingi GTD (kind='meva', o'stirilgan tuman).
    Qaytaradi: ([{district_code, district, value}], izoh)."""
    y, m1, m2, D = per["year"], per["m1"], per["m2"], per["last"]
    through = report.opening_date(con) or ""
    vals = defaultdict(float); note = ""
    om = int(through[5:7]) if through[:4] == str(y) else 0
    if om:
        if m1 == 1 and m2 >= om:
            for r in con.execute("SELECT district_code, ytd FROM karantin_opening WHERE year=?", (y,)): vals[r["district_code"]] += r["ytd"] or 0
        elif m1 <= om:
            note = f"Январь–{MONTHS[om - 1]} мева-сабзавоти (карантин) фақат жами маълум, ойларга бўлинмаган — бу даврга киритилмади"
    for r in con.execute("""SELECT district_code dc, sum(stat_usd) s FROM gtd_rows WHERE voided=0 AND kind='meva' AND report_date>? AND report_date<=?
                            AND substr(report_date,1,4)=? AND CAST(substr(report_date,6,2) AS INTEGER) BETWEEN ? AND ? GROUP BY dc""",
                         (through, D, str(y), m1, m2)):
        vals[r["dc"] or "?"] += r["s"] or 0
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    out = [{"district_code": k, "district": dn.get(k) or "Туман аниқланмаган", "value": round(v, 3)} for k, v in vals.items()
           if v > 0.0005 and (not dsel or k in dsel)]
    out.sort(key=lambda r: -r["value"])
    for i, r in enumerate(out, 1): r["no"] = i
    return out, note


def build(con, q: dict) -> dict:
    """Саноат — корхоналар (Номма-ном саҳифаси ҳисоби), мева-сабзавот — карантин қоидаси бўйича туманлар кесимида
    (Dashboard «Жами» = саноат + карантин мева билан бир хил; мева қисми божхона шаблонида корхона кесимида йўқ)."""
    per = _period(con, q)
    t = (q.get("t") or "all").strip()
    if t not in TYPES: t = "all"
    memo = str(q.get("memo") or "0") == "1"
    dsel = {x for x in (q.get("district") or "").split(",") if x}
    D, y, m1, m2 = per["last"], per["year"], per["m1"], per["m2"]
    acc = defaultdict(lambda: {"sanoat": 0.0, "memo": 0.0})
    if t in ("all", "sanoat"):
        for (inn, kind), a in report.company_amounts(con, D).items():
            if kind == "sanoat": acc[inn]["sanoat"] += sum(a["m"][m1 - 1:m2])
        if memo:
            for inn, mm in _memo_amounts(con, y, D).items(): acc[inn]["memo"] += sum(mm[m1 - 1:m2])
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, full_name, district_code, tarmoq, union_name FROM companies")}
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    t1 = {}   # tarmoq bo'sh bo'lsa — korxonaning shu yildagi eng katta tovar1 guruhi (bojxona bazasi)
    for r in con.execute("SELECT inn, tovar1, sum(value) v FROM items WHERE regime='ЭК' AND year=? AND tovar1 IS NOT NULL GROUP BY inn, tovar1", (y,)):
        if r["inn"] not in t1 or r["v"] > t1[r["inn"]][1]: t1[r["inn"]] = (r["tovar1"], r["v"])
    rows = []
    for inn, v in acc.items():
        total = v["sanoat"] + v["memo"]
        if total <= 0.0005: continue
        c = comp.get(inn) or {}
        if dsel and c.get("district_code") not in dsel: continue
        tar = (c.get("tarmoq") or "").strip() or (t1.get(inn) or ("",))[0] or "—"
        rows.append({"inn": inn, "name": c.get("name") or c.get("full_name") or inn, "full_name": c.get("full_name") or "",
                     "tarmoq": tar, "district_code": c.get("district_code") or "", "district": dn.get(c.get("district_code")) or "—",
                     "sanoat": round(v["sanoat"], 3), "memo": round(v["memo"], 3), "value": round(total, 3), "separate": v["memo"] > 0})
    rows.sort(key=lambda r: (-r["value"], r["name"]))
    meva, meva_note = _karantin(con, per, dsel) if t in ("all", "meva") else ([], "")
    ctot = sum(r["value"] for r in rows); mtot = sum(r["value"] for r in meva); total = ctot + mtot
    for i, r in enumerate(rows, 1):
        r["no"] = i; r["share"] = r["value"] / total * 100 if total else 0
    for r in meva: r["share"] = r["value"] / total * 100 if total else 0
    where = "Бухоро вилоятида"
    if len(dsel) == 1: where = "Бухоро вилояти " + (dn.get(next(iter(dsel))) or "").replace(" тумани", " туманида").replace(" шаҳри", " шаҳрида")
    tname = TYPES[t]
    title = (q.get("title") or "").strip() or \
        f"{where} {per['phrase']} экспорт қилган корхоналар рўйхати" + (f" ({tname})" if tname else "")
    return {"period": per, "type": t, "memo": memo, "title": title, "count": len(rows), "total": round(total, 3),
            "companies_total": round(ctot, 3), "sanoat": round(sum(r["sanoat"] for r in rows), 3), "meva": round(mtot, 3),
            "memo_total": round(sum(r["memo"] for r in rows), 3), "rows": rows, "meva_rows": meva, "meva_note": meva_note}


def meta(con) -> dict:
    last = report.last_data_date(con)
    return {"last_date": last, "year": int(last[:4]) if last else dt.date.today().year, "month": int(last[5:7]) if last else dt.date.today().month,
            "districts": [{"code": d["code"], "name": d["name_uz"]} for d in report.districts(con)]}


def export_xlsx(con, q: dict, out: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    d = build(con, q)
    if not d["rows"] and not d["meva_rows"]: raise ValueError("Танланган шартларда экспорт қилган корхона йўқ")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Корхоналар"
    thin = Side(style="thin", color="9AA5A0"); bd = Border(thin, thin, thin, thin)
    C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LW = Alignment(vertical="center", wrap_text=True)
    HF = PatternFill("solid", fgColor="D9E7E2"); TF = PatternFill("solid", fgColor="EEF4F1"); SF = PatternFill("solid", fgColor="F6F8F7")
    hdr = ["№", "Корхона номи", "Тармоқ номи", "Туман", "Экспорт ҳажми\n(минг АҚШ доллари)", "Жами экспортдаги улуши, %"]
    widths = [6, 50, 30, 22, 20, 16]
    ws.merge_cells("A1:F1"); ws["A1"] = d["title"]
    ws["A1"].font = Font(bold=True, size=13); ws["A1"].alignment = C; ws.row_dimensions[1].height = 36
    p = d["period"]
    ws.merge_cells("A2:F2"); ws["A2"] = f"({p['range']} маълумотлари)"
    ws["A2"].font = Font(italic=True, size=10); ws["A2"].alignment = C
    ws["F3"] = "минг АҚШ доллари"; ws["F3"].font = Font(italic=True, size=9); ws["F3"].alignment = Alignment(horizontal="right")
    HR = 4
    for i, (h, w) in enumerate(zip(hdr, widths), 1):
        c = ws.cell(HR, i, h); c.font = Font(bold=True, size=10); c.fill = HF; c.alignment = C; c.border = bd
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.row_dimensions[HR].height = 44
    def put(r, vals, bold=False, fill=None, size=10):
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = bd; c.font = Font(bold=bold, size=size)
            if fill: c.fill = fill
            c.alignment = C if i in (1, 4) else LW
        ws.cell(r, 5).number_format = "#,##0.0"
        c = ws.cell(r, 6, f"=IF($E${tr}=0,0,E{r}/$E${tr})"); c.border = bd; c.number_format = "0.00%"
        c.font = Font(bold=bold, size=size); c.alignment = C
        if fill: c.fill = fill
    tr = HR + 1; r = tr + 1; heads = []
    sections = []
    if d["type"] in ("all", "sanoat"):
        sections.append(("I" if d["type"] == "all" else "", f"Саноат маҳсулотлари экспорт қилган корхоналар ({d['count']} та)",
                         [[x["no"], x["name"] + (" (алоҳида ҳисоб)" if x["separate"] else ""), x["tarmoq"], x["district"], round(x["value"], 3)] for x in d["rows"]]))
    if d["type"] in ("all", "meva") and d["meva_rows"]:
        sections.append(("II" if d["type"] == "all" else "", "Мева-сабзавот маҳсулотлари (карантин рўйхати бўйича, туманлар кесимида)",
                         [[x["no"], "Мева-сабзавот экспорти", "Мева-сабзавот", x["district"], round(x["value"], 3)] for x in d["meva_rows"]]))
    for roman, name, lines in sections:
        hr_ = r; heads.append(hr_); r += 1
        f0 = r
        for ln in lines: put(r, ln); r += 1
        put(hr_, [roman, name, "", "", f"=SUM(E{f0}:E{r - 1})" if r > f0 else 0], bold=True, fill=SF)
    last = r - 1
    label = f"ЖАМИ ({d['count']} та корхона" + (" + мева-сабзавот" if d["meva_rows"] and d["type"] == "all" else "") + ")" if d["type"] != "meva" else "ЖАМИ"
    put(tr, ["", label, "", "", "=" + "+".join(f"E{h}" for h in heads) if heads else 0], bold=True, fill=TF, size=11)
    first = tr + 1
    if d.get("meva_note"):
        ws.cell(last + 2, 2, "Изоҳ: " + d["meva_note"]).font = Font(italic=True, size=9)
    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A{HR}:F{last}"
    ws.print_title_rows = f"{HR}:{HR}"; ws.print_area = f"A1:F{last}"
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
