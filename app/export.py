"""Excel output.

daily_workbook: one day's GTD rows and company summary.
The cumulative «… йил экспорт» workbook is written on the user's own template: see tpl.py.
"""
from __future__ import annotations
import copy, json, re, datetime as dt
from pathlib import Path
from collections import defaultdict
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from . import db, report
from .xl import _norm, find_sheet

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
def clean(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

def _inn_value(inn: str):
    return int(inn) if inn.isdigit() and len(inn) <= 15 else inn

# ------------------------------------------------------------------ daily workbook
THIN = Side(style="thin", color="C9CFCB")
def _hdr(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, clean(v)); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F6E63")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w

def daily_workbook(con, D: str, out_path: Path) -> dict:
    det = report.day_detail(con, D); dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Кунлик свод"
    ds = f"{D[8:10]}.{D[5:7]}.{D[:4]}"
    ws["A1"] = f"Бухоро вилояти экспорти · {ds} кунлик маълумот"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "минг АҚШ доллари · манба: божхона ГТД (Дата выгрузки)"; ws["A2"].font = Font(italic=True, color="5A6470")
    _hdr(ws, 4, ["Кўрсаткич", "Шу кун", "Йил бошидан"], [34, 16, 16])
    for i, (lbl, key) in enumerate([("Саноат маҳсулотлари", "sanoat"), ("Мева-сабзавот (карантин)", "meva"), ("Жами", "total")], 5):
        ws.cell(i, 1, lbl); ws.cell(i, 2, det["day"][key]); ws.cell(i, 3, det["ytd"][key])
        for c in (2, 3): ws.cell(i, c).number_format = "#,##0.00"
        if key == "total":
            for c in (1, 2, 3): ws.cell(i, c).font = Font(bold=True)
    # district breakdown for the day
    by_d = defaultdict(lambda: [0.0, 0.0])
    for c in det["companies"]: by_d[c["district_code"]][0] += c["day"]
    for m in det["meva_rows"]: by_d[m["district_code"]][1] += m["stat_usd"]
    ws["A10"] = "Туманлар кесимида"; ws["A10"].font = Font(bold=True)
    _hdr(ws, 11, ["Туман / шаҳар", "Саноат", "Мева-сабзавот", "Жами"])
    rr = 12
    for d in report.districts(con):
        s, mv = by_d.get(d["code"], [0, 0])
        ws.cell(rr, 1, d["name_uz"]); ws.cell(rr, 2, s or None); ws.cell(rr, 3, mv or None); ws.cell(rr, 4, f"=SUM(B{rr}:C{rr})")
        for c in (2, 3, 4): ws.cell(rr, c).number_format = "#,##0.00"
        rr += 1
    ws.cell(rr, 1, "Жами").font = Font(bold=True)
    for c in (2, 3, 4): ws.cell(rr, c, f"=SUM({L(c)}12:{L(c)}{rr-1})").number_format = "#,##0.00"
    ws.column_dimensions["D"].width = 16
    # companies
    wc = wb.create_sheet("Корхоналар")
    _hdr(wc, 1, ["ИНН", "Корхона", "Туман", "ГТД сони", "Шу кун", "Шу ой", "Йил бошидан", "Янги"], [14, 44, 22, 10, 14, 14, 16, 8])
    for i, c in enumerate(det["companies"], 2):
        vals = [_inn_value(c["inn"]), c["name"], dn.get(c["district_code"], "?"), c["gtd"], c["day"], c["month"], c["ytd"], "ҳа" if c["new"] else ""]
        for j, v in enumerate(vals, 1):
            cell = wc.cell(i, j, clean(v))
            if j in (5, 6, 7): cell.number_format = "#,##0.00"
    wc.freeze_panes = "A2"
    wm = wb.create_sheet("Мева-сабзавот")
    _hdr(wm, 1, ["Сана", "ИНН", "Экспортёр", "Ўстирилган туман", "Маҳсулот", "Давлат", "Нетто, кг", "минг $"], [11, 14, 36, 26, 44, 16, 12, 12])
    for i, m in enumerate(det["meva_rows"], 2):
        for j, v in enumerate([m["report_date"], _inn_value(m["inn"]), m["exporter"], m["grown_district"], m["product"], m["country"], m["netto"], m["stat_usd"]], 1):
            wm.cell(i, j, clean(v))
    wg = wb.create_sheet("ГТД қаторлари")
    cols = ["report_date", "kind", "inn", "exporter", "district_code", "gtd_no", "item_no", "tnved", "product", "country", "netto", "stat_usd", "invoice_usd", "regime"]
    _hdr(wg, 1, ["Сана", "Тур", "ИНН", "Экспортёр", "Туман", "ГТД", "Товар №", "ТН ВЭД", "Маҳсулот", "Давлат", "Нетто", "Стат. минг $", "Фактура минг $", "Режим"], [11, 8, 14, 34, 18, 26, 8, 12, 40, 14, 12, 12, 12, 10])
    kinds = {"sanoat": "саноат", "meva": "мева (карантин)", "meva_x": "мева · номманом", "memo": "алоҳида", "bnpz": "БНПЗ", "not_bukhara": "Бухоро эмас"}
    for i, r in enumerate(con.execute(f"SELECT {','.join(cols)} FROM gtd_rows WHERE voided=0 AND report_date=? ORDER BY kind, inn", (D,)), 2):
        for j, k in enumerate(cols, 1):
            v = r[k]
            if k == "kind": v = kinds.get(v, v)
            if k == "district_code": v = dn.get(v, v)
            if k == "inn": v = _inn_value(v)
            wg.cell(i, j, clean(v))
    wg.freeze_panes = "A2"
    wb.save(out_path)
    return {"file": str(out_path), "companies": len(det["companies"]), "meva_rows": len(det["meva_rows"])}
