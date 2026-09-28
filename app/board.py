"""Muammo va vazifalar — Trello uslubidagi doska.

Ustunlar cheksiz (issue_columns), har karta bir ustunda (issues.column_id, position). Belgilar (issue_labels) —
muammoga aloqador tashkilotlar; bitta kartaga bir nechta belgi qo'yiladi (issue_label_links). Fayllar
(issue_files) data\\issue_files\\<karta id>\\ papkasida saqlanadi, olib tashlangani _olib_tashlangan ga ko'chadi.

Ustunning is_done belgisi: unga ko'chirilgan karta «ҳал қилинган» (status=closed) bo'ladi va aksincha.

26.09.2026: ustuvorlik (priority), asos hujjat (basis_*), qadamlar (issue_checks), izohlar + avtomatik tarix
(issue_comments: kind='note' — foydalanuvchi yozuvi, 'log' — tizim), column_at — karta ustunga tushgan payt.
"""
from __future__ import annotations
import re, shutil, mimetypes, datetime as dt
from pathlib import Path
from collections import defaultdict
from . import db

# ko'k/yashil-ko'k yo'q: tanlangan chip va select ranglari bilan qo'shilib ketmasin
LABEL_COLORS = ["#B7912A", "#B23A3A", "#7A4DB0", "#E67E22", "#2E8B57", "#C2185B", "#8D6E63", "#6B8E23", "#5D6D7E", "#A0522D", "#B7791F", "#8E44AD"]
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
PRIORITY = {"urgent": "Шошилинч", "high": "Юқори", "normal": "Ўрта", "low": "Паст"}
BASIS = {"bayon": "Йиғилиш баёни", "topshiriq": "Топшириқ", "xat": "Кирувчи хат", "murojaat": "Корхона мурожаати", "boshqa": "Бошқа ҳужжат"}
STALE_DAYS = 14   # open card without any activity this long = «туриб қолган»

def _now() -> str: return dt.datetime.now().isoformat(sep=" ", timespec="seconds")

# ------------------------------------------------------------------ columns
def columns(con) -> list[dict]:
    return [dict(r) for r in con.execute("SELECT * FROM issue_columns ORDER BY position, id")]

def first_column(con, done: bool) -> int | None:
    r = con.execute("SELECT id FROM issue_columns WHERE is_done=? ORDER BY position, id LIMIT 1", (1 if done else 0,)).fetchone()
    if r: return r[0]
    r = con.execute("SELECT id FROM issue_columns ORDER BY position, id LIMIT 1").fetchone()
    return r[0] if r else None

def save_column(con, data: dict) -> dict:
    name = (data.get("name") or "").strip()[:60]
    if not name: raise ValueError("Устун номини ёзинг")
    color = _color(data.get("color")) or "#5D6D7E"
    is_done = 1 if data.get("is_done") else 0
    with con:
        if data.get("id"):
            cid = int(data["id"])
            old = con.execute("SELECT * FROM issue_columns WHERE id=?", (cid,)).fetchone()
            if not old: raise ValueError("Устун топилмади")
            con.execute("UPDATE issue_columns SET name=?, color=?, is_done=? WHERE id=?", (name, color, is_done, cid))
            if is_done != old["is_done"]:   # cards inside follow the column's meaning
                _sync_status_of_column(con, cid, bool(is_done))
        else:
            pos = (con.execute("SELECT coalesce(max(position), 0) FROM issue_columns").fetchone()[0] or 0) + 10
            cid = con.execute("INSERT INTO issue_columns(name, color, position, is_done) VALUES(?,?,?,?)", (name, color, pos, is_done)).lastrowid
        db.log(con, "board_column_saved", id=cid, name=name, is_done=is_done)
    return {"ok": True, "id": cid}

def delete_column(con, cid: int, move_to: int | None) -> dict:
    cols = columns(con)
    if len(cols) <= 1: raise ValueError("Камида битта устун қолиши керак")
    if not any(c["id"] == cid for c in cols): raise ValueError("Устун топилмади")
    n = con.execute("SELECT count(*) FROM issues WHERE column_id=?", (cid,)).fetchone()[0]
    with con:
        if n:
            if not move_to or move_to == cid or not any(c["id"] == move_to for c in cols):
                raise ValueError(f"Устунда {n} та карта бор — аввал уларни қайси устунга кўчиришни танланг")
            dest = next(c for c in cols if c["id"] == move_to)
            base = (con.execute("SELECT coalesce(max(position), 0) FROM issues WHERE column_id=?", (move_to,)).fetchone()[0] or 0)
            for i, r in enumerate(con.execute("SELECT id FROM issues WHERE column_id=? ORDER BY position, id", (cid,)).fetchall(), 1):
                con.execute("UPDATE issues SET column_id=?, position=?, column_at=? WHERE id=?", (move_to, base + i * 10, _now(), r[0]))
                log_issue(con, r[0], f"Устун ўчирилди → «{dest['name']}»")
            _sync_status_of_column(con, move_to, bool(dest["is_done"]))
        con.execute("DELETE FROM issue_columns WHERE id=?", (cid,))
        db.log(con, "board_column_deleted", id=cid, moved=n, to=move_to)
    return {"ok": True, "moved": n}

def reorder_columns(con, ids: list) -> dict:
    with con:
        for i, cid in enumerate(ids, 1):
            con.execute("UPDATE issue_columns SET position=? WHERE id=?", (i * 10, int(cid)))
    return {"ok": True}

def _sync_status_of_column(con, cid: int, done: bool):
    """Cards in a «done» column are closed, elsewhere open (called inside a transaction)."""
    if done:
        con.execute("UPDATE issues SET status='closed', closed_at=coalesce(closed_at, ?), updated_at=? WHERE column_id=? AND status<>'closed'", (_now(), _now(), cid))
    else:
        con.execute("UPDATE issues SET status='open', closed_at=NULL, updated_at=? WHERE column_id=? AND status<>'open'", (_now(), cid))

# ------------------------------------------------------------------ labels
def labels(con) -> list[dict]:
    used = {r[0]: r[1] for r in con.execute("SELECT label_id, count(*) FROM issue_label_links GROUP BY label_id")}
    return [{**dict(r), "used": used.get(r["id"], 0)} for r in con.execute("SELECT * FROM issue_labels ORDER BY position, lower(name)")]

def save_label(con, data: dict) -> dict:
    name = (data.get("name") or "").strip()[:60]
    if not name: raise ValueError("Белги номини ёзинг (масалан: Божхона, Солиқ, Банк)")
    color = _color(data.get("color"))
    with con:
        # SQLite lower() is ASCII-only: compare in Python so «Божхона» and «божхона» are one label
        dup = next((r for r in con.execute("SELECT id, name FROM issue_labels") if r["name"].casefold() == name.casefold()), None)
        if data.get("id"):
            lid = int(data["id"])
            if dup and dup[0] != lid: raise ValueError("Бундай белги аллақачон бор")
            con.execute("UPDATE issue_labels SET name=?, color=coalesce(?, color) WHERE id=?", (name, color, lid))
        else:
            if dup: raise ValueError("Бундай белги аллақачон бор: " + name)
            n = con.execute("SELECT count(*) FROM issue_labels").fetchone()[0]
            if not color:   # avtomatik: hali ishlatilmagan rang, hammasi band bo'lsa — navbatdagisi
                used = {r[0] for r in con.execute("SELECT color FROM issue_labels")}
                color = next((c for c in LABEL_COLORS if c not in used), LABEL_COLORS[n % len(LABEL_COLORS)])
            lid = con.execute("INSERT INTO issue_labels(name, color, position) VALUES(?,?,?)", (name, color, (n + 1) * 10)).lastrowid
        db.log(con, "board_label_saved", id=lid, name=name)
    return {"ok": True, "id": lid}

def delete_label(con, lid: int) -> dict:
    with con:
        if not con.execute("SELECT 1 FROM issue_labels WHERE id=?", (lid,)).fetchone(): raise ValueError("Белги топилмади")
        n = con.execute("SELECT count(*) FROM issue_label_links WHERE label_id=?", (lid,)).fetchone()[0]
        con.execute("DELETE FROM issue_label_links WHERE label_id=?", (lid,))
        con.execute("DELETE FROM issue_labels WHERE id=?", (lid,))
        db.log(con, "board_label_deleted", id=lid, unlinked=n)
    return {"ok": True, "unlinked": n}

def set_issue_labels(con, iid: int, label_ids) -> None:
    """Inside a transaction: replace the card's label set."""
    ids = sorted({int(x) for x in (label_ids or []) if str(x).strip()})
    con.execute("DELETE FROM issue_label_links WHERE issue_id=?", (iid,))
    for lid in ids:
        if con.execute("SELECT 1 FROM issue_labels WHERE id=?", (lid,)).fetchone():
            con.execute("INSERT OR IGNORE INTO issue_label_links(issue_id, label_id) VALUES(?,?)", (iid, lid))

# ------------------------------------------------------------------ enrich issue rows (labels, files)
def enrich(con, rows: list[dict]) -> list[dict]:
    if not rows: return rows
    ids = [r["id"] for r in rows]
    ph = ",".join("?" * len(ids))
    lab = defaultdict(list)
    for r in con.execute(f"SELECT l.issue_id, l.label_id FROM issue_label_links l JOIN issue_labels x ON x.id=l.label_id WHERE l.issue_id IN ({ph}) ORDER BY x.position, lower(x.name)", ids):
        lab[r[0]].append(r[1])
    files = defaultdict(list)
    for r in con.execute(f"SELECT id, issue_id, filename, size, mime, uploaded_at FROM issue_files WHERE deleted=0 AND issue_id IN ({ph}) ORDER BY id", ids):
        d = dict(r); d["is_image"] = Path(d["filename"]).suffix.lower() in IMAGE_EXT; files[r["issue_id"]].append(d)
    checks = defaultdict(list)
    for r in con.execute(f"SELECT id, issue_id, text, done, done_at FROM issue_checks WHERE issue_id IN ({ph}) ORDER BY position, id", ids):
        checks[r["issue_id"]].append(dict(r))
    cstat = {r[0]: (r[1] or 0, r[2]) for r in con.execute(f"SELECT issue_id, sum(kind='note'), max(at) FROM issue_comments WHERE issue_id IN ({ph}) GROUP BY issue_id", ids)}
    last_note = {r["issue_id"]: {"text": r["text"], "at": r["at"]} for r in con.execute(
        f"SELECT issue_id, text, at FROM issue_comments WHERE kind='note' AND issue_id IN ({ph}) ORDER BY at, id", ids)}
    today = dt.date.today()
    def days(v):
        try: return (today - dt.date.fromisoformat(str(v)[:10])).days
        except (TypeError, ValueError): return None
    for r in rows:
        r["labels"] = lab.get(r["id"], []); r["files"] = files.get(r["id"], []); r["checks"] = checks.get(r["id"], [])
        n_notes, last_c = cstat.get(r["id"], (0, None))
        r["n_notes"] = n_notes; r["last_note"] = last_note.get(r["id"])
        r["priority"] = r.get("priority") or "normal"
        seen = [x for x in (r.get("created_at"), r.get("updated_at"), r.get("column_at"), last_c) if x]
        r["last_at"] = max(seen) if seen else None
        r["idle_days"] = days(r["last_at"]); r["col_days"] = days(r.get("column_at") or r.get("created_at"))
        r["age_days"] = days(r.get("created_at"))
        r["stale"] = bool(r["status"] == "open" and (r["idle_days"] or 0) >= STALE_DAYS)
    return rows

def board(con) -> dict:
    from .companies import issues_list
    rows = issues_list(con, {"status": "all"})
    return {"columns": columns(con), "labels": labels(con), "issues": rows, "today": dt.date.today().isoformat(),
            "priority": PRIORITY, "basis": BASIS, "stale_days": STALE_DAYS}

# ------------------------------------------------------------------ move / delete
def move_issue(con, iid: int, column_id: int, position: float | None) -> dict:
    col = con.execute("SELECT * FROM issue_columns WHERE id=?", (int(column_id),)).fetchone()
    if not col: raise ValueError("Устун топилмади")
    old = con.execute("SELECT * FROM issues WHERE id=?", (int(iid),)).fetchone()
    if not old: raise ValueError("Карта топилмади")
    if position is None:
        position = (con.execute("SELECT coalesce(max(position), 0) FROM issues WHERE column_id=?", (col["id"],)).fetchone()[0] or 0) + 10
    moved = old["column_id"] != col["id"]
    with con:
        con.execute("UPDATE issues SET column_id=?, position=?, updated_at=? WHERE id=?", (col["id"], float(position), _now(), old["id"]))
        if moved: con.execute("UPDATE issues SET column_at=? WHERE id=?", (_now(), old["id"]))
        status = "closed" if col["is_done"] else "open"
        if status != old["status"]:
            con.execute("UPDATE issues SET status=?, closed_at=? WHERE id=?", (status, _now() if status == "closed" else None, old["id"]))
            from .companies import KIND_UZ
            db.event(con, old["inn"], "issue", f"{KIND_UZ.get(old['kind'], old['kind'])} {'ҳал қилинди' if status == 'closed' else 'қайта очилди'}: {old['title']}")
        if moved:
            db.event(con, old["inn"], "issue", f"«{old['title']}» → устун «{col['name']}»")
            src = con.execute("SELECT name FROM issue_columns WHERE id=?", (old["column_id"],)).fetchone()
            log_issue(con, old["id"], f"«{src['name'] if src else '—'}» → «{col['name']}»" + (" · ҳал қилинди" if status == "closed" and old["status"] != "closed" else " · қайта очилди" if status == "open" and old["status"] == "closed" else ""))
        _renumber(con, col["id"])
    return {"ok": True, "status": status}

def _renumber(con, cid: int):
    """Positions stay small integers so repeated drags never lose precision."""
    for i, r in enumerate(con.execute("SELECT id FROM issues WHERE column_id=? ORDER BY position, id", (cid,)).fetchall(), 1):
        con.execute("UPDATE issues SET position=? WHERE id=?", (i * 10, r[0]))

def delete_issue(con, iid: int) -> dict:
    old = con.execute("SELECT * FROM issues WHERE id=?", (int(iid),)).fetchone()
    if not old: raise ValueError("Карта топилмади")
    db.backup_if_stale("before_issue_delete")
    with con:
        for f in con.execute("SELECT * FROM issue_files WHERE issue_id=? AND deleted=0", (old["id"],)).fetchall():
            _trash(f)
        con.execute("UPDATE issue_files SET deleted=1 WHERE issue_id=?", (old["id"],))
        con.execute("DELETE FROM issue_label_links WHERE issue_id=?", (old["id"],))
        con.execute("DELETE FROM issue_checks WHERE issue_id=?", (old["id"],))
        con.execute("DELETE FROM issue_comments WHERE issue_id=?", (old["id"],))
        con.execute("DELETE FROM issues WHERE id=?", (old["id"],))
        db.event(con, old["inn"], "issue", f"Карта ўчирилди: {old['title']}")
        db.log(con, "issue_deleted", id=old["id"], inn=old["inn"], title=old["title"])
    return {"ok": True}

# ------------------------------------------------------------------ timeline (izohlar + tarix)
def log_issue(con, iid: int, text: str, kind: str = "log") -> None:
    """Inside a transaction: one line of the card's timeline (automatic history)."""
    con.execute("INSERT INTO issue_comments(issue_id, kind, text, at) VALUES(?,?,?,?)", (int(iid), kind, text[:500], _now()))

def _touch(con, iid: int):
    con.execute("UPDATE issues SET updated_at=? WHERE id=?", (_now(), int(iid)))

def timeline(con, iid: int) -> list[dict]:
    return [dict(r) for r in con.execute("SELECT * FROM issue_comments WHERE issue_id=? ORDER BY at, id", (int(iid),))]

def save_comment(con, data: dict) -> dict:
    text = (data.get("text") or "").strip()[:4000]
    if not text: raise ValueError("Изоҳ матнини ёзинг")
    with con:
        if data.get("id"):
            r = con.execute("SELECT * FROM issue_comments WHERE id=?", (int(data["id"]),)).fetchone()
            if not r or r["kind"] != "note": raise ValueError("Изоҳ топилмади")
            con.execute("UPDATE issue_comments SET text=?, edited_at=? WHERE id=?", (text, _now(), r["id"])); iid = r["issue_id"]; cid = r["id"]
        else:
            iid = int(data.get("issue_id") or 0)
            if not con.execute("SELECT 1 FROM issues WHERE id=?", (iid,)).fetchone(): raise ValueError("Аввал картани сақланг")
            cid = con.execute("INSERT INTO issue_comments(issue_id, kind, text, at) VALUES(?,?,?,?)", (iid, "note", text, _now())).lastrowid
        _touch(con, iid)
    return {"ok": True, "id": cid}

def delete_comment(con, cid: int) -> dict:
    r = con.execute("SELECT * FROM issue_comments WHERE id=?", (int(cid),)).fetchone()
    if not r or r["kind"] != "note": raise ValueError("Изоҳ топилмади (тизим ёзувлари ўчирилмайди)")
    with con: con.execute("DELETE FROM issue_comments WHERE id=?", (r["id"],))
    return {"ok": True}

# ------------------------------------------------------------------ checklist (қадамлар)
def save_check(con, data: dict) -> dict:
    with con:
        if data.get("id"):
            r = con.execute("SELECT * FROM issue_checks WHERE id=?", (int(data["id"]),)).fetchone()
            if not r: raise ValueError("Қадам топилмади")
            text = (data.get("text") if data.get("text") is not None else r["text"]).strip()[:300]
            if not text: raise ValueError("Қадам матнини ёзинг")
            done = (1 if data["done"] else 0) if "done" in data else r["done"]
            done_at = (r["done_at"] or _now()) if done else None
            con.execute("UPDATE issue_checks SET text=?, done=?, done_at=? WHERE id=?", (text, done, done_at, r["id"]))
            if done and not r["done"]: log_issue(con, r["issue_id"], f"✓ Қадам бажарилди: {text}")
            elif r["done"] and not done: log_issue(con, r["issue_id"], f"Қадам қайта очилди: {text}")
            _touch(con, r["issue_id"]); cid = r["id"]
        else:
            iid = int(data.get("issue_id") or 0)
            if not con.execute("SELECT 1 FROM issues WHERE id=?", (iid,)).fetchone(): raise ValueError("Аввал картани сақланг")
            texts = [t.strip()[:300] for t in str(data.get("text") or "").splitlines() if t.strip()]   # several lines pasted = several steps
            if not texts: raise ValueError("Қадам матнини ёзинг")
            pos = con.execute("SELECT coalesce(max(position), 0) FROM issue_checks WHERE issue_id=?", (iid,)).fetchone()[0] or 0
            for i, t in enumerate(texts, 1):
                cid = con.execute("INSERT INTO issue_checks(issue_id, text, position) VALUES(?,?,?)", (iid, t, pos + i * 10)).lastrowid
            _touch(con, iid)
    return {"ok": True, "id": cid}

def delete_check(con, cid: int) -> dict:
    r = con.execute("SELECT * FROM issue_checks WHERE id=?", (int(cid),)).fetchone()
    if not r: raise ValueError("Қадам топилмади")
    with con:
        con.execute("DELETE FROM issue_checks WHERE id=?", (r["id"],)); _touch(con, r["issue_id"])
    return {"ok": True}

def reorder_checks(con, iid: int, ids: list) -> dict:
    with con:
        for i, cid in enumerate(ids, 1):
            con.execute("UPDATE issue_checks SET position=? WHERE id=? AND issue_id=?", (i * 10, int(cid), int(iid)))
    return {"ok": True}

# ------------------------------------------------------------------ company card (modal side): contacts + this year's export
def company_brief(con, inn: str) -> dict:
    from .companies import _monthly, year, _contact_row, persons_list
    from . import report
    c = con.execute("SELECT inn, name, district_code, tarmoq FROM companies WHERE inn=?", (str(inn),)).fetchone()
    if not c: raise ValueError("Корхона топилмади")
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    m = _monthly(con, c["inn"]).get(c["inn"]) or {"m": [0.0] * 12, "last": None, "gtd": 0}
    ct = _contact_row(con.execute("SELECT * FROM company_contacts WHERE inn=?", (c["inn"],)).fetchone())
    return {"inn": c["inn"], "name": c["name"], "district": dn.get(c["district_code"]), "tarmoq": c["tarmoq"], "year": year(con),
            "ytd": round(sum(m["m"]), 3), "last": m.get("last"), "director": ct.get("director"), "phones": ct.get("phones") or [],
            "persons": [{"name": p["name"], "role": p["role"], "phones": p["phones"]} for p in persons_list(con, c["inn"])][:4],
            "open": con.execute("SELECT count(*) FROM issues WHERE inn=? AND status='open'", (c["inn"],)).fetchone()[0]}

# ------------------------------------------------------------------ attachments
def files_dir() -> Path:
    p = db.BASE / "issue_files"; p.mkdir(parents=True, exist_ok=True); return p

def add_file(con, iid: int, filename: str, data: bytes) -> dict:
    if not con.execute("SELECT 1 FROM issues WHERE id=?", (int(iid),)).fetchone(): raise ValueError("Аввал картани сақланг")
    safe = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", Path(filename).name)[:150] or "fayl"
    mime = mimetypes.guess_type(safe)[0] or "application/octet-stream"
    with con:
        cur = con.execute("INSERT INTO issue_files(issue_id, filename, stored, size, mime) VALUES(?,?,?,?,?)", (int(iid), safe, "", len(data), mime))
        fid = cur.lastrowid; rel = f"{int(iid)}/{fid}__{safe}"
        (files_dir() / str(int(iid))).mkdir(parents=True, exist_ok=True); (files_dir() / rel).write_bytes(data)
        con.execute("UPDATE issue_files SET stored=? WHERE id=?", (rel, fid))
        inn = con.execute("SELECT inn FROM issues WHERE id=?", (int(iid),)).fetchone()[0]
        db.event(con, inn, "issue", f"Картага файл бириктирилди: {safe}")
        log_issue(con, int(iid), f"Файл бириктирилди: {safe}")
    return {"ok": True, "id": fid, "filename": safe, "size": len(data), "mime": mime, "is_image": Path(safe).suffix.lower() in IMAGE_EXT}

def file_path(con, fid: int):
    r = con.execute("SELECT * FROM issue_files WHERE id=? AND deleted=0", (int(fid),)).fetchone()
    if not r: raise ValueError("Файл топилмади")
    p = files_dir() / r["stored"]
    if not p.exists(): raise ValueError("Файл дискда топилмади")
    return p, r["filename"]

def _trash(r):
    trash = files_dir() / "_olib_tashlangan"; trash.mkdir(exist_ok=True)
    src = files_dir() / r["stored"]
    if src.exists(): shutil.move(str(src), str(trash / f"{r['issue_id']}__{Path(r['stored']).name}"))

def remove_file(con, fid: int) -> dict:
    r = con.execute("SELECT * FROM issue_files WHERE id=? AND deleted=0", (int(fid),)).fetchone()
    if not r: raise ValueError("Файл топилмади")
    with con:
        _trash(r)
        con.execute("UPDATE issue_files SET deleted=1 WHERE id=?", (r["id"],))
        inn = con.execute("SELECT inn FROM issues WHERE id=?", (r["issue_id"],)).fetchone()
        db.event(con, inn[0] if inn else None, "issue", f"Картадан файл олиб ташланди: {r['filename']}")
        log_issue(con, r["issue_id"], f"Файл олиб ташланди: {r['filename']}")
    return {"ok": True}

# ------------------------------------------------------------------ svod (Excel)
def export_xlsx(con, out: Path) -> Path:
    import openpyxl
    from openpyxl.utils import get_column_letter as L
    from openpyxl.styles import PatternFill, Font
    from .companies import _styles, _clean, KIND_UZ
    from . import report
    st = _styles(); b = board(con); cols = b["columns"]; labs = b["labels"]; rows = b["issues"]
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    lname = {l["id"]: l["name"] for l in labs}; cname = {c["id"]: c["name"] for c in cols}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Свод"
    ws["A1"] = "Муаммо ва вазифалар — свод"; ws["A1"].font = st["title"]
    ws["A2"] = f"{dt.date.today():%d.%m.%Y} ҳолатига · карталар сони"; ws["A2"].font = st["sub"]

    def block(r0, title, keys, keyname, pick):
        ws.cell(r0, 1, title).font = st["sec"]
        hdr = [keyname, *[c["name"] for c in cols], "Жами", "Муддати ўтган"]
        for i, h in enumerate(hdr, 1):
            c = ws.cell(r0 + 1, i, h); c.font = st["h"]; c.fill = st["hf"]; c.alignment = st["c"]; c.border = st["b"]
        r = r0 + 2
        tot = [0] * (len(cols) + 2)
        for key, name in keys:
            sel = [x for x in rows if pick(x, key)]
            vals = [sum(1 for x in sel if x["column_id"] == c["id"]) for c in cols]
            line = [name, *vals, len(sel), sum(1 for x in sel if x["overdue"])]
            for i, v in enumerate(line, 1):
                c = ws.cell(r, i, _clean(v) if i == 1 else (v or None)); c.border = st["b"]
                if i > 1: c.alignment = st["c"]
            for i in range(len(cols) + 2): tot[i] += line[i + 1]
            r += 1
        for i, v in enumerate(["Жами", *tot], 1):
            c = ws.cell(r, i, v or (0 if i > 1 else v)); c.font = st["bold"]; c.fill = st["tot"]; c.border = st["b"]
            if i > 1: c.alignment = st["c"]
        return r + 2

    r = 4
    r = block(r, "Ташкилотлар (белгилар) кесимида", [(l["id"], l["name"]) for l in labs] + [(None, "Белгисиз")], "Ташкилот",
              lambda x, k: (k in x["labels"]) if k is not None else not x["labels"])
    r = block(r, "Устуворлик кесимида", list(PRIORITY.items()), "Устуворлик", lambda x, k: x["priority"] == k)
    dkeys = [(code, name) for code, name in sorted(dn.items(), key=lambda kv: kv[1])] + [(None, "Умумий (корхонасиз)")]
    r = block(r, "Туманлар кесимида", dkeys, "Туман", lambda x, k: (x["district_code"] == k) if k else not x["inn"])
    ws.column_dimensions["A"].width = 34
    for i in range(2, len(cols) + 4): ws.column_dimensions[L(i)].width = 14

    ws2 = wb.create_sheet("Рўйхат")
    hdr = ["№", "Карта", "Устун", "Устуворлик", "Қисқа мазмуни", "Корхона", "ИНН", "Туман", "Ташкилотлар", "Асос", "Масъул", "Муддат", "Ҳолат", "Ёзилган",
           "Ҳал қилинган", "Қадамлар", "Ҳаракатсиз (кун)", "Файллар", "Батафсил", "Охирги изоҳ", "Натижа"]
    widths = [5, 7, 14, 11, 40, 30, 12, 16, 26, 30, 16, 11, 13, 11, 12, 10, 11, 8, 50, 40, 40]
    DUE, STATE = 12, 13
    for i, (h, w) in enumerate(zip(hdr, widths), 1):
        c = ws2.cell(1, i, h); c.font = st["h"]; c.fill = st["hf"]; c.alignment = st["c"]; c.border = st["b"]; ws2.column_dimensions[L(i)].width = w
    order = {c["id"]: i for i, c in enumerate(cols)}
    for n, x in enumerate(sorted(rows, key=lambda x: (order.get(x["column_id"], 99), x["position"] or 0)), 2):
        basis = " · ".join(v for v in (BASIS.get(x.get("basis_kind")), x.get("basis_no"), _dmy(x.get("basis_date"))) if v) if x.get("basis_kind") or x.get("basis_no") else None
        ck = x["checks"]; ln = x.get("last_note")
        vals = [n - 1, f"#{x['id']}", cname.get(x["column_id"], "—"), PRIORITY.get(x["priority"], ""), x["title"], x.get("company") or ("умумий" if not x["inn"] else ""), x["inn"],
                dn.get(x.get("district_code"), ""), ", ".join(lname.get(i, "") for i in x["labels"]), basis, x["responsible"],
                dt.date.fromisoformat(x["due_date"]) if x["due_date"] else None,
                "ҳал қилинган" if x["status"] == "closed" else ("муддати ўтган" if x["overdue"] else "туриб қолган" if x["stale"] else "очиқ"),
                (x["created_at"] or "")[:10], (x["closed_at"] or "")[:10], f"{sum(1 for c in ck if c['done'])}/{len(ck)}" if ck else None,
                x["idle_days"] if x["status"] == "open" else None, len(x["files"]) or None, x["text"],
                f"{_dmy(ln['at'][:10])}: {ln['text']}" if ln else None, x["resolution"]]
        for i, v in enumerate(vals, 1):
            c = ws2.cell(n, i, _clean(v)); c.border = st["b"]; c.alignment = st["w"]
            if i == DUE: c.number_format = "DD.MM.YYYY"
            if i == STATE and (x["overdue"] or x["stale"]): c.font = st["red"]
    ws2.freeze_panes = "A2"
    wb.save(out); return out

def _dmy(v) -> str | None:
    return f"{v[8:10]}.{v[5:7]}.{v[:4]}" if v and len(v) >= 10 else None

def _color(v) -> str | None:
    s = str(v or "").strip()
    return s if re.fullmatch(r"#[0-9A-Fa-f]{6}", s) else None
