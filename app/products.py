"""Tarmoq va mahsulot daraxti: tarmoq -> sub-tarmoq -> mahsulot nomi -> HS-10 -> korxonalar.

Manba: items (bojxona bazasi, qoldiq sanasigacha) + keyingi kunlik GTD (gtd_rows). Har tugun korxona INN'lariga ulanadi.
Qoidalar dashboard bilan bir xil: «Бухоро эмас» (bukhara=0) hech qayerga, Buxoro davri sana bo'yicha, БНПЗ/SDK (excluded)
faqat so'ralganda (memo=1), voided GTD hech qachon. Sanoat — items/GTD kind='sanoat'. Meva-sabzavot — HAMMA joyda faqat KARANTIN qoidasi
(report.karantin bilan bir xil: karantin_opening svod + GTD kind='meva' o'stirilgan tuman bo'yicha, eksportyor kim bo'lishidan qat'i nazar);
items dagi kind='meva' qatorlari jamiga kirmaydi. «Ўтган йил» sarlavha/tuman uchun — faqat Sozlamalardagi baza (report.prev_year_amount).
"""
from __future__ import annotations
import datetime as dt, re
from collections import defaultdict
from . import db, report
from .companies import country_uz

OTHER = "Аниқланмаган"          # HS kodi lug'atda yo'q (faqat GTD da uchragan kod)
MEVA_T1 = "Мева-сабзавот махсулотлари"
MEVA_T1_S = "Мева-сабзавот махсулотлари (саноат ҳисобида)"   # kind=sanoat qatori, lekin HS lug'atida meva guruhi — saноат jamisida qoladi

def _vals(q, key):
    return {x.strip() for x in str(q.get(key) or "").split(",") if x.strip()}

def _period(con, q):
    """(year, from, to): joriy yil — oxirgi ma'lumot kunigacha; o'tgan yil — yil oxirigacha yoki ?to= gacha."""
    cur = int(db.get_setting(con, "year") or dt.date.today().year)
    y = int(q.get("year") or cur)
    last = report.last_data_date(con) or dt.date.today().isoformat()
    f = (q.get("from") or "").strip() or f"{y}-01-01"
    t = (q.get("to") or "").strip() or (last if y == cur else f"{y}-12-31")
    if not f.startswith(str(y)): f = f"{y}-01-01"
    if not t.startswith(str(y)): t = f"{y}-12-31"
    return y, f[:10], min(t[:10], f"{y}-12-31")

def _hs_dict(con) -> dict:
    return {r["hs10"]: dict(r) for r in con.execute("SELECT hs10, tovar1, tovar2, sprav1, sprav2, unit, name FROM hs_codes")}

def _hs4_dict(con) -> dict:
    """hs4 -> lug'at qatori (eng ko'p uchragan hs10 bo'yicha) — GTD dagi kod lug'atda bo'lmasa 4 xonali bo'yicha."""
    out = {}
    for r in con.execute("SELECT substr(hs10,1,4) k, tovar1, tovar2, sprav1, sprav2, count(*) n FROM hs_codes GROUP BY k, tovar1, tovar2, sprav1, sprav2 ORDER BY n"):
        out[r["k"]] = dict(r)
    return out

def _rows(con, y: int, f: str, t: str, comp: dict, keep, hs: dict, hs4: dict, with_gtd=True, inn: str | None = None):
    """Bir yil uchun tovar qatorlari: (inn, hs10, tovar1, tovar2, sprav1, sprav2, value, netto_t, gtd_key, country, src, rdate, kind sanoat|meva)."""
    through = db.get_setting(con, "opening_through_date") or ""
    out = []
    # «Йил боши манбаси = Шаблон»: those companies' sanoat items are replaced by the imported template months (dispute.template_items)
    from . import dispute
    tset = dispute.template_inns(con) if y == int(db.get_setting(con, "year") or y) else set()
    src_rows = con.execute("""SELECT i.inn, i.rdate, i.gtd, i.hs10, i.hs4, i.tovar1, i.tovar2, i.sprav1, i.sprav2, i.value, i.netto, i.kind, d.country
                            FROM items i JOIN declarations d ON d.decl_id=i.decl_id
                            WHERE i.regime='ЭК' AND i.year=? AND i.rdate>=? AND i.rdate<=?""" + (" AND i.inn=?" if inn else ""), (y, f, t, *((inn,) if inn else ())))
    for r in (list(src_rows) + (dispute.template_items(con, y, f, t, inn) if tset else [])):
        if r["inn"] in tset and r["kind"] == "sanoat" and not str(r["gtd"] or "").startswith("tpl:"): continue
        c = comp.get(r["inn"])
        if not keep(c, r["rdate"], r["kind"]): continue
        code = r["hs10"] or ""
        d = hs.get(code) or hs4.get(code[:4]) or {}
        t1 = r["tovar1"] or d.get("tovar1") or (MEVA_T1 if r["kind"] == "meva" else OTHER)
        if t1 == MEVA_T1 and r["kind"] != "meva": t1 = MEVA_T1_S
        out.append((r["inn"], code, t1, r["tovar2"] or d.get("tovar2") or code[:4] or "—", r["sprav1"] or d.get("sprav1") or t1,
                    r["sprav2"] or d.get("sprav2") or (code[:6] or "—"), r["value"] or 0.0, r["netto"] or 0.0, r["gtd"], country_uz(r["country"]), "baza", r["rdate"], "meva" if r["kind"] == "meva" else "sanoat"))
    if not with_gtd: return out
    # X1: davrning birinchi kuni tushib qolmasin — ikki alohida shart: report_date >= f VA report_date > through (opening bo'lsa)
    for r in con.execute("""SELECT inn, report_date d, gtd_no, tnved, product, country, netto, stat_usd, kind FROM gtd_rows
                            WHERE voided=0 AND kind IN ('sanoat','meva','meva_x','bnpz','memo') AND report_date>=? AND report_date<=?"""
                         + (" AND report_date>?" if through else "") + (" AND inn=?" if inn else ""), (f, t, *((through,) if through else ()), *((inn,) if inn else ()))):
        c = comp.get(r["inn"])
        if not keep(c, r["d"], r["kind"]): continue
        code = str(r["tnved"] or "").strip()
        d = hs.get(code) or hs4.get(code[:4]) or {}
        t1 = d.get("tovar1") or (MEVA_T1 if r["kind"] in ("meva", "meva_x") else OTHER)
        if t1 == MEVA_T1 and r["kind"] not in ("meva", "meva_x"): t1 = MEVA_T1_S
        out.append((r["inn"], code, t1, d.get("tovar2") or code[:4] or "—", d.get("sprav1") or t1, d.get("sprav2") or (code[:6] or "—"),
                    r["stat_usd"] or 0.0, (r["netto"] or 0.0) / 1000, r["gtd_no"], country_uz(r["country"]), "gtd", r["d"], "meva" if r["kind"] in ("meva", "meva_x") else "sanoat"))
    return out

def _keeper(con, q, comp, meva_inns):
    from .daily import is_bukhara
    dists, kinds = _vals(q, "district"), _vals(q, "kind") or {"sanoat", "meva"}
    with_memo = str(q.get("memo") or "0") == "1"
    def keep(c, d, kind):
        if c is None: return False
        if c.get("bukhara") == 0: return False
        if d and not ((c.get("bukhara_from") or "") <= d <= (c.get("bukhara_to") or "9999")): return False
        if c.get("excluded") or kind in ("bnpz", "memo"):
            if not with_memo: return False
        if kind == "meva" and not is_bukhara(c, c.get("inn"), meva_inns, d): return False     # boshqa viloyat mevasi
        k = "meva" if kind in ("meva", "meva_x") else "sanoat"
        if k not in kinds: return False
        if dists and c.get("district_code") not in dists: return False
        return True
    return keep

def sanoat_only(keep):
    """items/GTD dagi meva qatorlari tashlab yuboriladi — meva HAMMA joyda faqat karantin qoidasi bilan (karantin_* yordamchilari)."""
    return lambda c, d, k: k not in ("meva", "meva_x") and keep(c, d, k)

def prev_kind(kinds) -> str:
    """Filtr turi -> Sozlamalardagi ўтган йил базаси turi (all | sanoat | meva)."""
    ks = set(kinds or ())
    if ks == {"sanoat"}: return "sanoat"
    if ks == {"meva"}: return "meva"
    return "all"

def prev_settings(con, code: str, kind: str, f: str, t: str) -> float:
    """«Ўтган йил шу даврга» — ФАҚАТ Sozlamalardagi baza (report.prev_year_amount): davr f..t kunlariga mutanosib.
    f = 1-yanvar bo'lsa aynan prev_year_amount(t); aks holda t gacha va f dan oldingi kungacha bo'lgan qiymatlar farqi."""
    v = report.prev_year_amount(con, code, kind, t)
    if f > f"{t[:4]}-01-01":
        d0 = (dt.date.fromisoformat(f) - dt.timedelta(days=1)).isoformat()
        v -= report.prev_year_amount(con, code, kind, d0)
    return v

def prev_full_settings(con, code: str, kind: str, y: int) -> float:
    """«Ўтган йил жами» — Sozlamalar → Даврий натижалар (12 ойлик) бўлса шундан; бўлмаса эски база (prev_year_base)."""
    from . import periods
    v = periods.full_prev(con, y, code, kind)
    if v is not None: return float(v)
    r = con.execute("SELECT prev_year_base FROM prognoz WHERE year=? AND district_code=? AND kind=?", (y, code, kind)).fetchone()
    return float(r["prev_year_base"] or 0) if r else 0.0

# ------------------------------------------------------------------ meva-sabzavot: KARANTIN qoidasi (report.karantin bilan bir xil)
KAR_LABEL = "Мева-сабзавот (карантин рўйхати, база даври)"
KAR_ROW = "Мева-сабзавот (карантин рўйхати)"

def _opening(con) -> str: return db.get_setting(con, "opening_through_date") or ""

def karantin_base(con, y: int, f: str, t: str) -> dict:
    """{district_code: meva} — karantin_opening (bazadagi svod) ning f..t davriga tegishli qismi. Yil boshidan (f=01.01, t>=opening)
    bo'lsa ytd; opening oyi davr ichida bo'lsa shu oyning month_amt; boshqa davrlar uchun bazada oy tafsiloti yo'q -> 0."""
    through = _opening(con); out = defaultdict(float)
    if not through or not through.startswith(str(y)): return out
    om = int(through[5:7]); m0, m1 = f"{y}-{om:02d}-01", f"{y}-{om:02d}-31"
    ytd_mode = f <= f"{y}-01-01" and t >= through
    for r in con.execute("SELECT district_code, ytd, month, month_amt FROM karantin_opening WHERE year=?", (y,)):
        if ytd_mode: out[r["district_code"]] += r["ytd"] or 0
        elif r["month"] and f <= m1 and t >= m0: out[r["district_code"]] += r["month_amt"] or 0
    return out

def karantin_gtd(con, y: int, f: str, t: str) -> list:
    """Opening sanasidan keyingi GTD meva qatorlari (kind='meva', o'stirilgan tuman bo'yicha, eksportyor reyestrda bo'lishi shart emas)."""
    through = _opening(con)
    return con.execute("""SELECT inn, exporter, report_date d, gtd_no, tnved, product, country, netto, stat_usd, district_code dc FROM gtd_rows
                          WHERE voided=0 AND kind='meva' AND report_date>=? AND report_date<=? AND substr(report_date,1,4)=?""" + (" AND report_date>?" if through else ""),
                       (f, t, str(y), *((through,) if through else ()))).fetchall()

def karantin_period(con, y: int, f: str, t: str, dists=None) -> dict:
    """{district_code: {"base", "gtd", "total", "inns": set, "gtds": set}} — davr uchun karantin mevasi (report.karantin bilan mos)."""
    out = defaultdict(lambda: {"base": 0.0, "gtd": 0.0, "total": 0.0, "inns": set(), "gtds": set()})
    for dc, v in karantin_base(con, y, f, t).items():
        if dists and dc not in dists: continue
        out[dc]["base"] += v; out[dc]["total"] += v
    for r in karantin_gtd(con, y, f, t):
        dc = r["dc"] or "?"
        if dists and dc not in dists: continue
        out[dc]["gtd"] += r["stat_usd"] or 0; out[dc]["total"] += r["stat_usd"] or 0; out[dc]["inns"].add(r["inn"]); out[dc]["gtds"].add(r["gtd_no"])
    return out

def karantin_months(con, y: int, D: str | None = None) -> dict:
    """{district_code: [12]} — oylar bo'yicha karantin mevasi (report.karantin_monthly kabi: opening oyi + GTD oylari; oldingi oylar noma'lum -> 0)."""
    through = _opening(con); out = defaultdict(lambda: [0.0] * 12)
    if through and through.startswith(str(y)):
        om = int(through[5:7])
        for r in con.execute("SELECT district_code, month_amt FROM karantin_opening WHERE year=?", (y,)): out[r["district_code"]][om - 1] += r["month_amt"] or 0
    for r in karantin_gtd(con, y, f"{y}-01-01", D or f"{y}-12-31"): out[r["dc"] or "?"][int(r["d"][5:7]) - 1] += r["stat_usd"] or 0
    return out

def karantin_rows(con, y: int, f: str, t: str, hs: dict, hs4: dict, dn: dict, dists=None, names: dict | None = None) -> list:
    """_rows formatidagi meva qatorlari: GTD meva (haqiqiy mahsulot/davlat bilan) + har tuman uchun BITTA sintetik «база даври» qatori."""
    out = []
    for r in karantin_gtd(con, y, f, t):
        dc = r["dc"] or "?"
        if dists and dc not in dists: continue
        if names is not None and r["inn"] not in names: names[r["inn"]] = {"name": r["exporter"] or r["inn"], "district_code": dc}
        code = str(r["tnved"] or "").strip(); d = hs.get(code) or hs4.get(code[:4]) or {}
        out.append((r["inn"], code, MEVA_T1, d.get("tovar2") or code[:4] or "—", d.get("sprav1") or MEVA_T1, d.get("sprav2") or (code[:6] or "—"),
                    r["stat_usd"] or 0.0, (r["netto"] or 0.0) / 1000, r["gtd_no"], country_uz(r["country"]), "gtd", r["d"], "meva"))
    for dc, v in karantin_base(con, y, f, t).items():
        if (dists and dc not in dists) or not v: continue
        key = "karantin:" + dc; nm = dn.get(dc, dc)
        if names is not None: names[key] = {"name": nm, "district_code": dc}
        out.append((key, nm, MEVA_T1, KAR_LABEL, MEVA_T1, KAR_LABEL, v, 0.0, None, "", "baza", _opening(con), "meva"))
    return out

def _levels(cls: str):
    # A: tarmoq -> sub-tarmoq -> mahsulot nomi -> HS-10 ; B: mahsulot guruhi -> mahsulot nomi -> HS-10
    return (2, 3, 5, 1) if cls != "B" else (4, 5, 1)

def _tree(rows, levels):
    """Ichma-ich lug'at: key -> {v, n(netto), inns:set, gtd:set, ch:{...}}; oxirgi daraja (HS-10) da korxonalar."""
    root = {"v": 0.0, "n": 0.0, "inns": set(), "gtd": set(), "ch": {}}
    def cnt(node, r):   # sintetik karantin qatori korxona/GTD soniga kirmaydi
        node["v"] += r[6]; node["n"] += r[7]
        if not str(r[0]).startswith("karantin:"): node["inns"].add(r[0])
        if r[8]: node["gtd"].add(r[8])
    for r in rows:
        node = root; cnt(node, r)
        for li, idx in enumerate(levels):
            key = r[idx] or "—"
            node = node["ch"].setdefault(key, {"v": 0.0, "n": 0.0, "inns": set(), "gtd": set(), "ch": {}, "co": {}})
            cnt(node, r)
            if li == len(levels) - 1:
                co = node["co"].setdefault(r[0], {"v": 0.0, "n": 0.0, "gtd": set(), "countries": set()})
                co["v"] += r[6]; co["n"] += r[7]; co["gtd"].add(r[8])
                if r[9]: co["countries"].add(r[9])
    return root

def _walk(node, prev, comp, dn, labels, depth, max_items=None):
    out = []
    for key, ch in node["ch"].items():
        p = (prev or {}).get("ch", {}).get(key) if prev else None
        item = {"key": key, "name": labels.get(key, key) if depth == "hs" else key, "value": round(ch["v"], 3), "netto_t": round(ch["n"], 1),
                "companies": len(ch["inns"]), "gtd": len(ch["gtd"]), "prev": round(p["v"], 3) if p else 0.0}
        if ch["ch"]:
            item["children"] = _walk(ch, p, comp, dn, labels, depth)
        else:
            cos = []
            for inn, co in ch["co"].items():
                c = comp.get(inn) or {}
                pc = (p or {}).get("co", {}).get(inn) if p else None
                cos.append({"inn": inn, "name": c.get("name") or inn, "district": dn.get(c.get("district_code")), "value": round(co["v"], 3),
                            "netto_t": round(co["n"], 1), "gtd": len(co["gtd"]), "countries": sorted(co["countries"])[:6], "prev": round(pc["v"], 3) if pc else 0.0})
            cos.sort(key=lambda x: -x["value"]); item["companies_list"] = cos
        out.append(item)
    out.sort(key=lambda x: -x["value"])
    return out

def tree(con, q: dict) -> dict:
    """q: year, cls (A|B), district (a,b), kind (sanoat,meva), memo (0/1), compare (0/1), from, to"""
    y, f, t = _period(con, q)
    cls = (q.get("cls") or "A").upper()
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, district_code, tarmoq, kind, excluded, in_base, bukhara, bukhara_from, bukhara_to, source FROM companies")}
    meva_inns = {i for (i,) in con.execute("SELECT inn FROM template_inns WHERE category='meva'")} or \
                {i for (i,) in con.execute("SELECT inn FROM companies WHERE kind='meva' AND source='nomma-nom'")}
    keep = _keeper(con, q, comp, meva_inns)
    hs, hs4 = _hs_dict(con), _hs4_dict(con)
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    levels = _levels(cls)
    kinds, dists = _vals(q, "kind") or {"sanoat", "meva"}, _vals(q, "district")
    # sanoat — items + GTD (kind sanoat / bnpz / memo); meva — FAQAT karantin qoidasi (o'stirilgan tuman), dashboard bilan bir xil
    rows = _rows(con, y, f, t, comp, sanoat_only(keep), hs, hs4)
    names = dict(comp)
    if "meva" in kinds: rows += karantin_rows(con, y, f, t, hs, hs4, dn, dists, names)
    root = _tree(rows, levels)
    prev = None; prev_total = None; prev_src = None
    if str(q.get("compare") or "1") != "0":
        pf, pt = f"{y - 1}{f[4:]}", f"{y - 1}{t[4:]}"
        if pt.endswith("02-29"): pt = pt[:-2] + "28"
        prev = _tree(_rows(con, y - 1, pf, pt, comp, keep, hs, hs4), levels)   # tugunlar bo'yicha taqqoslash: bojxona bazasi (qo'lda manba yo'q)
        # sarlavhadagi «ўтган йил» — faqat Sozlamalardagi baza (vilyoat / bitta tuman tanlanganda o'sha tuman)
        code = next(iter(dists)) if len(dists) == 1 else "1706"
        prev_total = prev_settings(con, code, prev_kind(kinds), f, t); prev_src = "settings"
    comp = names
    # HS-10 nomlari: lug'atdagi rasmiy nom, bo'lmasa GTD tovar matnidan namuna
    labels = {}
    for code, d in hs.items():
        if d.get("name"): labels[code] = d["name"]
    for r in con.execute("SELECT tnved, product, sum(stat_usd) v FROM gtd_rows WHERE voided=0 AND product IS NOT NULL GROUP BY tnved, product ORDER BY v"):
        code = str(r["tnved"] or "").strip()
        if code and not (hs.get(code) or {}).get("name"):
            txt = re.sub(r"^\s*\d+[.)]\s*", "", str(r["product"])).split(";")[0].split(":")[0].strip()
            if txt: labels[code] = txt[:90]
    items = _walk(root, prev, comp, dn, labels, "hs")
    years = [r[0] for r in con.execute("SELECT DISTINCT year FROM items WHERE regime='ЭК' ORDER BY year")]
    cur = int(db.get_setting(con, "year") or y)
    if cur not in years: years.append(cur)
    return {"year": y, "from": f, "to": t, "cls": cls, "levels": ["Тармоқ", "Соҳа", "Маҳсулот", "ТН ВЭД"] if cls != "B" else ["Маҳсулот гуруҳи", "Маҳсулот", "ТН ВЭД"],
            "total": round(root["v"], 3), "netto_t": round(root["n"], 1), "companies": len(root["inns"]), "gtd": len(root["gtd"]),
            "prev_total": round(prev_total, 3) if prev else None, "prev_source": prev_src, "prev_total_customs": round(prev["v"], 3) if prev else None,
            "items_prev_source": "customs" if prev else None, "meva_rule": "karantin", "prev_period": [f"{y - 1}{f[4:]}", f"{y - 1}{t[4:]}"] if prev else None,
            "years": sorted(years), "items": items, "rows": len(rows),
            "districts": [{"code": c, "name": n} for c, n in sorted(dn.items(), key=lambda x: x[1])],
            "hs_named": sum(1 for d in hs.values() if d.get("name")), "hs_total": len(hs)}

def company_products(con, inn: str, year: int | None = None) -> dict:
    """Bitta korxonaning mahsulot daraxti (pasport uchun): tarmoq -> sub -> mahsulot -> HS-10."""
    q = {"year": year} if year else {}
    y, f, t = _period(con, q)
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, district_code, kind, excluded, in_base, bukhara, bukhara_from, bukhara_to, source FROM companies WHERE inn=?", (inn,))}
    if not comp: return {"year": y, "items": [], "total": 0}
    keep = _keeper(con, {"memo": "1"}, comp, set())
    hs, hs4 = _hs_dict(con), _hs4_dict(con)
    rows = _rows(con, y, f, t, comp, keep, hs, hs4, inn=inn)
    root = _tree(rows, _levels("A"))
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    return {"year": y, "total": round(root["v"], 3), "items": _walk(root, None, comp, dn, {}, "hs")}

def export_xlsx(con, q: dict, out):
    """Daraxt bitta varaqda (darajalar ustunlarda) + korxonalar varag'i."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    d = tree(con, q)
    thin = Side(style="thin", color="D6DBD8"); bd = Border(thin, thin, thin, thin)
    H = Font(bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor="0F6E63")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Маҳсулотлар"
    ws["A1"] = f"Бухоро вилояти экспорти — тармоқ ва маҳсулот кесимида, {d['year']} йил"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{d['from'][8:10]}.{d['from'][5:7]}.{d['from'][:4]} — {d['to'][8:10]}.{d['to'][5:7]}.{d['to'][:4]} · минг АҚШ доллари"; ws["A2"].font = Font(italic=True, color="5A6470", size=10)
    lv = d["levels"]
    hdr = lv + ["Қиймати", "Ўтган йил", "Ўсиш %", "Нетто, т", "Корхоналар", "ГТД"]
    widths = [30] * len(lv) + [14, 14, 10, 12, 12, 8]
    for i, (h, w) in enumerate(zip(hdr, widths), 1):
        c = ws.cell(4, i, h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center"); c.border = bd
        ws.column_dimensions[L(i)].width = w
    row = [5]
    def put(items, path):
        for it in items:
            vals = path + [it["name"] if not it.get("children") else it["key"]]
            vals = vals + [""] * (len(lv) - len(vals))
            growth = (it["value"] / it["prev"] - 1) if it.get("prev") else None
            for i, v in enumerate(vals + [it["value"], it.get("prev") or 0, growth, it["netto_t"], it["companies"], it["gtd"]], 1):
                c = ws.cell(row[0], i, v); c.border = bd
                if i == len(lv) + 1 or i == len(lv) + 2 or i == len(lv) + 4: c.number_format = "#,##0.0"
                if i == len(lv) + 3: c.number_format = "0.0%"
                if len(path) == 0: c.font = Font(bold=True)
            row[0] += 1
            if it.get("children"): put(it["children"], path + [it["key"]])
    put(d["items"], [])
    ws.freeze_panes = "A5"
    ws2 = wb.create_sheet("Корхоналар")
    hdr2 = lv + ["ИНН", "Корхона", "Туман", "Қиймати", "Ўтган йил", "Нетто, т", "ГТД", "Давлатлар"]
    for i, (h, w) in enumerate(zip(hdr2, [26] * len(lv) + [12, 40, 18, 14, 14, 12, 8, 30]), 1):
        c = ws2.cell(1, i, h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center"); c.border = bd
        ws2.column_dimensions[L(i)].width = w
    r0 = [2]
    def put2(items, path):
        for it in items:
            if it.get("children"): put2(it["children"], path + [it["key"]]); continue
            for co in it.get("companies_list", []):
                for i, v in enumerate(path + [it["key"]] + [co["inn"], co["name"], co["district"], co["value"], co["prev"], co["netto_t"], co["gtd"], ", ".join(co["countries"])], 1):
                    c = ws2.cell(r0[0], i, v); c.border = bd
                    if i in (len(lv) + 4, len(lv) + 5, len(lv) + 6): c.number_format = "#,##0.0"
                r0[0] += 1
    put2(d["items"], [])
    ws2.freeze_panes = "A2"
    wb.save(out); return out


# ---------------------------------------------------------------- Динамика: oylar kesimida, tanlangan kesim bo'yicha
def dynamics(con, q: dict) -> dict:
    """q: year, by (total|district|tarmoq|country|union), district, kind, memo. Joriy yil va o'tgan yil oylar bo'yicha; ytd — kun-kuniga."""
    y, f, t = _period(con, q)
    by = (q.get("by") or "total").lower()
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, district_code, tarmoq, kind, excluded, in_base, bukhara, bukhara_from, bukhara_to, source, union_name FROM companies")}
    meva_inns = {i for (i,) in con.execute("SELECT inn FROM template_inns WHERE category='meva'")} or \
                {i for (i,) in con.execute("SELECT inn FROM companies WHERE kind='meva' AND source='nomma-nom'")}
    keep = _keeper(con, q, comp, meva_inns)
    hs, hs4 = _hs_dict(con), _hs4_dict(con)
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    def key_of(r):
        c = comp.get(r[0]) or {}
        if by == "district": return dn.get(c.get("district_code")) or "—"
        if by == "tarmoq": return r[2]
        if by == "country": return r[9] or "—"
        if by == "union": return c.get("union_name") or "Худудий ташкилотлар"
        if by == "company": return f"{c.get('name') or r[0]}"
        return "Жами"
    kinds, dists = _vals(q, "kind") or {"sanoat", "meva"}, _vals(q, "district")
    def meva_key(dc):
        if by == "district": return dn.get(dc) or "—"
        if by == "tarmoq": return MEVA_T1
        if by == "total": return "Жами"
        return KAR_ROW      # davlat / uyushma / korxona kesimida karantin mevasi bitta alohida qator
    kinn = {}       # korxona kesimida: qator -> INN (karantin qatori «Мева-сабзавот (карантин рўйхати)» — inn yo'q)
    def series(yy, f_, t_, kp):
        m = defaultdict(lambda: [0.0] * 12); inns = defaultdict(set); ytd = defaultdict(float)
        for r in _rows(con, yy, f"{yy}-01-01", f"{yy}-12-31", comp, kp, hs, hs4):
            k = key_of(r); mi = int(r[11][5:7]) - 1
            m[k][mi] += r[6]; inns[k].add(r[0])
            if by == "company": kinn.setdefault(k, r[0])
            if f_ <= r[11] <= t_: ytd[k] += r[6]
        return m, inns, ytd
    pf, pt = f"{y - 1}{f[4:]}", f"{y - 1}{t[4:]}"
    if pt.endswith("02-29"): pt = pt[:-2] + "28"
    # joriy yil: sanoat (items + GTD) + meva karantin qoidasi bilan; o'tgan yil: bojxona bazasi (qatorlar uchun boshqa manba yo'q)
    cm, ci, cy = series(y, f, t, sanoat_only(keep)); pm, pi, py = series(y - 1, pf, pt, keep)
    if "meva" in kinds:
        for dc, mm in karantin_months(con, y, t).items():
            if dists and dc not in dists: continue
            k = meva_key(dc)
            for i, v in enumerate(mm): cm[k][i] += v
        for dc, a in karantin_period(con, y, f, t, dists).items():
            cy[meva_key(dc)] += a["total"]; ci[meva_key(dc)] |= a["inns"]
    # «ўтган йил шу даврга»: vilyoat va tumanlar uchun Sozlamalardagi baza; boshqa kesimlarda bojxona bazasi (2025 items)
    pk = prev_kind(kinds); code_of = {n: c for c, n in dn.items()}
    keys = sorted(set(cm) | set(pm), key=lambda k: -(cy.get(k, 0)))
    rows = []
    for k in keys:
        src = "customs"; pv = py.get(k, 0)
        if by == "district" and k in code_of: pv = prev_settings(con, code_of[k], pk, f, t); src = "settings"
        elif by == "total": pv = prev_settings(con, "1706", pk, f, t); src = "settings"
        rows.append({"key": k, "inn": kinn.get(k) if by == "company" else None,
                     "months": [round(v, 3) for v in cm.get(k, [0.0] * 12)], "prev_months": [round(v, 3) for v in pm.get(k, [0.0] * 12)],
                     "ytd": round(cy.get(k, 0), 3), "prev_ytd": round(pv, 3), "prev_ytd_customs": round(py.get(k, 0), 3), "prev_source": src,
                     "prev_year": round(sum(pm.get(k, [0.0] * 12)), 3),
                     "companies": len(ci.get(k, ())), "prev_companies": len(pi.get(k, ()))})
    tot = [round(sum(r["months"][i] for r in rows), 3) for i in range(12)]; ptot = [round(sum(r["prev_months"][i] for r in rows), 3) for i in range(12)]
    years = [r[0] for r in con.execute("SELECT DISTINCT year FROM items WHERE regime='ЭК' ORDER BY year")]
    cur = int(db.get_setting(con, "year") or y)
    if cur not in years: years.append(cur)
    hcode = next(iter(dists)) if len(dists) == 1 else "1706"
    return {"year": y, "from": f, "to": t, "prev_from": pf, "prev_to": pt, "by": by, "rows": rows, "total_months": tot, "prev_total_months": ptot,
            "ytd": round(sum(cy.values()), 3), "prev_ytd": round(prev_settings(con, hcode, pk, f, t), 3), "prev_source": "settings",
            "prev_ytd_customs": round(sum(py.values()), 3), "prev_year": round(sum(ptot), 3), "meva_rule": "karantin",
            "years": sorted(years), "cur_month": int(t[5:7]) if y == cur else 12,
            "districts": [{"code": c, "name": n} for c, n in sorted(dn.items(), key=lambda x: x[1])]}

def export_dynamics(con, q: dict, out):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    d = dynamics(con, q)
    M = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]
    thin = Side(style="thin", color="D6DBD8"); bd = Border(thin, thin, thin, thin)
    H = Font(bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor="0F6E63")
    wb = openpyxl.Workbook()
    for sheet, yy, mk in ((f"{d['year']}", d["year"], "months"), (f"{d['year'] - 1}", d["year"] - 1, "prev_months")):
        ws = wb.active if sheet == str(d["year"]) else wb.create_sheet(); ws.title = sheet
        ws["A1"] = f"Динамика — {yy} йил · {{'total': 'жами', 'district': 'туманлар', 'tarmoq': 'тармоқлар', 'country': 'давлатлар', 'union': 'уюшмалар'}}.get(d['by'], d['by']) · минг АҚШ доллари"
        ws["A1"].font = Font(bold=True, size=13)
        hdr = ["Кесим"] + M + ["Йил жами", "Корхоналар"]
        for i, h in enumerate(hdr, 1):
            c = ws.cell(3, i, h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center"); c.border = bd
            ws.column_dimensions[L(i)].width = 34 if i == 1 else 11
        for n, r in enumerate(d["rows"], 1):
            vals = [r["key"]] + r[mk] + [sum(r[mk]), r["companies"] if mk == "months" else r["prev_companies"]]
            for i, v in enumerate(vals, 1):
                c = ws.cell(3 + n, i, v); c.border = bd
                if 2 <= i <= 14: c.number_format = "#,##0.0"
        ws.freeze_panes = "B4"
    wb.save(out); return out
