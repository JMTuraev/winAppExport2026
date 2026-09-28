"""Geografiya: eksport davlatlar kesimida (xarita uchun).
Manba: bojxona bazasi (items, qoldiq sanasigacha) + undan keyingi kunlik GTD (gtd_rows)."""
from __future__ import annotations
import datetime as dt
from collections import defaultdict
from . import db, report

# O'zbekcha nom -> ISO alpha-2 (xarita kodi). Nomlar companies.country_uz() dan keladi.
ISO = {
    "Афғонистон": "AF", "Озарбайжон": "AZ", "Арманистон": "AM", "Беларусь": "BY", "Германия": "DE", "Грузия": "GE",
    "Ҳиндистон": "IN", "Эрон": "IR", "Қозоғистон": "KZ", "Хитой": "CN", "Корея": "KR", "Қирғизистон": "KG",
    "Латвия": "LV", "Литва": "LT", "Марокаш": "MA", "Миср": "EG", "БАА": "AE", "Уммон": "OM", "Покистон": "PK",
    "Польша": "PL", "Россия": "RU", "Словакия": "SK", "АҚШ": "US", "Тожикистон": "TJ", "Туркманистон": "TM",
    "Туркия": "TR", "Украина": "UA", "Франция": "FR", "Хорватия": "HR", "Япония": "JP", "Ветнам": "VN",
    "Венгрия": "HU", "Жанубий Африка": "ZA", "Италия": "IT", "Испания": "ES", "Нидерландия": "NL",
    "Буюк Британия": "GB", "Исроил": "IL", "Саудия Арабистони": "SA", "Қатар": "QA", "Қувайт": "KW", "Ироқ": "IQ",
    "Сурия": "SY", "Мўғулистон": "MN", "Молдова": "MD", "Эстония": "EE", "Болгария": "BG", "Руминия": "RO",
    "Чехия": "CZ", "Австрия": "AT", "Швейцария": "CH", "Белгия": "BE", "Канада": "CA", "Малайзия": "MY",
    "Индонезия": "ID", "Таиланд": "TH", "Бангладеш": "BD", "Шри-Ланка": "LK", "Сербия": "RS", "Греция": "GR",
    "Ўзбекистон": "UZ", "Швеция": "SE", "Норвегия": "NO", "Финляндия": "FI", "Дания": "DK", "Португалия": "PT",
    "Словения": "SI", "Албания": "AL", "Кипр": "CY", "Иордания": "JO", "Ливан": "LB", "Туркия ": "TR",
    "Сингапур": "SG", "Филиппин": "PH", "Австралия": "AU", "Бразилия": "BR", "Аргентина": "AR", "Мексика": "MX",
}
KIND_SKIP = ("not_bukhara",)          # «Бухоро эмас» belgilangan qatorlar hech qayerga kirmaydi

def _period(con, q) -> tuple[str, str]:
    y = int(db.get_setting(con, "year") or dt.date.today().year)
    last = report.last_data_date(con) or dt.date.today().isoformat()
    f = (q.get("from") or "").strip() or f"{y}-01-01"
    t = (q.get("to") or "").strip() or last
    return (f[:10], t[:10]) if f <= t else (t[:10], f[:10])

def _vals(q, key):
    return {x.strip() for x in str(q.get(key) or "").split(",") if x.strip()}

def rows(con, q: dict) -> dict:
    """Davlatlar kesimi. q: from,to,district,tarmoq,kind,country,memo(0/1)"""
    from .companies import country_uz, has_base_detail
    f, t = _period(con, q)
    through = db.get_setting(con, "opening_through_date") or ""
    dists, tarmoqs, kinds, countries = _vals(q, "district"), _vals(q, "tarmoq"), _vals(q, "kind"), _vals(q, "country")
    with_memo = str(q.get("memo") or "1") != "0"
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, district_code, tarmoq, kind, excluded, bukhara, bukhara_from, bukhara_to FROM companies")}
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    from . import dispute
    tset = dispute.template_inns(con)      # «Йил боши манбаси = Шаблон»: items skipped, template months added (davlat: shablon)

    def keep(inn, c, d=None):
        if c is None: return False
        if c.get("bukhara") == 0: return False               # «Бухоро эмас»: base rows and GTD alike
        if d and not ((c.get("bukhara_from") or "") <= d <= (c.get("bukhara_to") or "9999")): return False   # outside the Bukhara period
        if c.get("excluded") and not with_memo: return False
        if dists and c.get("district_code") not in dists: return False
        if tarmoqs and (c.get("tarmoq") or "").strip() not in tarmoqs: return False
        if kinds and c.get("kind") not in kinds: return False
        return True

    def collect(f, t, karantin: bool):
        """karantin=True (joriy davr): items/GTD dan faqat sanoat (+ bnpz/memo alohida hisob so'ralganda), meva — FAQAT karantin qoidasi
        (GTD kind='meva', o'stirilgan tuman bo'yicha, eksportyor reyestrda bo'lishi shart emas). karantin=False (o'tgan yil, bojxona bazasi): eski qoida."""
        agg = defaultdict(lambda: {"value": 0.0, "inns": set(), "gtd": set(), "netto": 0.0})
        per_inn = defaultdict(lambda: defaultdict(float))      # country -> inn -> value
        months = defaultdict(lambda: [0.0] * 12)
        def kind_ok(k):
            if k in ("bnpz", "memo"): return with_memo
            if k in ("meva", "meva_x"): return not karantin      # joriy davrda meva faqat karantin qatorlaridan
            return k == "sanoat"
        if has_base_detail(con):
            src_rows = list(con.execute("""SELECT i.inn, i.rdate, i.gtd, d.country, i.value, i.netto, i.kind FROM items i JOIN declarations d ON d.decl_id=i.decl_id
                                    WHERE i.regime='ЭК' AND i.rdate>=? AND i.rdate<=?""", (f, t)))
            for r in src_rows + (dispute.template_items(con, int(t[:4]), f, t) if tset else []):
                if not kind_ok(r["kind"]): continue
                if r["inn"] in tset and r["kind"] == "sanoat" and not str(r["gtd"] or "").startswith("tpl:"): continue
                c = comp.get(r["inn"])
                if not keep(r["inn"], c, r["rdate"]): continue
                nm = country_uz(r["country"]) if r["country"] else dispute.TPL_T1; a = agg[nm]
                a["value"] += r["value"] or 0; a["inns"].add(r["inn"]); a["gtd"].add(r["gtd"]); a["netto"] += r["netto"] or 0
                per_inn[nm][r["inn"]] += r["value"] or 0
                months[nm][int(r["rdate"][5:7]) - 1] += r["value"] or 0
        # X1: davrning birinchi kuni tushib qolmasin — report_date >= f VA report_date > through (ikki alohida shart)
        for r in con.execute("SELECT inn, exporter, district_code dc, report_date, gtd_no, country, stat_usd, netto, kind FROM gtd_rows WHERE voided=0 AND report_date>=? AND report_date<=?"
                             + (" AND report_date>?" if through else ""), (f, t, *((through,) if through else ()))):
            if r["kind"] in KIND_SKIP: continue
            if karantin and r["kind"] == "meva":
                if "meva" not in kinds and kinds: continue
                if dists and (r["dc"] or "?") not in dists: continue
                if tarmoqs and "Мева-сабзавот" not in tarmoqs: continue
                names.setdefault(r["inn"], r["exporter"])
            else:
                if not kind_ok(r["kind"]): continue
                c = comp.get(r["inn"])
                if not keep(r["inn"], c, r["report_date"]): continue
            nm = country_uz(r["country"]); a = agg[nm]
            a["value"] += r["stat_usd"] or 0; a["inns"].add(r["inn"]); a["gtd"].add(r["gtd_no"]); a["netto"] += (r["netto"] or 0) / 1000
            per_inn[nm][r["inn"]] += r["stat_usd"] or 0
            months[nm][int(r["report_date"][5:7]) - 1] += r["stat_usd"] or 0
        return agg, per_inn, months
    names = {}
    agg, per_inn, months = collect(f, t, True)
    # karantin mevasining baza davri (svod) — davlat kesimi yo'q: jamiga kiradi, qatorlarga emas
    from . import products
    y = int(t[:4]); meva_base = 0.0
    if (not kinds or "meva" in kinds) and (not tarmoqs or "Мева-сабзавот" in tarmoqs):
        meva_base = sum(v for dc, v in products.karantin_base(con, y, f, t).items() if not dists or dc in dists)
    # the same period one year earlier: new markets (no export then), lost markets (export then, none now)
    pf, pt = f"{int(f[:4]) - 1}{f[4:]}", f"{int(t[:4]) - 1}{t[4:]}"
    if pt.endswith("02-29"): pt = pt[:-2] + "28"
    pagg, pper, _ = collect(pf, pt, False)
    total = sum(a["value"] for a in agg.values()) + meva_base
    comp = {**{i: {"name": n} for i, n in names.items()}, **comp}
    hcode = next(iter(dists)) if len(dists) == 1 else "1706"
    pk = products.prev_kind(kinds)
    if pk == "all" and tarmoqs == {"Мева-сабзавот"}: pk = "meva"
    prev_total = products.prev_settings(con, hcode, pk, f, t)
    out = []
    for nm, a in agg.items():
        if countries and nm not in countries: continue
        top = sorted(per_inn[nm].items(), key=lambda x: -x[1])[:3]
        pv = pagg[nm]["value"] if nm in pagg else 0.0
        out.append({"country": nm, "code": ISO.get(nm), "value": round(a["value"], 3),
                    "share": round(a["value"] / total * 100, 2) if total else 0,
                    "companies": len(a["inns"]), "gtd": len(a["gtd"]), "netto_t": round(a["netto"], 1),
                    "prev": round(pv, 3), "new": pv == 0, "new_companies": sorted(a["inns"] - (pagg[nm]["inns"] if nm in pagg else set())),
                    "months": [round(v, 3) for v in months[nm]],
                    "top": [{"inn": i, "name": (comp.get(i) or {}).get("name"), "value": round(v, 3)} for i, v in top]})
    out.sort(key=lambda x: -x["value"])
    lost = sorted(({"country": nm, "code": ISO.get(nm), "prev": round(a["value"], 3), "companies": len(a["inns"]),
                    "top": [{"inn": i, "name": (comp.get(i) or {}).get("name"), "value": round(v, 3)} for i, v in sorted(pper[nm].items(), key=lambda x: -x[1])[:3]]}
                   for nm, a in pagg.items() if nm not in agg and a["value"] > 0), key=lambda x: -x["prev"])
    shown = sum(x["value"] for x in out)
    return {"rows": out, "from": f, "to": t, "prev_from": pf, "prev_to": pt, "prev_total": round(prev_total, 3), "prev_source": "settings",
            "prev_total_customs": round(sum(a["value"] for a in pagg.values()), 3), "rows_prev_source": "customs",
            "meva_rule": "karantin", "meva_base": round(meva_base, 3),      # karantin mevasining baza davri: jamida bor, davlat qatorlarida yo'q
            "lost": lost, "new_count": sum(1 for x in out if x["new"]), "total": round(total, 3), "shown": round(shown, 3),
            "countries": len(out), "companies": len({i for nm, d in per_inn.items() for i in d if not countries or nm in countries}),
            "gtd": sum(x["gtd"] for x in out), "unmapped": sorted({x["country"] for x in out if not x["code"]}),
            "districts": [{"code": c, "name": n} for c, n in sorted(dn.items(), key=lambda x: x[1])],
            "tarmoqlar": sorted({(c.get("tarmoq") or "").strip() for c in comp.values() if c.get("tarmoq")}),
            "all_countries": sorted(agg.keys())}

def country(con, q: dict) -> dict:
    """Bitta davlat: korxonalar, mahsulotlar, GTD lar."""
    from .companies import country_uz, has_base_detail
    f, t = _period(con, q)
    through = db.get_setting(con, "opening_through_date") or ""
    target = (q.get("country") or "").strip()
    code = (q.get("code") or "").strip().upper()
    if not target and code:
        target = next((nm for nm, cc in ISO.items() if cc == code), "")
    if not target: raise ValueError("Davlat tanlanmadi")
    dists, tarmoqs, kinds = _vals(q, "district"), _vals(q, "tarmoq"), _vals(q, "kind")
    with_memo = str(q.get("memo") or "1") != "0"
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, district_code, tarmoq, kind, excluded, bukhara, bukhara_from, bukhara_to FROM companies")}
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    def keep(c, d=None):
        if c is None: return False
        if c.get("bukhara") == 0: return False               # «Бухоро эмас»: base rows and GTD alike
        if d and not ((c.get("bukhara_from") or "") <= d <= (c.get("bukhara_to") or "9999")): return False   # outside the Bukhara period
        if c.get("excluded") and not with_memo: return False
        if dists and c.get("district_code") not in dists: return False
        if tarmoqs and (c.get("tarmoq") or "").strip() not in tarmoqs: return False
        if kinds and c.get("kind") not in kinds: return False
        return True
    by_inn = defaultdict(lambda: {"value": 0.0, "gtd": set(), "netto": 0.0, "last": None, "products": defaultdict(float)})
    prods = defaultdict(lambda: {"value": 0.0, "netto": 0.0, "name": None})
    def kind_ok(k):      # sanoat (+ bnpz/memo alohida hisobda); items dagi meva yo'q — meva faqat karantin (GTD kind='meva') qoidasi bilan
        if k in ("bnpz", "memo"): return with_memo
        return k == "sanoat"
    names = {}
    from . import dispute
    tset = dispute.template_inns(con)
    if has_base_detail(con):
        src_rows = list(con.execute("""SELECT i.inn, i.rdate, i.gtd, d.country, i.value, i.netto, i.hs10 hs, i.tovar1, i.tovar2, i.kind FROM items i JOIN declarations d ON d.decl_id=i.decl_id
                                WHERE i.regime='ЭК' AND i.rdate>=? AND i.rdate<=?""", (f, t)))
        for r in src_rows + [{**x, "hs": ""} for x in (dispute.template_items(con, int(t[:4]), f, t) if tset else [])]:
            if r["inn"] in tset and r["kind"] == "sanoat" and not str(r["gtd"] or "").startswith("tpl:"): continue
            if (country_uz(r["country"]) if r["country"] else dispute.TPL_T1) != target or not kind_ok(r["kind"]): continue
            c = comp.get(r["inn"])
            if not keep(c, r["rdate"]): continue
            o = by_inn[r["inn"]]; o["value"] += r["value"] or 0; o["gtd"].add(r["gtd"]); o["netto"] += r["netto"] or 0
            o["last"] = max(filter(None, [o["last"], r["rdate"]]), default=None)
            k = (r["hs"] or "")[:4] or "—"; p = prods[k]; p["value"] += r["value"] or 0; p["netto"] += r["netto"] or 0
            p["name"] = p["name"] or r["tovar2"] or r["tovar1"]
            o["products"][k] += r["value"] or 0
    # X1: report_date >= f VA report_date > through (ikki alohida shart)
    for r in con.execute("SELECT inn, exporter, district_code dc, report_date, gtd_no, country, stat_usd, netto, tnved, product, kind FROM gtd_rows WHERE voided=0 AND report_date>=? AND report_date<=?"
                         + (" AND report_date>?" if through else ""), (f, t, *((through,) if through else ()))):
        if r["kind"] in KIND_SKIP or country_uz(r["country"]) != target: continue
        if r["kind"] == "meva":       # karantin qoidasi: o'stirilgan tuman bo'yicha, eksportyor reyestrda bo'lishi shart emas
            if kinds and "meva" not in kinds: continue
            if dists and (r["dc"] or "?") not in dists: continue
            if tarmoqs and "Мева-сабзавот" not in tarmoqs: continue
            names.setdefault(r["inn"], {"name": r["exporter"], "district_code": r["dc"], "tarmoq": "Мева-сабзавот"})
        else:
            if not kind_ok(r["kind"]): continue
            c = comp.get(r["inn"])
            if not keep(c, r["report_date"]): continue
        o = by_inn[r["inn"]]; o["value"] += r["stat_usd"] or 0; o["gtd"].add(r["gtd_no"]); o["netto"] += (r["netto"] or 0) / 1000
        o["last"] = max(filter(None, [o["last"], r["report_date"]]), default=None)
        k = str(r["tnved"] or "")[:4] or "—"; p = prods[k]; p["value"] += r["stat_usd"] or 0; p["netto"] += (r["netto"] or 0) / 1000
        p["name"] = p["name"] or (r["product"] or "")[:60]
        o["products"][k] += r["stat_usd"] or 0
    tot = sum(o["value"] for o in by_inn.values())
    comp = {**names, **comp}
    comps = sorted(({"inn": i, "name": (comp.get(i) or {}).get("name"), "district": dn.get((comp.get(i) or {}).get("district_code")),
                     "tarmoq": (comp.get(i) or {}).get("tarmoq"), "value": round(o["value"], 3),
                     "share": round(o["value"] / tot * 100, 1) if tot else 0, "gtd": len(o["gtd"]),
                     "netto_t": round(o["netto"], 1), "last": o["last"],
                     "product": (prods.get(max(o["products"], key=o["products"].get))or{}).get("name") if o["products"] else None}
                    for i, o in by_inn.items()), key=lambda x: -x["value"])
    plist = sorted(({"hs4": k, "name": v["name"], "value": round(v["value"], 3), "netto_t": round(v["netto"], 1),
                     "share": round(v["value"] / tot * 100, 1) if tot else 0} for k, v in prods.items()), key=lambda x: -x["value"])
    return {"country": target, "code": ISO.get(target), "from": f, "to": t, "total": round(tot, 3),
            "companies": comps, "products": plist[:12], "gtd": sum(len(o["gtd"]) for o in by_inn.values()),
            "netto_t": round(sum(o["netto"] for o in by_inn.values()), 1)}

def export_xlsx(con, q: dict, out) -> object:
    """Davlatlar kesimi + har bir davlat bo'yicha korxonalar — ikki varaqli Excel."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    d = rows(con, q)
    thin = Side(style="thin", color="D6DBD8"); bd = Border(thin, thin, thin, thin)
    H = Font(bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor="0F6E63")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Давлатлар"
    ws["A1"] = "Бухоро вилояти экспорти — давлатлар кесимида"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{d['from'][8:10]}.{d['from'][5:7]}.{d['from'][:4]} — {d['to'][8:10]}.{d['to'][5:7]}.{d['to'][:4]} · минг АҚШ доллари"
    ws["A2"].font = Font(italic=True, color="5A6470", size=10)
    hdr = ["№", "Давлат", "Код", "Қиймати", "Улуши", "Корхоналар", "ГТД", "Нетто, т"]
    for i, (h, w) in enumerate(zip(hdr, [5, 26, 7, 15, 10, 13, 9, 13]), 1):
        c = ws.cell(4, i, h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center"); c.border = bd
        ws.column_dimensions[L(i)].width = w
    for n, r in enumerate(d["rows"], 1):
        for i, v in enumerate([n, r["country"], r["code"], r["value"], (r["share"] or 0) / 100, r["companies"], r["gtd"], r["netto_t"]], 1):
            c = ws.cell(4 + n, i, v); c.border = bd
            if i in (4, 8): c.number_format = "#,##0.0"
            if i == 5: c.number_format = "0.0%"
    t = 5 + len(d["rows"])
    ws.cell(t, 2, f"Жами: {d['countries']} та давлат").font = Font(bold=True)
    c = ws.cell(t, 4, f"=SUM(D5:D{max(t - 1, 5)})"); c.number_format = "#,##0.0"; c.font = Font(bold=True)
    ws.freeze_panes = "A5"
    ws2 = wb.create_sheet("Корхоналар")
    hdr2 = ["Давлат", "ИНН", "Корхона", "Туман", "Тармоқ", "Қиймати", "Давлат ичида %", "ГТД", "Охирги"]
    for i, (h, w) in enumerate(zip(hdr2, [20, 13, 42, 18, 22, 14, 15, 8, 12]), 1):
        c = ws2.cell(1, i, h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center"); c.border = bd
        ws2.column_dimensions[L(i)].width = w
    r0 = 2
    for r in d["rows"]:
        det = country(con, {**q, "country": r["country"]})
        for x in det["companies"]:
            for i, v in enumerate([r["country"], x["inn"], x["name"], x["district"], x["tarmoq"], x["value"], (x["share"] or 0) / 100, x["gtd"], x["last"]], 1):
                c = ws2.cell(r0, i, v); c.border = bd
                if i == 6: c.number_format = "#,##0.0"
                if i == 7: c.number_format = "0.0%"
            r0 += 1
    ws2.freeze_panes = "A2"
    wb.save(out); return out
