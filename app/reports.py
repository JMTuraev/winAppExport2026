"""Ҳисобот тузувчи — rasmiy Excel jadvallar (2-jadval uslubida), istalgan kesim va davr bo'yicha.

Qatorlar: tumanlar (saноат/мева qatorlari bilan), tarmoqlar, davlatlar, uyushmalar, korxonalar.
Ustunlar: o'tgan yil amalda (shu davr), o'tgan yil jami, prognoz (reja), amalda, ijro % va farqi, o'sish % va farqi,
korxonalar soni, GTD, netto. Davr: yil boshidan / oy / chorak / sana oralig'i. Reja — Sozlamalardagi versiyali reja.
Ma'lumot: sanoat — items + kunlik GTD (products._rows, kind sanoat); meva — karantin qoidasi (products.karantin_period = report.karantin).
«Ўтган йил» vilyoat/tumanlar uchun — faqat Sozlamalardagi baza (prev_source='settings'); boshqa kesimlarda bojxona bazasi ('customs').
"""
from __future__ import annotations
import datetime as dt, calendar
from collections import defaultdict
from . import db, report, products, plans

ROW_KINDS = {"district": "туманлар", "tarmoq": "тармоқлар", "country": "давлатлар", "union": "уюшмалар", "company": "корхоналар"}
COL_DEFS = {   # key: (label, group, kind)  kind: money | pct | int
    "prev_fact": ("Амалда", "{py} йил", "money"), "prev_full": ("Йил жами", "{py} йил", "money"),
    "plan": ("Прогноз", "{y} йил", "money"), "fact": ("Амалда", "{y} йил", "money"),
    "exec_pct": ("%да", "Ижро", "pct"), "exec_diff": ("Фарқи", "Ижро", "money"),
    "growth_pct": ("%да", "Ўсиши", "pct"), "growth_diff": ("Фарқи", "Ўсиши", "money"),
    "share": ("Улуши, %", "", "pct"), "companies": ("Корхоналар", "", "int"), "gtd": ("ГТД", "", "int"), "netto": ("Нетто, т", "", "money"),
}
DEFAULT_COLS = ["prev_fact", "plan", "fact", "exec_pct", "exec_diff", "growth_pct", "growth_diff"]
MONTHS = plans.MONTHS

def period(con, p: dict) -> tuple[int, str, str, str]:
    """(year, from, to, label)"""
    cur = int(db.get_setting(con, "year") or dt.date.today().year)
    y = int(p.get("year") or cur)
    last = report.last_data_date(con) or dt.date.today().isoformat()
    t = p.get("type") or "ytd"
    if t == "month":
        m = int(p.get("month") or (int(last[5:7]) if y == cur else 12))
        f, to = f"{y}-{m:02d}-01", f"{y}-{m:02d}-{calendar.monthrange(y, m)[1]:02d}"; label = f"{MONTHS[m - 1].lower()} ойи"
    elif t == "quarter":
        qn = int(p.get("quarter") or 1); m0 = (qn - 1) * 3 + 1; m1 = m0 + 2
        f, to = f"{y}-{m0:02d}-01", f"{y}-{m1:02d}-{calendar.monthrange(y, m1)[1]:02d}"; label = f"{qn}-чорак"
    elif t == "range":
        f, to = (p.get("from") or f"{y}-01-01")[:10], (p.get("to") or last)[:10]
        if not f.startswith(str(y)): f = f"{y}-01-01"
        if not to.startswith(str(y)): to = f"{y}-12-31"
        label = f"{f[8:10]}.{f[5:7]} — {to[8:10]}.{to[5:7]}"
    elif t == "year":
        f, to = f"{y}-01-01", f"{y}-12-31"; label = "йил якуни"
    else:
        f = f"{y}-01-01"; to = last if y == cur else f"{y}-12-31"; label = f"01.01 — {to[8:10]}.{to[5:7]}"
    if y == cur and to > last: to = last
    return y, f, min(to, f"{y}-12-31"), label

def _key(kind_rows, r, comp, dn):
    c = comp.get(r[0]) or {}
    if kind_rows == "district": return c.get("district_code") or "—"
    if kind_rows == "tarmoq": return r[2]
    if kind_rows == "country": return r[9] or "—"
    if kind_rows == "union": return c.get("union_name") or "Худудий ташкилотлар"
    return r[0]

def build(con, cfg: dict) -> dict:
    rows_kind = cfg.get("rows") or "district"
    if rows_kind not in ROW_KINDS: raise ValueError("Qator turi noto'g'ri")
    cols = [c for c in (cfg.get("cols") or DEFAULT_COLS) if c in COL_DEFS]
    split = bool(cfg.get("split_kind", rows_kind == "district"))
    y, f, t, plabel = period(con, cfg.get("period") or {})
    pf, pt = f"{y - 1}{f[4:]}", f"{y - 1}{t[4:]}"
    if pt.endswith("02-29"): pt = pt[:-2] + "28"
    q = {"district": cfg.get("district") or "", "kind": cfg.get("kind") or "", "memo": "1" if cfg.get("memo") else "0"}
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, district_code, tarmoq, kind, excluded, in_base, bukhara, bukhara_from, bukhara_to, source, union_name FROM companies")}
    meva_inns = {i for (i,) in con.execute("SELECT inn FROM template_inns WHERE category='meva'")} or \
                {i for (i,) in con.execute("SELECT inn FROM companies WHERE kind='meva' AND source='nomma-nom'")}
    keep = products._keeper(con, q, comp, meva_inns)
    hs, hs4 = products._hs_dict(con), products._hs4_dict(con)
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    order = {d["code"]: i for i, d in enumerate(report.districts(con))}
    kinds, dists = products._vals(q, "kind") or {"sanoat", "meva"}, products._vals(q, "district")
    def meva_key(dc):
        if rows_kind == "district": return dc
        if rows_kind == "tarmoq": return products.MEVA_T1
        return products.KAR_ROW
    def collect(yy, f_, t_, kp, with_kar):
        """sanoat — items + GTD (kind sanoat); meva — FAQAT karantin qoidasi (o'stirilgan tuman; dashboard/svod bilan bir xil)."""
        agg = defaultdict(lambda: {"sanoat": 0.0, "meva": 0.0, "netto": 0.0, "inns": set(), "gtd": set()})
        for r in products._rows(con, yy, f_, t_, comp, kp, hs, hs4):
            k = _key(rows_kind, r, comp, dn); a = agg[k]
            a[r[12]] += r[6]; a["netto"] += r[7]; a["inns"].add(r[0]); a["gtd"].add(r[8])
        if with_kar and "meva" in kinds:
            for dc, x in products.karantin_period(con, yy, f_, t_, dists).items():
                a = agg[meva_key(dc)]; a["meva"] += x["total"]; a["inns"] |= x["inns"]; a["gtd"] |= x["gtds"]
        return agg
    cur = collect(y, f, t, products.sanoat_only(keep), True)
    prev = collect(y - 1, pf, pt, keep, False)                                   # faqat tarmoq/davlat/uyushma/korxona qatorlari uchun (bojxona bazasi)
    prev_full = collect(y - 1, f"{y - 1}-01-01", f"{y - 1}-12-31", keep, False) if "prev_full" in cols and rows_kind != "district" else {}
    # «ўтган йил шу даврга» va «йил жами» — vilyoat va tumanlar uchun FAQAT Sozlamalardagi baza (report.prev_year_amount)
    # filtr: bitta tuman tanlangan bo'lsa vilyoat qatori ham o'sha tuman bazasidan; tur filtri bo'lsa 'all' o'rniga o'sha tur
    def _ck(code, kind):
        if code == "1706" and len(dists) == 1: code = next(iter(dists))
        if kind == "all": kind = products.prev_kind(kinds)
        return code, kind
    def prev_set(code, kind): return products.prev_settings(con, *_ck(code, kind), f, t)
    def prev_full_set(code, kind): return products.prev_full_settings(con, *_ck(code, kind), y)
    # reja: faqat tumanlar kesimida (viloyat jami = yig'indi)
    has_plan = "plan" in cols or "exec_pct" in cols or "exec_diff" in cols
    plan_v = plans.effective_version(con, t) if has_plan else None
    old_pr = report.prognoz(con, y) if has_plan and not plan_v else {}
    ptype = (cfg.get("period") or {}).get("type") or "ytd"
    def plan_of(code, kind):
        if not has_plan or rows_kind != "district": return None
        # reja: avvalo versiyali reja, so'ng СВОД prognozi (davriy natijalar endi faqat tugagan yillar — rejasiz)
        if plan_v: return plans.period_plan(con, f, t, None if code == "1706" else code, kind)
        o = (old_pr.get(code) or {}).get(kind) or {}
        if ptype == "year": return o.get("year_plan")
        if ptype == "ytd": return o.get("period_plan")
        if ptype == "month": return o.get("cur_month_plan")
        return None
    keys = set(cur) | set(prev)
    if rows_kind == "district": keys = [c for c in sorted(keys, key=lambda k: order.get(k, 99))]
    else: keys = sorted(keys, key=lambda k: -(cur[k]["sanoat"] + cur[k]["meva"]) if k in cur else 0)
    top = int(cfg.get("top") or 0)
    if top and rows_kind != "district": keys = keys[:top]
    total_fact = sum(cur[k]["sanoat"] + cur[k]["meva"] for k in cur)
    def cells(code, kind, a, p, pfull, settings=False):
        fv = (a["sanoat"] + a["meva"]) if kind == "all" else a[kind]
        if settings:      # vilyoat / tuman: Sozlamalardagi ўтган йил базаси
            pv, pfv = prev_set(code, kind), prev_full_set(code, kind)
        else:             # boshqa kesimlar: bojxona bazasi (2025 items) — qo'lda manba yo'q
            pv = (p["sanoat"] + p["meva"]) if kind == "all" else p[kind]
            pfv = (pfull["sanoat"] + pfull["meva"]) if kind == "all" else pfull[kind]
        pl = plan_of(code, kind)
        out = {"fact": fv, "prev_fact": pv, "prev_full": pfv, "plan": pl,
               "exec_pct": (fv / pl * 100) if pl else None, "exec_diff": (fv - pl) if pl is not None else None,
               "growth_pct": (fv / pv * 100) if pv else None, "growth_diff": fv - pv,
               "share": (fv / total_fact * 100) if total_fact else None,
               "companies": len(a["inns"]) if kind == "all" else None, "gtd": len(a["gtd"]) if kind == "all" else None, "netto": a["netto"] if kind == "all" else None}
        return {k: (round(v, 3) if isinstance(v, float) else v) for k, v in out.items()}
    empty = {"sanoat": 0.0, "meva": 0.0, "netto": 0.0, "inns": set(), "gtd": set()}
    def merged(aggs):
        m = {"sanoat": 0.0, "meva": 0.0, "netto": 0.0, "inns": set(), "gtd": set()}
        for a in aggs: m["sanoat"] += a["sanoat"]; m["meva"] += a["meva"]; m["netto"] += a["netto"]; m["inns"] |= a["inns"]; m["gtd"] |= a["gtd"]
        return m
    out_rows = []
    tot_c, tot_p, tot_pf = merged(cur.values()), merged(prev.values()), merged(prev_full.values()) if prev_full else empty
    out_rows.append({"n": "", "label": "Бухоро вилояти", "level": 0, "prev_source": "settings", "cells": cells("1706", "all", tot_c, tot_p, tot_pf, True)})
    if split:
        out_rows.append({"n": "", "label": "саноат маҳсулотлари", "level": 1, "prev_source": "settings", "cells": cells("1706", "sanoat", tot_c, tot_p, tot_pf, True)})
        out_rows.append({"n": "", "label": "мева-сабзавотлар", "level": 1, "prev_source": "settings", "cells": cells("1706", "meva", tot_c, tot_p, tot_pf, True)})
    for i, k in enumerate(keys, 1):
        a, p, pfull = cur.get(k, empty), prev.get(k, empty), prev_full.get(k, empty) if prev_full else empty
        is_d = rows_kind == "district" and k in dn; src = "settings" if is_d else "customs"
        if rows_kind == "district": label = dn.get(k, k)
        elif rows_kind == "company": label = (comp.get(k) or {}).get("name") or k
        else: label = k
        row = {"n": i, "label": label, "level": 0, "prev_source": src, "cells": cells(k, "all", a, p, pfull, is_d)}
        if rows_kind == "company": row["inn"] = k if k in comp else None; row["district"] = dn.get((comp.get(k) or {}).get("district_code"))
        out_rows.append(row)
        if split:
            out_rows.append({"n": "", "label": "саноат маҳсулотлари", "level": 1, "prev_source": src, "cells": cells(k, "sanoat", a, p, pfull, is_d)})
            out_rows.append({"n": "", "label": "мева-сабзавотлар", "level": 1, "prev_source": src, "cells": cells(k, "meva", a, p, pfull, is_d)})
    columns = [{"key": c, "label": COL_DEFS[c][0], "group": COL_DEFS[c][1].format(y=y, py=y - 1), "kind": COL_DEFS[c][2]} for c in cols]
    title = cfg.get("title") or f"Бухоро вилоятининг {y} йил {plabel} экспорт кўрсаткичлари ({ROW_KINDS[rows_kind]} кесимида)"
    sign_pos = db.get_setting(con, "report_sign_pos", "Бошқарма бошлиғи"); sign_name = db.get_setting(con, "report_sign_name", "")
    return {"title": title, "year": y, "from": f, "to": t, "prev_from": pf, "prev_to": pt, "period_label": plabel, "rows_kind": rows_kind, "split": split,
            "columns": columns, "rows": out_rows, "unit": "минг АҚШ долл.", "plan_source": ("versions" if plan_v else ("workbook" if old_pr else None)) if has_plan else None,
            "plan_version": plan_v, "sign": {"position": sign_pos, "name": sign_name}, "date": dt.date.today().isoformat(),
            "prev_source": "settings" if rows_kind == "district" else "mixed", "meva_rule": "karantin",
            "note": f"{y} йил {f[8:10]}.{f[5:7]}–{t[8:10]}.{t[5:7]}; ўтган йил {pf[8:10]}.{pf[5:7]}–{pt[8:10]}.{pt[5:7]}. Манба: божхона базаси + кундалик ГТД; "
                    f"мева-сабзавот — карантин рўйхати (1706 ҳудудда етиштирилган); ўтган йил — Созламалардаги база (кунларга мутаносиб)."}

# ------------------------------------------------------------------ Excel
def export_xlsx(con, cfg: dict, out, style: str = "official"):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    d = build(con, cfg)
    off = style == "official"
    FN = "Calibri" if off else "Times New Roman"
    fs = {"title": 24 if off else 16, "unit": 18 if off else 11, "head": 21 if off else 12, "body": 20 if off else 12, "sign": 20 if off else 13}
    med = Side(style="medium", color="000000"); thin = Side(style="thin", color="000000")
    B = lambda l=thin, r=thin, tp=thin, b=thin: Border(left=l, right=r, top=tp, bottom=b)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ҳисобот"
    cols = d["columns"]; ncol = 2 + len(cols)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws["A1"]; c.value = d["title"]; c.font = Font(name=FN, size=fs["title"], bold=True); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 74 if off else 44
    c = ws.cell(2, ncol, d["unit"]); c.font = Font(name=FN, size=fs["unit"]); c.alignment = Alignment(horizontal="right")
    # header: 2 rows (group / label), groups merged
    ws.cell(3, 1, "Т/р"); ws.cell(3, 2, ROW_KINDS[d["rows_kind"]].capitalize())
    ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1); ws.merge_cells(start_row=3, start_column=2, end_row=4, end_column=2)
    i = 3
    while i < 3 + len(cols):
        g = cols[i - 3]["group"]; j = i
        while j + 1 < 3 + len(cols) and cols[j + 1 - 3]["group"] == g and g: j += 1
        if g:
            ws.cell(3, i, g)
            if j > i: ws.merge_cells(start_row=3, start_column=i, end_row=3, end_column=j)
            for k in range(i, j + 1): ws.cell(4, k, cols[k - 3]["label"])
        else:
            ws.cell(3, i, cols[i - 3]["label"]); ws.merge_cells(start_row=3, start_column=i, end_row=4, end_column=i)
        i = j + 1
    for rr in (3, 4):
        ws.row_dimensions[rr].height = 34 if off else 24
        for cc in range(1, ncol + 1):
            c = ws.cell(rr, cc); c.font = Font(name=FN, size=fs["head"], bold=True); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = B(med, med, med if rr == 3 else thin, med if rr == 4 else thin)
    r0 = 5
    for row in d["rows"]:
        lvl = row["level"]; bold = lvl == 0
        ws.cell(r0, 1, row["n"] if row["n"] != "" else None); ws.cell(r0, 2, ("   " if lvl else "") + row["label"])
        ws.cell(r0, 1).alignment = Alignment(horizontal="center", vertical="center"); ws.cell(r0, 2).alignment = Alignment(horizontal="left" if lvl else "center", vertical="center", wrap_text=True)
        for cc, col in enumerate(cols, 3):
            v = row["cells"].get(col["key"]); c = ws.cell(r0, cc, v)
            c.number_format = {"money": "#,##0.0", "pct": "0.0", "int": "#,##0"}[col["kind"]]
            c.alignment = Alignment(horizontal="center" if col["kind"] != "money" else "right", vertical="center")
        for cc in range(1, ncol + 1):
            c = ws.cell(r0, cc); c.font = Font(name=FN, size=fs["body"], bold=bold); c.border = B(med if cc in (1, 2) or cc == ncol else thin, med if cc == ncol or cc == 1 else thin, thin, thin)
        ws.row_dimensions[r0].height = 35 if off else 20
        r0 += 1
    for cc in range(1, ncol + 1): ws.cell(r0 - 1, cc).border = B(thin, thin, thin, med)
    # imzo
    r0 += 2
    sg = d["sign"]
    if sg.get("position") or sg.get("name"):
        c = ws.cell(r0, 2, sg.get("position") or ""); c.font = Font(name=FN, size=fs["sign"], bold=True)
        c = ws.cell(r0, ncol, sg.get("name") or ""); c.font = Font(name=FN, size=fs["sign"], bold=True); c.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 8 if off else 6; ws.column_dimensions["B"].width = 58 if off else 40
    for cc in range(3, ncol + 1): ws.column_dimensions[L(cc)].width = 22 if off else 15
    ws.print_options.horizontalCentered = True; ws.page_setup.orientation = "landscape" if ncol > 7 else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "C5"
    wb.save(out); return out
