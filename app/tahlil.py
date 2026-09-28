"""Ҳисоботлар → «Тармоқлар таҳлили»: tanlangan tarmoq (tovar1) korxonalari — joriy yil va o'tgan yil shu davr, korxona -> mahsulot (sprav2).

Guruhlar (Jafar shabloni «Текстиль 8 ой 2026 сараланган.xlsx» bilan aynan bir xil):
  I   Экспорти кўпайган корхоналар   (o'tgan yil > 0, joriy >= o'tgan)     — o'sish bo'yicha kamayish tartibida
  II  Янги экспорт қилган корхоналар (o'tgan yil = 0)                       — o'sish bo'yicha kamayish tartibida
  III Экспорти камайган корхоналар   (joriy < o'tgan, joriy > 0)            — o'sish bo'yicha o'sish tartibida (eng ko'p kamaygan birinchi)
  IV  {y} йилда экспорт қилмаган     (joriy = 0)                             — o'sish bo'yicha kamayish tartibida
Korxona ichida mahsulotlar ham guruh tartibida (III — o'sish tartibida, qolganlari — kamayish).
Ma'lumot manbai va qoidalar — products._rows (Тармоқ ва маҳсулот sahifasi bilan bir xil): bojxona bazasi + keyingi kunlik GTD,
«Бухоро эмас» chiqariladi, Buxoro davri, БНПЗ/SDK faqat memo=1 da, meva — faqat sanoat hisobidagi qatorlar (karantin sintetik qatori korxonaga bog'lanmagan).
Excel — app/templates/tarmoq_tahlil.xlsx shablonidan (uslub, mavzu ranglari, formulalar shablondagidek).
"""
from __future__ import annotations
import copy, datetime as dt, re
from collections import defaultdict
from pathlib import Path
from . import db, report, products as P

TPL = Path(__file__).resolve().parent / "templates" / "tarmoq_tahlil.xlsx"
MONTHS_GEN = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
PRESETS = [{"id": "textile", "label": "Тўқимачилик", "title": "тўқимачилик", "t1": ["Тукимачилик махсулотлари", "Калава ип"],
            "hint": "Тўқимачилик маҳсулотлари + Калава ип"}]
G_UP, G_NEW, G_DOWN, G_STOP = "up", "new", "down", "stop"
G_ORDER = [G_UP, G_NEW, G_DOWN, G_STOP]
ROMAN = {G_UP: "I", G_NEW: "II", G_DOWN: "III", G_STOP: "IV"}

def _status(g: str, y: int) -> str:
    return {G_UP: "кўпайган", G_NEW: "янги корхона", G_DOWN: "камайган", G_STOP: f"{y} йил экспорт қилмаган"}[g]

def _gtitle(g: str, y: int) -> str:
    return {G_UP: "Экспорти кўпайган корхоналар", G_NEW: "Янги экспорт қилган корхоналар", G_DOWN: "Экспорти камайган корхоналар",
            G_STOP: f"{y} йилда экспорт қилмаган корхоналар"}[g]

def _cur_year(con) -> int: return int(db.get_setting(con, "year") or dt.date.today().year)

def _month_end(y: int, m: int) -> str:
    return (dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)).isoformat()

def default_months(con) -> int:
    """Oxirgi to'liq oy: ma'lumot 22.09 gacha bo'lsa -> 8 (shablondagi «8 ойда»)."""
    last = report.last_data_date(con) or dt.date.today().isoformat()
    d = dt.date.fromisoformat(last)
    return d.month if _month_end(d.year, d.month) == last else max(1, d.month - 1)

def _period(con, q) -> dict:
    """q.ptype: months (1..N ой) | ytd (охирги маълумот кунигача) | range (from..to). Ikkala yil uchun bir xil kunlar."""
    cy = _cur_year(con)
    y = int(q.get("year") or cy)
    last = report.last_data_date(con) or dt.date.today().isoformat()
    pt = (q.get("ptype") or "months").strip()
    if pt == "months":
        n = max(1, min(12, int(q.get("months") or default_months(con))))
        f, t = f"{y}-01-01", _month_end(y, n)
        lbl = f"{n} ойда" if n < 12 else "йил давомида"
    else:
        if pt == "range" and q.get("from") and q.get("to"):
            f, t = str(q["from"])[:10], str(q["to"])[:10]
            f, t = f"{y}{f[4:]}", f"{y}{t[4:]}"
        else:
            pt = "ytd"; f = f"{y}-01-01"; t = last if y == cy else f"{y}-12-31"
            if y == cy and not t.startswith(str(y)): t = f"{y}-12-31"
        if t < f: f, t = t, f
        d0, d1 = dt.date.fromisoformat(f), dt.date.fromisoformat(t)
        lbl = (f"{d0.day} {MONTHS_GEN[d0.month - 1]} – {d1.day} {MONTHS_GEN[d1.month - 1]}" if (d0.month, d0.day) != (1, 1)
               else f"1 январь – {d1.day} {MONTHS_GEN[d1.month - 1]}")
    pf, ptt = f"{y - 1}{f[4:]}", f"{y - 1}{t[4:]}"
    if ptt.endswith("02-29"): ptt = ptt[:-2] + "28"
    return {"year": y, "prev_year": y - 1, "type": pt, "from": f, "to": t, "prev_from": pf, "prev_to": ptt, "label": lbl,
            "data_to": min(t, last) if y == cy else t, "partial": y == cy and last < t}

def categories(con) -> list:
    """Tarmoqlar (tovar1) — joriy va o'tgan yil qiymati bilan (tanlash ro'yxati uchun)."""
    cy = _cur_year(con)
    vals = defaultdict(lambda: [0.0, 0.0])
    for r in con.execute("SELECT tovar1, year, sum(value) v FROM items WHERE regime='ЭК' AND kind='sanoat' AND year IN (?,?) GROUP BY tovar1, year", (cy, cy - 1)):
        if r["tovar1"]: vals[r["tovar1"]][0 if r["year"] == cy else 1] += r["v"] or 0
    out = [{"id": p["id"], "label": p["label"], "hint": p["hint"], "t1": p["t1"], "preset": True,
            "value": round(sum(vals[t][0] for t in p["t1"]), 1)} for p in PRESETS]
    for t1, (v, pv) in sorted(vals.items(), key=lambda x: -x[1][0]):
        out.append({"id": "t1:" + t1, "label": t1, "t1": [t1], "value": round(v, 1), "prev": round(pv, 1)})
    return out

def _cat(q) -> tuple[list, str, str]:
    """(tovar1 ro'yxati, sahifa nomi, sarlavhadagi ibora)."""
    cid = (q.get("cat") or "textile").strip()
    for p in PRESETS:
        if p["id"] == cid: return p["t1"], p["label"], f"{p['title']} корхоналари"
    t1 = cid[3:] if cid.startswith("t1:") else cid
    return [t1], t1, f"«{t1}» тармоғи корхоналари"

def build(con, q: dict) -> dict:
    t1s, cat_label, cat_phrase = _cat(q)
    per = _period(con, q)
    y, py = per["year"], per["prev_year"]
    comp = {r["inn"]: dict(r) for r in con.execute("SELECT inn, name, full_name, district_code, tarmoq, kind, excluded, in_base, bukhara, bukhara_from, bukhara_to, source FROM companies")}
    meva_inns = {i for (i,) in con.execute("SELECT inn FROM template_inns WHERE category='meva'")} or \
                {i for (i,) in con.execute("SELECT inn FROM companies WHERE kind='meva' AND source='nomma-nom'")}
    qk = {"district": q.get("district") or "", "memo": q.get("memo") or "0", "kind": "sanoat"}
    keep = P.sanoat_only(P._keeper(con, qk, comp, meva_inns))
    hs, hs4 = P._hs_dict(con), P._hs4_dict(con)
    want = set(t1s)
    # data[inn][product] = [prev_n, prev_v, cur_n, cur_v]
    data = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0, 0.0, 0.0]))
    for yy, f, t, off in ((py, per["prev_from"], per["prev_to"], 0), (y, per["from"], per["to"], 2)):
        for r in P._rows(con, yy, f, t, comp, keep, hs, hs4):
            if r[2] not in want: continue
            prod = r[5] if r[5] and not re.fullmatch(r"[\d—-]+", str(r[5])) else (r[4] or r[3] or "Бошқа")
            cell = data[r[0]][prod]; cell[off] += r[7] or 0.0; cell[off + 1] += r[6] or 0.0
    dn = {d["code"]: d["name_uz"] for d in report.districts(con)}
    groups = {g: [] for g in G_ORDER}
    for inn, prods in data.items():
        e = sum(p[1] for p in prods.values()); g_ = sum(p[3] for p in prods.values())
        if abs(e) < 1e-9 and abs(g_) < 1e-9: continue
        grp = G_NEW if e == 0 else G_STOP if g_ == 0 else G_UP if g_ >= e else G_DOWN
        asc = grp == G_DOWN
        plist = [{"name": k, "prev_n": v[0], "prev_v": v[1], "cur_n": v[2], "cur_v": v[3]} for k, v in prods.items() if v[1] or v[3]]
        plist.sort(key=lambda p: ((p["cur_v"] - p["prev_v"]) if asc else -(p["cur_v"] - p["prev_v"]), p["name"]))
        c = comp.get(inn) or {}
        groups[grp].append({"inn": inn, "name": c.get("full_name") or c.get("name") or inn, "short": c.get("name") or "", "district": dn.get(c.get("district_code")) or "",
                            "prev_n": sum(p["prev_n"] for p in plist), "prev_v": e, "cur_n": sum(p["cur_n"] for p in plist), "cur_v": g_,
                            "status": _status(grp, y), "products": plist})
    for g, lst in groups.items():
        asc = g == G_DOWN
        lst.sort(key=lambda c: ((c["cur_v"] - c["prev_v"]) if asc else -(c["cur_v"] - c["prev_v"]), c["name"]))
    def tot(lst):
        return {k: sum(c[k] for c in lst) for k in ("prev_n", "prev_v", "cur_n", "cur_v")}
    out_groups = [{"key": g, "roman": ROMAN[g], "title": _gtitle(g, y), "status": _status(g, y), "count": len(groups[g]), **tot(groups[g]),
                   "companies": groups[g]} for g in G_ORDER]
    allc = [c for g in G_ORDER for c in groups[g]]
    title = (q.get("title") or "").strip() or \
        f"Бухоро вилоятидаги {cat_phrase} томонидан амалга оширилган экспорт ҳажмлари\n({py}-{y} йиллар {per['label']}) "
    prods_all = defaultdict(lambda: [0.0, 0.0])
    for c in allc:
        for p in c["products"]: prods_all[p["name"]][0] += p["prev_v"]; prods_all[p["name"]][1] += p["cur_v"]
    return {"category": cat_label, "t1": t1s, "title": title, "period": per, "unit": "минг долл", "count": len(allc), **tot(allc),
            "groups": out_groups, "products": sorted([{"name": k, "prev_v": v[0], "cur_v": v[1]} for k, v in prods_all.items()], key=lambda x: -x["cur_v"]),
            "last_date": report.last_data_date(con)}

# ------------------------------------------------------------------ Excel (shablon bo'yicha)
def _inn_val(inn):
    s = str(inn or "")
    return int(s) if s.isdigit() and len(s) < 16 else s

def export_xlsx(con, q: dict, out: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font
    d = build(con, q)
    y, py, lbl = d["period"]["year"], d["period"]["prev_year"], d["period"]["label"]
    wb = openpyxl.load_workbook(TPL)
    ws, wk = wb["Соҳалар билан"], wb["Корхоналар"]
    def styles(sh, rows):
        return {r: [copy.copy(sh.cell(r, c)._style) for c in range(1, 11)] for r in rows}
    S = styles(ws, (5, 6, 7, 8)); K = styles(wk, (5, 6, 7))
    kh6 = wk.row_dimensions[6].height
    ws.delete_rows(5, ws.max_row); wk.delete_rows(5, wk.max_row)
    for sh in (ws, wk):
        sh["A1"].value = d["title"]
        sh["D3"].value = f"{py} йил {lbl}"; sh["F3"].value = f"{y} йил {lbl}"
    def put(sh, r, st, vals):
        for c in range(1, 11):
            cell = sh.cell(r, c); cell._style = copy.copy(st[c - 1])
            v = vals[c - 1] if c - 1 < len(vals) else None
            if v is not None: cell.value = v
    def color(sh, r, cols, rgb):
        for c in cols:
            cell = sh.cell(r, c); f = copy.copy(cell.font); f.color = rgb; cell.font = f
    red, green = "FFFF0000", "FF00B050"
    hi = lambda r: [f"=G{r}-E{r}", f"=IF(E{r}=0,IF(G{r}=0,0,100),G{r}/E{r}*100)"]
    stf = lambda r: (f'=IF(E{r}=0,"янги корхона",IF(G{r}=0,"{y} йил экспорт қилмаган",IF(G{r}>=E{r},"кўпайган","камайган")))')
    num = lambda v: round(v, 6) if v else None
    # ---- «Соҳалар билан»
    r = 6; grow_s = []; comp_rows = []    # comp_rows: (group, sheet row) — «Корхоналар» havolalari uchun
    for g in d["groups"]:
        gr = r; grow_s.append(gr); r += 1; first = r
        for c in g["companies"]:
            cr = r; r += 1
            p0 = r
            for p in c["products"]:
                put(ws, r, S[8], [None, p["name"], p["name"], num(p["prev_n"]), num(p["prev_v"]), num(p["cur_n"]), num(p["cur_v"]), *hi(r)]); r += 1
            p1 = r - 1
            sums = [f"=SUM({L}{p0}:{L}{p1})" for L in "DEFG"] if p1 >= p0 else [0, 0, 0, 0]
            put(ws, cr, S[7], [f'=IF(J{cr}<>"",COUNTIF(J${first}:J{cr},J{cr}),"")', _inn_val(c["inn"]), c["name"], *sums, *hi(cr), stf(cr)])
            if g["key"] == G_NEW: color(ws, cr, (10,), green)
            elif g["key"] in (G_DOWN, G_STOP): color(ws, cr, (8, 9, 10), red)
            comp_rows.append((g["key"], cr))
        last = r - 1; st = g["status"]
        if last >= first:
            vals = [g["roman"], None, f'="{g["title"]} ("&COUNTIF(J{first}:J{last},"{st}")&" та)"',
                    *[f'=SUMIF(J{first}:J{last},"{st}",{L}{first}:{L}{last})' for L in "DEFG"], *hi(gr)]
        else:
            vals = [g["roman"], None, f"{g['title']} (0 та)", 0, 0, 0, 0, *hi(gr)]
        put(ws, gr, S[6], vals)
    put(ws, 5, S[5], [None, None, "ЖАМИ", *["=" + "+".join(f"{L}{x}" for x in grow_s) for L in "DEFG"], *hi(5)])
    ws.auto_filter.ref = f"A4:J{r - 1}"; ws.print_area = f"A1:J{r - 1}"; ws.print_title_rows = "3:4"
    # ---- «Корхоналар» (havolalar «Соҳалар билан» varag'iga — shablondagidek)
    r = 6; grow_k = []; ref = "'Соҳалар билан'!"
    for g in d["groups"]:
        gr = r; grow_k.append(gr); wk.row_dimensions[gr].height = kh6; r += 1; first = r
        for key, sr in comp_rows:
            if key != g["key"]: continue
            hk = hi(r)
            if key == G_STOP: hk[0] = 0          # shablonda «Корхоналар» varag'ida to'xtaganlar o'sishi 0 deb yozilgan
            if key == G_NEW: hk[1] = "-"         # yangi korxona — foiz yo'q (shablondagidek «-»)
            put(wk, r, K[7], [f'=IF(J{r}<>"",COUNTIF(J${first}:J{r},J{r}),"")', *[f"={ref}{L}{sr}" for L in "BCDEFG"], *hk, stf(r)])
            if key == G_NEW:
                from openpyxl.styles import Alignment
                wk.cell(r, 9).alignment = Alignment(horizontal="center", vertical="center")
            if key == G_NEW: color(wk, r, (10,), green)
            elif key in (G_DOWN, G_STOP): color(wk, r, (8, 9, 10), red)
            r += 1
        last = r - 1; st = g["status"]
        if last >= first:
            vals = [g["roman"], None, f'="{g["title"]} ("&COUNTIF(J{first}:J{last},"{st}")&" та)"',
                    *[f'=SUMIF(J{first}:J{last},"{st}",{L}{first}:{L}{last})' for L in "DEFG"], *hi(gr)]
        else:
            vals = [g["roman"], None, f"{g['title']} (0 та)", 0, 0, 0, 0, *hi(gr)]
        put(wk, gr, K[6], vals)
    put(wk, 5, K[5], [None, None, f"ЖАМИ корхоналар ({d['count']} та)", *["=" + "+".join(f"{L}{x}" for x in grow_k) for L in "DEFG"], *hi(5)])
    wk.auto_filter.ref = f"A4:J{r - 1}"; wk.print_area = f"A1:J{r - 1}"; wk.print_title_rows = "3:4"; wk.freeze_panes = "D5"
    try:
        from openpyxl.workbook.properties import CalcProperties
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
    except Exception: pass
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
